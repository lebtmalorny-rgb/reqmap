"""Context is full, local spans stay local, and every visible copy is checked."""
from copy import deepcopy
from dataclasses import replace
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from reqmap.agent_service import AgentService
from reqmap.agent_knowledge import VerifiedKnowledge
from reqmap.agent_finalize import build_agent_run
from reqmap.binding_catalog import knowledge_digest
from reqmap.config import AnalysisProfile
from reqmap.export_json import canonical_json_bytes
from reqmap.knowledge import KnowledgeBase
from reqmap.models import to_dict
from reqmap.source_context import capture_source_document
from tests.agent_support import make_agent_config
from tests.source_context_support import DETACHED_TEXT, text_document, map_for, link_for, sign_test_map
from tests.test_binding_sessions import select_source
from tests.test_binding_mapping import setup_binding, kb


def context_report(tmp_path, monkeypatch, setup_binding, *, interface='api', quote='через API', duration=False):
    knowledge, catalog, proposal = setup_binding
    deep = type(knowledge) is not KnowledgeBase
    config = replace(make_agent_config(tmp_path/'agent'), analysis_profile=AnalysisProfile.DEEP if deep else AnalysisProfile.LEGACY)
    doc = capture_source_document(**text_document((DETACHED_TEXT, quote), 'agent-texts'))
    links = [link_for(doc, value=interface)]
    if duration:
        links.append(link_for(doc, field='constraint', value=dict(name='duration', operator='within', value='1', unit='ms'), link_id='duration'))
    raw = map_for(doc, disposition='linked', links=links)
    raw['rows'].append(map_for(doc, target=1, reviewed=False)['rows'][0])
    path = tmp_path/'map.json';path.write_bytes(canonical_json_bytes(raw))
    trust = sign_test_map(path) if deep else None
    config = replace(config, source_context_path=path, knowledge_trust=trust,
                     limits=replace(config.limits, max_response_bytes=8*1024*1024))
    monkeypatch.setattr('reqmap.agent_service.load_agent_knowledge', lambda config: VerifiedKnowledge(knowledge, knowledge_digest(knowledge), knowledge.snapshot_id if deep else None, catalog))
    svc = AgentService(config)
    started = svc.call('reqmap_start_session', dict(request_id='start', source=dict(kind='texts', texts=[r.text for r in doc.requirements])))
    assert started.ok, started
    sid = started.data['session_id']
    ctx = select_source(svc, sid, DETACHED_TEXT)
    assert ctx.ok, ctx
    proposal = {**proposal, 'proposal_schema_version':2, 'obligation_id':'REQ-0001-O001', 'predicate_ids':['P-NOVA-CREATE']}
    reply = svc.call('reqmap_submit_mapping', dict(session_id=sid,request_id='map',expected_revision=1,
                     atom_id='REQ-0001-A001',context_id=ctx.data['context_id'],proposal=proposal))
    assert reply.ok, reply
    view, verified = svc.verified_view(svc.store.read(sid))
    run = build_agent_run(view, verified)
    final = svc.call('reqmap_finalize', dict(session_id=sid,request_id='finish',expected_revision=2,allow_partial=True))
    assert final.ok, final
    return run, Path(final.data['artifacts']['result.json']).parent, svc, sid


def restored_cell(book, sheet, header):
    headers = [c.value for c in book[sheet][1]]
    value = book[sheet].cell(2, headers.index(header)+1).value
    ref = json.loads(value)
    if type(ref) is dict and 'xlsx_text_key' in ref:
        chunks = {r[0].value:r[1].value for r in list(book['Запуск'].rows)[1:]}
        value = ''.join(chunks[f"{ref['xlsx_text_key']}:{i}"] for i in range(1,ref['parts']+1))
        assert hashlib.sha256(value.encode()).hexdigest() == ref['sha256']
    return json.loads(value)


def checker(run, output):
    from reqmap.crosscheck import crosscheck
    from reqmap.crosscheck_deep import crosscheck_deep
    return (crosscheck_deep if run.schema_version=='2.2' else crosscheck)(run, output/'result.json', output/'result.xlsx', output/'report.md')


@pytest.mark.parametrize('duration', [False, True])
def test_five_artifacts_preserve_context_and_local_spans(tmp_path,monkeypatch,setup_binding,duration):
    run, output, _, _ = context_report(tmp_path,monkeypatch,setup_binding,duration=duration)
    from reqmap.binding_export import source_context_display, source_context_summary
    row = run.requirements[0]
    atom = row.atom_results[0]
    display = source_context_display(row.source_binding.context_decision, atom.atom.obligation_id)
    assert display['own_quote'] == DETACHED_TEXT
    assert display['effective_interface'] == 'api'
    assert display['origins'][0]['origin']['row']['coordinate']['row'] == 2
    assert display['origins'][0]['origin']['span']['quote'] == 'через API'
    assert atom.support_status.value == ('insufficient_evidence' if duration else 'supported')
    assert all(DETACHED_TEXT[s.start:s.end] == s.quote for s in (*atom.atom.source_spans,*atom.binding_decision.uncovered))
    payload=json.loads((output/'result.json').read_bytes())
    assert len(payload['metadata']['source_context']['loaded']['mapping']['rows']) == 2
    book=load_workbook(output/'result.xlsx')
    assert len(book.sheetnames)==(7 if run.schema_version=='2.2' else 5)
    assert restored_cell(book,'Атомарные утверждения','Source context')==display
    book.close()
    assert canonical_json_bytes(display).decode().strip() in (output/'report.md').read_text()
    summary=source_context_summary(run)
    manifest=json.loads((output/'manifest.json').read_bytes())
    assert manifest['source_context']==summary
    logs=[json.loads(x) for x in (output/'run.jsonl').read_text().splitlines()]
    assert any(x.get('source_context')==summary for x in logs)
    for name, sha in manifest['artifact_hashes'].items():
        assert hashlib.sha256((output/name).read_bytes()).hexdigest()==sha
    assert not checker(run,output)


@pytest.mark.parametrize('artifact', ['json','xlsx','markdown'])
def test_export_tamper_with_same_status_fails_crosscheck(tmp_path,monkeypatch,setup_binding,artifact):
    run,output,_,_=context_report(tmp_path,monkeypatch,setup_binding,duration=True)
    if artifact=='json':
        path=output/'result.json';raw=json.loads(path.read_bytes())
        raw['metadata']['source_context']['loaded']['map_sha256']='0'*64
        path.write_bytes(canonical_json_bytes(raw))
    elif artifact=='xlsx':
        book=load_workbook(output/'result.xlsx');sheet=book['Атомарные утверждения']
        headers=[c.value for c in sheet[1]]
        sheet.cell(2,headers.index('Source context')+1).value='{}'
        book.save(output/'result.xlsx');book.close()
    else:
        from reqmap.binding_export import source_context_display
        path=output/'report.md';display=source_context_display(run.requirements[0].source_binding.context_decision,'REQ-0001-O001')
        text=canonical_json_bytes(display).decode().strip()
        assert text in path.read_text()
        path.write_text(path.read_text().replace(text,'{}',1))
    assert checker(run,output)


def test_snapshot_rows_cannot_be_relabelled_in_export(tmp_path,monkeypatch,setup_binding):
    run,_,_,_=context_report(tmp_path,monkeypatch,setup_binding)
    from reqmap.binding_export import validate_binding_run
    row=run.requirements[0]
    forged=replace(run,requirements=(replace(row,requirement=replace(row.requirement,source_id='forged')),*run.requirements[1:]))
    with pytest.raises(ValueError,match='SOURCE_CONTEXT'):
        validate_binding_run(forged)


def test_long_and_formula_like_context_quotes(tmp_path,monkeypatch,setup_binding):
    quote='=HYPERLINK("https://example.invalid")'+('😀контекст'*5000)
    run,output,svc,sid=context_report(tmp_path,monkeypatch,setup_binding,quote=quote,interface='gui')
    from reqmap.binding_export import source_context_display
    display=source_context_display(run.requirements[0].source_binding.context_decision,'REQ-0001-O001')
    book=load_workbook(output/'result.xlsx')
    assert restored_cell(book,'Атомарные утверждения','Source context')==display
    assert display['origins'][0]['origin']['span']['quote']==quote
    assert all(c.data_type!='f' for sheet in book for row in sheet for c in row)
    assert all(len(c.value.encode('utf-16-le'))<=32767*2 for sheet in book for row in sheet for c in row if isinstance(c.value,str))
    book.close()
    assert not checker(run,output)
    small=AgentService(replace(svc.config,limits=replace(svc.config.limits,max_response_bytes=8192)))
    reply=small.call('reqmap_get_result',dict(session_id=sid))
    assert reply.error.code=='RESPONSE_TOO_LARGE'
