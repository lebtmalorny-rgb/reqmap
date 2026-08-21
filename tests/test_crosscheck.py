"""Cross-artifact проверка JSON, XLSX и Markdown относительно RunResult."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from reqmap.crosscheck import crosscheck
from reqmap.export_json import write_canonical_json
from reqmap.export_markdown import write_markdown
from reqmap.export_xlsx import write_xlsx
from tests.test_export_xlsx import mixed_run


def write_artifacts(tmp_path: Path):
    run = mixed_run()
    json_path = tmp_path / "result.json"
    xlsx_path = tmp_path / "result.xlsx"
    markdown_path = tmp_path / "report.md"
    write_canonical_json(run, json_path)
    write_xlsx(run, xlsx_path)
    write_markdown(run, markdown_path)
    return run, json_path, xlsx_path, markdown_path


def issue_codes(issues) -> set[str]:
    return {item.code for item in issues}


def test_crosscheck_accepts_consistent_artifacts(tmp_path: Path) -> None:
    run, json_path, xlsx_path, markdown_path = write_artifacts(tmp_path)

    assert crosscheck(run, json_path, xlsx_path, markdown_path) == ()


def test_crosscheck_reports_deleted_xlsx_mapping_row(tmp_path: Path) -> None:
    run, json_path, xlsx_path, markdown_path = write_artifacts(tmp_path)
    workbook = load_workbook(xlsx_path)
    workbook["Сопоставления"].delete_rows(2)
    workbook.save(xlsx_path)
    workbook.close()

    issues = crosscheck(run, json_path, xlsx_path, markdown_path)

    assert "CROSSCHECK_MAPPING_COUNT" in issue_codes(issues)


def test_crosscheck_compares_ids_statuses_and_run_status(tmp_path: Path) -> None:
    run, json_path, xlsx_path, markdown_path = write_artifacts(tmp_path)
    workbook = load_workbook(xlsx_path)
    requirement_headers = {
        cell.value: cell.column for cell in workbook["Требования"][1]
    }
    workbook["Требования"].cell(
        row=2,
        column=requirement_headers["Поддержка: код"],
    ).value = "partial"
    workbook["Атомарные утверждения"].cell(row=2, column=1).value = "forged-atom"
    for row in workbook["Запуск"].iter_rows(min_row=2):
        if row[0].value == "run_status":
            row[1].value = "FAILED"
    workbook.save(xlsx_path)
    workbook.close()

    codes = issue_codes(crosscheck(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_XLSX_IDS" in codes
    assert "CROSSCHECK_XLSX_STATUSES" in codes
    assert "CROSSCHECK_XLSX_RUN_STATUS" in codes


def test_crosscheck_reports_json_content_and_hash_changes(tmp_path: Path) -> None:
    run, json_path, xlsx_path, markdown_path = write_artifacts(tmp_path)
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    payload["run_status"] = "FAILED"
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )

    codes = issue_codes(crosscheck(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_JSON_RUN_STATUS" in codes
    assert "CROSSCHECK_JSON_HASH" in codes


def test_crosscheck_reports_markdown_marker_and_hash_changes(
    tmp_path: Path,
) -> None:
    run, json_path, xlsx_path, markdown_path = write_artifacts(tmp_path)
    text = markdown_path.read_text(encoding="utf-8")
    markdown_path.write_text(
        text.replace('"mappings":2', '"mappings":99'),
        encoding="utf-8",
    )

    codes = issue_codes(crosscheck(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_MARKDOWN_COUNTS" in codes
    assert "CROSSCHECK_MARKDOWN_HASH" in codes


def test_crosscheck_detects_xlsx_hash_change_without_data_change(
    tmp_path: Path,
) -> None:
    run, json_path, xlsx_path, markdown_path = write_artifacts(tmp_path)
    workbook = load_workbook(xlsx_path)
    workbook["Запуск"].sheet_view.showGridLines = True
    workbook.save(xlsx_path)
    workbook.close()

    codes = issue_codes(crosscheck(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_XLSX_HASH" in codes
    assert "CROSSCHECK_XLSX_VALUES" not in codes


def test_crosscheck_returns_issues_for_unreadable_artifacts(tmp_path: Path) -> None:
    run, json_path, xlsx_path, markdown_path = write_artifacts(tmp_path)
    json_path.write_bytes(b"\xff")
    xlsx_path.write_bytes(b"not-an-xlsx")
    markdown_path.write_bytes(b"\xff")

    codes = issue_codes(crosscheck(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_JSON_READ" in codes
    assert "CROSSCHECK_XLSX_VALUES" in codes
    assert "CROSSCHECK_XLSX_READ" in codes
    assert "CROSSCHECK_MARKDOWN_READ" in codes


def test_crosscheck_rejects_intermediate_symlink_without_following_it(
    tmp_path: Path,
) -> None:
    real_output = tmp_path / "real-output"
    real_output.mkdir()
    run, json_path, xlsx_path, markdown_path = write_artifacts(real_output)
    before = {
        path.name: path.read_bytes()
        for path in (json_path, xlsx_path, markdown_path)
    }
    linked_output = tmp_path / "linked-output"
    linked_output.symlink_to(real_output, target_is_directory=True)

    codes = issue_codes(
        crosscheck(
            run,
            linked_output / "result.json",
            linked_output / "result.xlsx",
            linked_output / "report.md",
        )
    )

    assert {
        "CROSSCHECK_JSON_READ",
        "CROSSCHECK_XLSX_READ",
        "CROSSCHECK_MARKDOWN_READ",
    }.issubset(codes)
    assert before == {
        path.name: path.read_bytes()
        for path in (json_path, xlsx_path, markdown_path)
    }


def test_crosscheck_does_not_modify_read_only_artifact_directory(
    tmp_path: Path,
) -> None:
    output = tmp_path / "read-only-output"
    output.mkdir()
    run, json_path, xlsx_path, markdown_path = write_artifacts(output)
    before = {
        path.name: path.read_bytes()
        for path in (json_path, xlsx_path, markdown_path)
    }
    output.chmod(0o500)
    try:
        assert crosscheck(run, json_path, xlsx_path, markdown_path) == ()
    finally:
        output.chmod(0o700)
    assert before == {
        path.name: path.read_bytes()
        for path in (json_path, xlsx_path, markdown_path)
    }
