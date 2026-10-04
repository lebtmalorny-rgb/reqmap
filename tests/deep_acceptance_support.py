"""Synthetic inputs and HTTP replies, independent of frozen expected outcomes.

All claims below are fabricated contract fixtures. They do not establish any
OpenStack/Kolla/Rocky capability. Keys live only under pytest temporary paths.
"""
from contextlib import contextmanager
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import shutil
import socket
from threading import Thread

from openpyxl import load_workbook
import pytest

from reqmap.cli import main
from tests.deep_factories import sign_existing_v2_snapshot
from tests.test_acceptance import _FakeOpenAIServer

ROOT = Path(__file__).resolve().parents[1]
SCOPE = dict(source_release="2025.1", target_release="2025.1",
             kolla_ansible_release="2025.1", host_profile="rocky_linux_9",
             version_constraint="2025.1")
# name: component, contour, target, lifecycle, search terms
ACTIONS = {
    "CREATE": ("nova", "openstack_runtime", "SERVER", "runtime", "Создание сервера Nova API"),
    "RECONFIGURE": ("kolla_ansible", "kolla_ansible", "CONFIG", "reconfigure", "Применить конфигурацию сервисов Kolla-Ansible reconfigure"),
    "HOST": ("rocky_linux_9", "host_os", "SYSCTL", "reconfigure", "Изменить параметр ядра Rocky Linux вручную"),
    "SYSCTL": ("kolla_ansible", "kolla_ansible", "SYSCTL", "reconfigure", "Применить sysctl Kolla-Ansible"),
    "HA": ("nova", "openstack_runtime", "SERVER", "recover", "Восстановить HA виртуальную машину после отказа"),
    "UPGRADE": ("kolla_ansible", "kolla_ansible", "CONFIG", "upgrade", "Обновить upgrade OpenStack 2025.1 2026.1"),
    "ROLLBACK": ("nova", "openstack_runtime", "SERVER", "runtime", "Создание сервера rollback откат"),
    "NEGATIVE": ("nova", "openstack_runtime", "SERVER", "runtime", "Запрещённое создание сервера forbidden"),
    "CONFLICT": ("nova", "openstack_runtime", "SERVER", "runtime", "Противоречивое создание сервера conflict"),
}
for _kind, _label in (
    ("VM", "ВМ"), ("VOLUME", "том"), ("NETWORK", "сеть"), ("IMAGE", "образ"),
    ("CONTROL-PLANE", "control-plane"), ("DATABASE", "базу данных"),
):
    ACTIONS[f"MIGRATE-{_kind}"] = (
        "nova", "openstack_runtime", "SERVER", "migrate", f"Перенести {_label} migrate-{_kind.lower()}"
    )


def _json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")


def _jsonl(path, values):
    path.write_text("".join(json.dumps(value, ensure_ascii=False, sort_keys=True) + "\n" for value in values), encoding="utf-8")


def gold_snapshot(tmp_path):
    root = tmp_path / "snapshot"
    shutil.copytree(ROOT / "tests/fixtures/kb_v2_minimal", root)
    _json(root / "components.json", {"components": [
        dict(id=component, display_name=component, kind=kind, releases=["2025.1"])
        for component, kind in (("nova", "openstack_service"), ("kolla_ansible", "deployment_tool"), ("rocky_linux_9", "host_os_subsystem"))
    ]})
    _json(root / "actors.json", {"actors": [
        dict(id=f"ACTOR-{component}", component_ref=component, kind="automation")
        for component in ("nova", "kolla_ansible", "rocky_linux_9")
    ]})
    _jsonl(root / "targets.jsonl", [
        dict(id="TARGET-SERVER", component_ref="nova", contour="openstack_runtime", host_profile=None),
        dict(id="TARGET-CONFIG", component_ref="kolla_ansible", contour="kolla_ansible", host_profile=None),
        dict(id="TARGET-SYSCTL", component_ref="rocky_linux_9", contour="host_os", host_profile="rocky_linux_9"),
    ])
    capabilities, actions, effects, evidence, procedures = [], [], [], [], []
    for name, (component, contour, target, phase, terms) in ACTIONS.items():
        ids = [f"EV-{name}"] + (["EV-CONFLICT-NEGATIVE"] if name == "CONFLICT" else ["EV-NEGATIVE-POSITIVE"] if name == "NEGATIVE" else [])
        scope = {**SCOPE, **(dict(target_release="2026.1", version_constraint="2026.1") if name == "UPGRADE" else {})}
        capabilities.append(dict(id=f"CAP-{name}", component_ref=component, name_ru=terms, terms=terms.split(), evidence_ids=ids))
        actions.append(dict(id=f"ACTION-{name}", component_ref=component, contour=contour,
                            target_ref=f"TARGET-{target}", effect_refs=[f"EFFECT-{name}"],
                            interface_type={"openstack_runtime": "openstack_api", "kolla_ansible": "kolla_ansible_action", "host_os": "host_command"}[contour],
                            operation=f"synthetic-{name.lower()}", version_scope=scope,
                            procedure_required=name not in {"CREATE", "HOST", "RECONFIGURE", "NEGATIVE", "CONFLICT"}, evidence_ids=ids))
        effects.append(dict(id=f"EFFECT-{name}", target_ref=f"TARGET-{target}", before_state="before", after_state="after", verification_criteria=["synthetic state observed"], reversible=name != "UPGRADE", evidence_ids=ids))
        for evidence_id in ids:
            evidence.append(dict(id=evidence_id, claim=f"Synthetic contract: {terms}", claim_kind="capability",
                                 polarity="negative" if evidence_id == "EV-NEGATIVE" or evidence_id == "EV-CONFLICT-NEGATIVE" else "positive",
                                 strength="direct", source_id="SRC-MINIMAL", locator=f"SRC-MINIMAL:{name}",
                                 version_constraint="2026.1" if evidence_id == "EV-NEGATIVE-POSITIVE" else scope["version_constraint"], applicable_contours=[contour] + (["host_os"] if name == "SYSCTL" else []),
                                 supports_entity_refs=[f"CAP-{name}", f"ACTION-{name}", f"EFFECT-{name}"],
                                 local_excerpt=f"Synthetic contract: {terms}", review_state="reviewed"))
        if name in {"SYSCTL", "UPGRADE", "ROLLBACK"}:
            steps = []
            phases = [("before", "preflight", []), ("action", phase, ["before"]), ("verify", "verify", ["action"])]
            if name == "UPGRADE":
                phases = [("action", "upgrade", [])]
            elif name != "ROLLBACK":
                phases.append(("rollback", "rollback", []))
            for local_id, step_phase, dependencies in phases:
                steps.append(dict(local_step_id=local_id, phase=step_phase, contour=contour,
                                  executor_ref=f"ACTOR-{component}", target_ref=f"TARGET-{target}", action_ref=f"ACTION-{name}",
                                  preconditions=["synthetic precondition"], success_criteria=["synthetic criterion"],
                                  evidence_ids=ids, depends_on=dependencies,
                                  rollback_step_local_id="rollback" if local_id == "action" and name not in {"ROLLBACK", "UPGRADE"} else None))
            procedures.append(dict(id=f"PROC-{name}", lifecycle_phase=phase, action_refs=[f"ACTION-{name}"], steps=steps))
    for filename, values in (("capabilities", capabilities), ("actions", actions), ("effects", effects), ("evidence", evidence), ("procedures", procedures)):
        _jsonl(root / f"{filename}.jsonl", values)
    # Locators refer to explicit synthetic sections, never downloaded upstream content.
    corpus = root / "corpus/SRC-MINIMAL.md"
    corpus.write_text("# Synthetic contract fixture; not upstream evidence\n\n" + "\n".join(f"## {name}\nSynthetic contract: {item[4]}\n" for name, item in ACTIONS.items()), encoding="utf-8")
    manifest = json.loads((root / "source-manifest.json").read_text())
    manifest["sources"][0]["content_sha256"] = hashlib.sha256(corpus.read_bytes()).hexdigest()
    _json(root / "source-manifest.json", manifest)
    _json(root / "synonyms.json", {})
    return root, sign_existing_v2_snapshot(root, tmp_path / "trust")


def _selection(name):
    component, contour, target, phase, _terms = ACTIONS[name]
    scope = {**SCOPE, **(dict(target_release="2026.1", version_constraint="2026.1") if name == "UPGRADE" else {})}
    return dict(contour=contour, component_ref=component, executor_ref=f"ACTOR-{component}",
                target_contour="host_os" if target == "SYSCTL" else contour, target_ref=f"TARGET-{target}",
                action_ref=f"ACTION-{name}", effect_ref=f"EFFECT-{name}", lifecycle_phase=phase,
                version_scope=scope, evidence_ids=[f"EV-{name}"], support_status="supported", related_indexes=[])


def scripted_response(payload):
    if "original_payload" in payload:
        return scripted_response(payload["original_payload"])
    if "requirement_text" in payload:
        text = payload["requirement_text"]
        quotes = text.split(" и ") if text.startswith("Создание сервера через Nova API и ") else [text]
        return {"atoms": [dict(text=quote, source_quote=quote, mandatory=True) for quote in quotes]}
    text = payload["atom"]["source_quote"]
    response = dict(support_status="supported", supported_aspects=[text], unconfirmed_aspects=[], responsibilities=[], procedure_template_ids=[])
    if text in {"Обеспечить работу системы.", "Поддержать квантовый режим X."}:
        return {**response, "support_status": "insufficient_evidence", "supported_aspects": [], "unconfirmed_aspects": [text]}
    if "migrate-" in text:
        name = "MIGRATE-" + text.split("migrate-", 1)[1].rstrip(".").upper()
    elif "upgrade" in text:
        name = "UPGRADE"
    elif "sysctl" in text:
        name = "SYSCTL"
    elif "вручную" in text:
        name = "HOST"
    elif "reconfigure" in text:
        name = "RECONFIGURE"
    elif "HA" in text:
        name = "HA"
    elif "rollback" in text:
        name = "ROLLBACK"
    elif "forbidden" in text:
        name = "NEGATIVE"
    elif "conflict" in text:
        name = "CONFLICT"
    else:
        name = "CREATE"
    record = _selection(name)
    if "без доказательств" in text:
        record["evidence_ids"] = []
    if "runtime 2026.1" in text:
        record["version_scope"]["target_release"] = "2026.1"
    # The conflict reply deliberately hides the negative position; the gate must recover it.
    response["responsibilities"] = [record]
    if name == "SYSCTL":
        host = deepcopy(record)
        host.update(contour="host_os", component_ref="rocky_linux_9", related_indexes=[1])
        record["related_indexes"] = [2]
        response["responsibilities"].append(host)
    if name in {"SYSCTL", "UPGRADE", "ROLLBACK"}:
        response["procedure_template_ids"] = [f"PROC-{name}"]
    return response


@contextmanager
def model_server():
    server = _FakeOpenAIServer(scripted_response)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    attempted = []
    original_connect = socket.socket.connect
    allowed_address = server.server_address

    def guarded_connect(sock, address):
        attempted.append(address)
        assert address == allowed_address, f"External connection forbidden in acceptance: {address}"
        return original_connect(sock, address)

    try:
        with pytest.MonkeyPatch.context() as patch:
            patch.setattr(socket.socket, "connect", guarded_connect)
            yield server
            assert server.errors == [], server.errors
            assert all(address == allowed_address for address in attempted)
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def cli_arguments(tmp_path, root, allowed_signers, server, texts):
    config = tmp_path / "config.json"
    _json(config, dict(analysis_profile="deep", model=dict(base_url=server.base_url, model=server.model_name, timeout_seconds=5, retries=0, seed=19),
                       knowledge_path=str(root), knowledge_trust=dict(allowed_signers_path=str(allowed_signers)), top_k=50))
    arguments = ["analyze", "--config", str(config), "--output", str(tmp_path / "out")]
    for text in texts:
        arguments.extend(["--requirement", text])
    return arguments


def run_gold(tmp_path):
    cases = json.loads((ROOT / "tests/fixtures/deep_gold.json").read_text(encoding="utf-8"))["cases"]
    root, signers = gold_snapshot(tmp_path)
    with model_server() as server:
        code = main(cli_arguments(tmp_path, root, signers, server, [case["requirement"] for case in cases]))
        assert code == 4, (code, server.errors)
        requests = list(server.requests)
    output = tmp_path / "out"
    result = json.loads((output / "result.json").read_text(encoding="utf-8"))
    return cases, result, output, requests


def assert_artifacts(output, result):
    names = {"result.json", "result.xlsx", "report.md", "run.jsonl", "manifest.json"}
    assert {path.name for path in output.iterdir() if path.is_file()} == names
    manifest = json.loads((output / "manifest.json").read_text())
    assert manifest["run_id"] == result["run_id"]
    assert manifest["run_status"] == result["run_status"]
    assert manifest["artifact_hashes"] == {name: hashlib.sha256((output / name).read_bytes()).hexdigest() for name in names - {"manifest.json"}}
    workbook = load_workbook(output / "result.xlsx", read_only=True)
    try:
        expected = {
            "Требования": [item["requirement"]["requirement_id"] for item in result["requirements"]],
            "Атомарные утверждения": [atom["atom"]["atom_id"] for item in result["requirements"] for atom in item["atom_results"]],
            "Ответственность": [item["record_id"] for item in result["responsibility_records"]],
            "Доказательства": [item["evidence_id"] for item in result["evidence"]],
        }
        for sheet, ids in expected.items():
            assert [row[0] for row in workbook[sheet].iter_rows(min_row=2, values_only=True)] == ids
        for row, item in zip(workbook["Требования"].iter_rows(min_row=2, values_only=True), result["requirements"], strict=True):
            assert row[6] == item["requirement"]["text"]
            assert row[11] == item["analysis_state"]
            assert row[12] == (item["support_status"] or "") or (row[12] is None and item["support_status"] is None)
        for row, item in zip(workbook["Ответственность"].iter_rows(min_row=2, values_only=True), result["responsibility_records"], strict=True):
            assert row[3:8] == (item["contour"], item["component_ref"], item["executor_ref"], item["target_contour"], item["target_ref"])
    finally:
        workbook.close()
    report = (output / "report.md").read_text(encoding="utf-8")
    for item in result["responsibility_records"]:
        assert item["record_id"] in report
        assert item["contour"] in report
    for item in result["requirements"]:
        assert item["requirement"]["requirement_id"] in report
        for diagnostic in item["diagnostics"]:
            assert diagnostic in report
