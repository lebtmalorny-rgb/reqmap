"""Explicit synthetic source documents; no implicit approval of real inputs."""
from io import BytesIO
from functools import lru_cache
from pathlib import Path
import tempfile
import subprocess

from openpyxl import Workbook

from reqmap.export_json import canonical_json_bytes
from reqmap.ids import generated_requirement_id
from reqmap.models import Requirement, SourceCoordinate


API_TEXT = "Nova должна создавать ВМ через API"
DETACHED_TEXT = "Nova должна создавать ВМ"


def text_document(texts=(API_TEXT,), source_name="synthetic-texts"):
    rows = tuple(Requirement(generated_requirement_id(i), None, text, i,
                            SourceCoordinate(source_name, None, i))
                 for i, text in enumerate(texts, 1))
    return dict(content=canonical_json_bytes(list(texts)), source_kind="texts",
                source_name=source_name, requirements=rows, input_profile=None)


def workbook_bytes():
    book = Workbook()
    for i, name in enumerate(("Первый", "Второй")):
        sheet = book.active if i == 0 else book.create_sheet()
        sheet.title = name
        sheet.append(["ID", "Требование", "Условие", "Подсказка"])
        sheet.append(["same-id", "😀GUI и GUI", "при отказе узла", "контекст"])
        sheet.append(["same-id", DETACHED_TEXT, "через GUI", "Nova"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def map_for(document, *, disposition="independent", target=0, links=(), reviewed=True):
    """Approve only the explicitly selected synthetic row, never all inputs."""
    import hashlib
    from reqmap.models import to_dict
    row = document.requirements[target]
    ref = dict(requirement_id=row.requirement_id, coordinate=to_dict(row.coordinate),
               source_sha256=hashlib.sha256(row.text.encode()).hexdigest())
    entry = dict(target=ref, disposition=disposition,
                 review_state="reviewed" if reviewed else "draft",
                 reviewed_by="synthetic-maintainer" if reviewed else None,
                 reviewed_at="2026-10-08T00:00:00Z" if reviewed else None,
                 reason_ru="Проверка синтетического примера целиком",
                 context_complete=reviewed and disposition != "unresolved", links=list(links))
    return dict(schema_version="1.0", map_id="synthetic-map",
                source_kind=document.source_kind, source_name=document.source_name,
                input_sha256=document.input_sha256, input_profile_sha256=document.input_profile_sha256,
                requirements_sha256=document.requirements_sha256, rows=[entry])


def link_for(document, origin=1, *, field="interface", value="gui", link_id="link-1", start=0, end=None):
    text = document.requirements[origin].text
    end = len(text) if end is None else end
    return dict(link_id=link_id, field=field, value=value,
                origin=dict(row=map_for(document, target=origin)["rows"][0]["target"],
                            span=dict(start=start, end=end, quote=text[start:end])))


@lru_cache(maxsize=128)
def document_for_binding(binding):
    """Reproduce only this explicitly supplied synthetic row's actual coordinate."""
    from reqmap.config import InputProfile
    from reqmap.input_xlsx import load_xlsx_bytes
    from reqmap.source_context import capture_source_document
    if binding.coordinate.sheet is None:
        data=text_document((binding.source_text,),binding.coordinate.source_name)
        return capture_source_document(**data)
    book=Workbook();sheet=book.active;sheet.title=binding.coordinate.sheet
    sheet.append(['ID','Требование'])
    for _ in range(2,binding.coordinate.row):sheet.append([None,None])
    sheet.append(['source-row-1',binding.source_text])
    stream=BytesIO();book.save(stream);book.close();content=stream.getvalue()
    profile=__import__('reqmap.config',fromlist=['InputProfile']).InputProfile((binding.coordinate.sheet,),(),1,'ID','Требование')
    rows=load_xlsx_bytes(content,profile,binding.coordinate.source_name)
    assert rows[0].requirement_id==binding.requirement_id
    return capture_source_document(content=content,source_kind='xlsx',source_name=binding.coordinate.source_name,input_profile=profile,requirements=rows)


_temporary_trust = []


def sign_test_map(path, trust=None):
    from reqmap.config import KnowledgeTrustConfig
    if trust is None:
        directory=tempfile.TemporaryDirectory(prefix='reqmap-synthetic-context-')
        _temporary_trust.append(directory)
        root=Path(directory.name).resolve();key=root/'signing_key'
        subprocess.run(['ssh-keygen','-q','-t','ed25519','-N','','-f',str(key)],check=True,capture_output=True)
        allowed=root/'allowed_signers'
        allowed.write_text('reqmap-snapshot '+key.with_suffix('.pub').read_text())
        trust=KnowledgeTrustConfig(allowed)
    else:
        key=trust.allowed_signers_path.parent/'signing_key'
    signature=Path(str(path)+'.sig')
    if signature.exists():signature.unlink()
    subprocess.run(['ssh-keygen','-Y','sign','-f',str(key),'-n','reqmap-source-context',str(path)],check=True,capture_output=True)
    return trust


def reviewed_context(binding, catalog):
    """Explicit test-call approval; production callers never invoke this helper."""
    from reqmap.config import AnalysisProfile
    from reqmap.binding_runtime import requirement_context
    from reqmap.source_context import load_source_context_map, freeze_source_context
    doc=document_for_binding(binding)
    raw=map_for(doc)
    directory=tempfile.TemporaryDirectory(prefix='reqmap-synthetic-map-')
    _temporary_trust.append(directory)
    path=Path(directory.name).resolve()/'map.json'
    path.write_bytes(canonical_json_bytes(raw))
    deep=catalog is not None and any(p.action_ref is not None for p in catalog.predicates.values())
    profile=AnalysisProfile.DEEP if deep else AnalysisProfile.LEGACY
    trust=sign_test_map(path) if deep else None
    frozen=freeze_source_context(doc,load_source_context_map(path,doc,profile=profile,trust=trust))
    return requirement_context(doc.requirements[0],catalog,frozen,source_context_trust=trust)


def with_reviewed_texts(config, texts):
    """Review exactly the caller's listed synthetic texts for an MCP fixture."""
    from dataclasses import replace
    from reqmap.config import AnalysisProfile
    from reqmap.source_context import capture_source_document
    doc=capture_source_document(**text_document(tuple(texts),'agent-texts'))
    raw=map_for(doc)
    raw['rows']=[map_for(doc,target=i)['rows'][0] for i in range(len(texts))]
    path=config.input_root/'reviewed-synthetic-map.json'
    path.write_bytes(canonical_json_bytes(raw))
    trust=config.knowledge_trust
    if config.analysis_profile is AnalysisProfile.DEEP:trust=sign_test_map(path,trust)
    return replace(config,source_context_path=path,knowledge_trust=trust)


def request_with_hint(tmp_path, text, hint):
    from reqmap.config import InputProfile
    from reqmap.input_xlsx import load_xlsx_bytes
    from reqmap.source_context import capture_source_document
    from reqmap.models import AnalysisRequest
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    sheet.append(["ID", "Requirement", "Hint"])
    sheet.append(["source-1", text, hint])
    stream = BytesIO()
    book.save(stream)
    book.close()
    content = stream.getvalue()
    profile = InputProfile(("Synthetic",), (), 1, "ID", "Requirement", hint_columns=("Hint",))
    rows = load_xlsx_bytes(content, profile, "synthetic.xlsx")
    document = capture_source_document(content=content, source_kind="xlsx", source_name="synthetic.xlsx",
                                       requirements=rows, input_profile=profile)
    return (AnalysisRequest(rows, document.input_sha256, "xlsx", None, tmp_path / "output", source_document=document), profile)


def with_reviewed_request(config, request):
    """Review exactly the synthetic document explicitly supplied by this test."""
    from dataclasses import replace
    from reqmap.config import AnalysisProfile
    document = request.source_document
    raw = map_for(document)
    raw["rows"] = [map_for(document, target=i)["rows"][0] for i in range(len(document.requirements))]
    path = request.output_dir.parent / "reviewed-synthetic-map.json"
    path.write_bytes(canonical_json_bytes(raw))
    trust = config.knowledge_trust
    if config.analysis_profile is AnalysisProfile.DEEP:
        trust = sign_test_map(path, trust)
    return replace(config, source_context_path=path, knowledge_trust=trust)
