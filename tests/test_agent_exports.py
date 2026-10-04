"""External proposals must have explicit, cross-artifact, unattested origin."""
import importlib
import json
from dataclasses import replace
import pytest
from openpyxl import load_workbook
from tests.test_export_deep_json import deep_run
from tests.test_json_manifest import completed_run


def origin():
    return dict(mode='external_agent',session_id='0d5c9c22-348f-4d4c-b9f4-ff03fe67db1e',revision=3,
                tool_contract_version='1.0',workflow_version='1.0',proposal_journal_sha256='a'*64,
                reported_client='codex',identity_verified=False)


def agent_run(deep=False):
    run = deep_run() if deep else completed_run()
    metadata = {**run.metadata,'model':'external-agent','seed':None,'analysis_origin':origin()}
    metadata.pop('endpoint_origin',None)
    return replace(run,metadata=metadata)


@pytest.mark.parametrize('deep',[False,True])
def test_origin_is_published_in_all_reports_and_detects_tampering(tmp_path,deep):
    assert importlib.util.find_spec('reqmap.publication'), 'shared publisher is missing'
    from reqmap.publication import publish_artifacts
    from reqmap.crosscheck import crosscheck
    from reqmap.crosscheck_deep import crosscheck_deep
    run = agent_run(deep)
    assert publish_artifacts(run,tmp_path) == ()
    canonical = json.loads((tmp_path/'result.json').read_text())
    manifest = json.loads((tmp_path/'manifest.json').read_text())
    assert canonical['metadata']['analysis_origin'] == origin() == manifest['analysis_origin']
    assert canonical['metadata']['model'] == 'external-agent'
    assert canonical['metadata']['seed'] is None
    assert 'endpoint_origin' not in manifest
    wb = load_workbook(tmp_path/'result.xlsx')
    rows = dict(wb['Запуск'].iter_rows(min_row=2,values_only=True))
    assert json.loads(rows['analysis_origin']) == origin()
    wb.close()
    path = tmp_path/'report.md'
    text = path.read_text()
    assert 'identity_verified' in text
    path.write_text(text.replace('"identity_verified":false','"identity_verified":true'))
    checker = crosscheck_deep if deep else crosscheck
    assert checker(run,tmp_path/'result.json',tmp_path/'result.xlsx',path)


@pytest.mark.parametrize('changes',[{'identity_verified':True},{'unknown':'field'},{'proposal_journal_sha256':'bad'},{'reported_model':'https://private.example/key'},{'reported_model':'sk-secret'}])
def test_untrusted_origin_is_rejected(changes):
    assert importlib.util.find_spec('reqmap.analysis_origin')
    from reqmap.analysis_origin import validate_analysis_origin
    with pytest.raises(ValueError):
        validate_analysis_origin({**origin(),**changes})


@pytest.mark.parametrize('deep',[False,True])
@pytest.mark.parametrize('changes',[{'seed':1},{'model':'pretend-verified-model'},{'endpoint_origin':'https://example.com'}])
def test_origin_cannot_claim_backend_model_control(tmp_path,deep,changes):
    from reqmap.export_json import write_canonical_json
    from reqmap.export_deep_json import write_deep_canonical_json
    run = agent_run(deep)
    with pytest.raises(ValueError):
        (write_deep_canonical_json if deep else write_canonical_json)(replace(run,metadata={**run.metadata,**changes}),tmp_path/'result.json')


def test_proposal_hash_ignores_reads_and_finalization_but_keeps_rejections(tmp_path):
    from tests.test_agent_store import store_setup, command, accepted
    assert importlib.util.find_spec('reqmap.analysis_origin')
    from reqmap.analysis_origin import proposal_journal_sha256, build_analysis_origin
    store,sid,seed,_ = store_setup(tmp_path)
    store.transact(command(sid),accepted)
    record = store.read(sid)
    first = proposal_journal_sha256(record)
    assert proposal_journal_sha256(store.read(sid)) == first
    assert build_analysis_origin(record,2)['identity_verified'] is False
    store.transact(command(sid,'stale'),accepted)
    assert proposal_journal_sha256(store.read(sid)) != first
