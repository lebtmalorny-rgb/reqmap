"""Explicit positional imports shared by CLI configuration and durable MCP input."""
from dataclasses import asdict, replace
import hashlib
import json

from openpyxl import Workbook
import pytest

from reqmap.config import load_config
from reqmap.config_common import parse_input_profile
from reqmap.errors import ConfigError, InputProfileError
from reqmap.input_xlsx import load_xlsx
from reqmap.models import SourceField, SourceHint
from tests.test_config import BASE_CONFIG, write_config
from tests.test_agent_config import agent_module, config_file


def positional(**changes):
    return dict(include_sheets=['First', 'Second'], column_mode='position',
                data_start_row=3, id_column='B', text_column='C', **changes)


def book(tmp_path):
    wb = Workbook()
    first = wb.active
    first.title = 'First'
    first['A1'] = 'Ignored title'
    first.append([])
    for row, text in ((3, ' First\nline '), (5, 'Continuation'), (6, 'Continuation')):
        first.cell(row, 2, 'repeat')
        first.cell(row, 3, text)
        first.cell(row, 4, 'P1')
        first.cell(row, 5, 'Expected')
        first.cell(row, 6, 'parent')
        first.cell(row, 7, 'Untrusted hint')
    second = wb.create_sheet('Second')
    second['B3'] = 'repeat'
    second['C3'] = 'Independent'
    second['G3'] = 'Other hint'
    path = tmp_path / 'synthetic.xlsx'
    wb.save(path)
    wb.close()
    return path


def test_position_profile_shared_by_cli_and_agent_and_roundtrips(tmp_path):
    raw = positional(priority_column='D', expected_result_column='E',
                     parent_column='F', hint_columns=['G'])
    cli = load_config(write_config(tmp_path, {**BASE_CONFIG, 'input_profile': raw}), {})
    agent = agent_module().load_agent_config(config_file(tmp_path, input_profile=raw))
    assert cli.input_profile == agent.input_profile
    profile = cli.input_profile
    assert profile.header_row is None
    assert profile.data_start_row == 3
    assert parse_input_profile(json.loads(json.dumps(asdict(profile)))) == profile


@pytest.mark.parametrize('changes', [
    {'column_mode': 'guess'}, {'column_mode': True}, {'column_mode': []},
    {'header_row': 1}, {'data_start_row': None}, {'data_start_row': True},
    {'data_start_row': 0}, {'data_start_row': 1048577},
    {'id_column': '1'}, {'text_column': 'c'}, {'text_column': 'XFE'},
    {'text_column': 'C3'}, {'text_column': '$C'}, {'priority_column': 'Priority'},
    {'expected_result_column': 'XFE'}, {'parent_column': '0'},
    {'hint_columns': ['D', 'wrong']},
])
def test_position_profile_rejects_ambiguous_or_invalid_addressing(changes):
    with pytest.raises(ConfigError):
        parse_input_profile({**positional(), **changes})


def test_position_profile_requires_explicit_start_and_legacy_does_not_mix_modes():
    raw = positional()
    del raw['data_start_row']
    with pytest.raises(ConfigError):
        parse_input_profile(raw)
    legacy = dict(include_sheets=['First'], header_row=1, id_column='ID', text_column='Text')
    profile = parse_input_profile(legacy)
    assert profile.column_mode == 'header'
    assert parse_input_profile(json.loads(json.dumps(asdict(profile)))) == profile
    with pytest.raises(ConfigError):
        parse_input_profile({**legacy, 'data_start_row': 3})


def test_position_import_preserves_duplicates_all_coordinates_and_fields(tmp_path):
    path = book(tmp_path)
    before = path.read_bytes()
    profile = parse_input_profile(positional(priority_column='D', expected_result_column='E',
                                              parent_column='F', hint_columns=['G']))
    rows = load_xlsx(path, profile)
    assert [(r.coordinate.sheet, r.coordinate.row, r.source_id, r.text) for r in rows] == [
        ('First', 3, 'repeat', ' First\nline '), ('First', 5, 'repeat', 'Continuation'),
        ('First', 6, 'repeat', 'Continuation'), ('Second', 3, 'repeat', 'Independent')]
    assert [r.requirement_id for r in rows] == ['REQ-0001', 'REQ-0002', 'REQ-0003', 'REQ-0004']
    assert rows[0].source_fields == (SourceField('D', 'P1'), SourceField('E', 'Expected'), SourceField('F', 'parent'))
    assert rows[0].source_hints == (SourceHint('Untrusted hint', 'G'),)
    assert all(r.coordinate.source_name == path.name for r in rows)
    assert rows[3].source_fields == tuple(SourceField(c, '') for c in ('D', 'E', 'F'))
    assert before == path.read_bytes()


def test_headerless_is_never_guessed_and_missing_position_is_diagnostic(tmp_path):
    path = book(tmp_path)
    with pytest.raises(InputProfileError) as error:
        load_xlsx(path, None)
    assert error.value.code == 'XLSX_PROFILE_AMBIGUOUS'
    profile = parse_input_profile({**positional(), 'text_column': 'AA'})
    with pytest.raises(InputProfileError) as error:
        load_xlsx(path, profile)
    assert error.value.code == 'XLSX_PROFILE_COLUMN_MISSING'


def test_position_import_uses_multiletter_columns_and_rejects_formulas(tmp_path):
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'First'
    sheet['Z3'] = 17
    sheet['AA3'] = 'Last column text'
    path = tmp_path / 'wide.xlsx'
    wb.save(path)
    profile = parse_input_profile({**positional(), 'include_sheets': ['First'], 'id_column': 'Z', 'text_column': 'AA'})
    assert load_xlsx(path, profile)[0].text == 'Last column text'
    sheet['AA3'] = '=1+1'
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    with pytest.raises(InputProfileError) as error:
        load_xlsx(path, profile)
    assert error.value.code == 'XLSX_TEXT_FORMULA'
    assert error.value.details == {'sheet': 'First', 'row': 3}
    assert path.read_bytes() == before


def test_position_mcp_snapshot_replays_without_original_and_preserves_profile(tmp_path):
    from reqmap.agent_input import import_agent_input
    from reqmap.agent_store import SessionStore
    from reqmap.agent_types import SessionSeed, SessionSettings
    profile = parse_input_profile(positional())
    config = agent_module().load_agent_config(config_file(tmp_path))
    config = replace(config, input_profile=profile)
    path = book(tmp_path)
    snapshot = import_agent_input({'kind': 'xlsx', 'path': str(path)}, config)
    assert snapshot.input_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    settings = SessionSettings(config.analysis_profile, 8, profile, '0' * 64, None, '2.0', '2.0')
    seed = SessionSeed(snapshot, settings, (), None, 'unknown')
    store = SessionStore(config.session_root)
    reply = store.create('position-start', {'source': 'synthetic'}, lambda: seed)
    assert reply.ok, reply.error
    path.unlink()
    restored = SessionStore(config.session_root).read(reply.data['session_id']).seed
    assert restored.input_snapshot == snapshot
    assert restored.settings.input_profile == profile


def test_duplicate_continuation_does_not_inherit_subject_or_parent_context(tmp_path):
    from reqmap.binding_source import bind_source
    from reqmap.decomposition import prepare_decomposition
    from reqmap.grouping import build_groups
    from tests.test_input_xlsx import make_workbook, profile
    path = make_workbook(tmp_path, headers=['ID', 'Text'], rows=[
        ['7', 'Nova должна создавать виртуальную машину.'],
        ['7', 'за одну миллисекунду'],
    ])
    grouped = build_groups(load_xlsx(path, profile('ID', 'Text')))
    assert grouped[0].group_ids == grouped[1].group_ids
    assert grouped[1].parent_id is None
    context = prepare_decomposition(grouped[1], None)
    assert context['parent_text'] is None
    binding = bind_source(grouped[1])
    assert binding.unresolved_fragments


def test_position_import_without_dimensions_preserves_sparse_rows(tmp_path):
    wb = Workbook(write_only=True)
    sheet = wb.create_sheet('Data')
    sheet.append([None, '7', 'First', 'P1'])
    sheet.append([None, '7'])
    sheet.append([None, '7', 'Last'])
    path = tmp_path / 'streaming.xlsx'
    wb.save(path)
    wb.close()
    profile = parse_input_profile({**positional(), 'include_sheets': ['Data'],
                                   'data_start_row': 1, 'priority_column': 'D'})
    before = path.read_bytes()
    rows = load_xlsx(path, profile)
    assert [(r.coordinate.row, r.source_id, r.text) for r in rows] == [(1, '7', 'First'), (3, '7', 'Last')]
    assert rows[0].source_fields == (SourceField('D', 'P1'),)
    assert rows[1].source_fields == (SourceField('D', ''),)
    assert path.read_bytes() == before


@pytest.mark.parametrize('empty', [False, True])
def test_position_import_without_dimensions_rejects_missing_columns(tmp_path, empty):
    wb = Workbook(write_only=True)
    sheet = wb.create_sheet('Data')
    if not empty:
        sheet.append([None, '7'])
    path = tmp_path / 'missing.xlsx'
    wb.save(path)
    wb.close()
    profile = parse_input_profile({**positional(), 'include_sheets': ['Data'], 'data_start_row': 1})
    with pytest.raises(InputProfileError) as error:
        load_xlsx(path, profile)
    assert error.value.code == 'XLSX_PROFILE_COLUMN_MISSING'
