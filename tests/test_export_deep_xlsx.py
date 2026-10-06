"""Deterministic, normalized XLSX export for canonical deep results."""

from dataclasses import replace
import hashlib
from pathlib import Path

from openpyxl import load_workbook
import pytest

from reqmap.export_deep_xlsx import SHEET_HEADERS, verify_deep_xlsx, write_deep_xlsx
from tests.test_export_deep_json import deep_run


EXPECTED_HEADERS = {
    "Требования": (
        "Requirement ID", "Source ID", "Ordinal", "Файл", "Лист", "Строка",
        "Текст требования", "Parent ID", "Group IDs", "Source fields",
        "Source hints", "Состояние: код", "Поддержка: код", "Atom IDs",
        "Responsibility IDs", "Procedure graph IDs", "Диагностика", "Source binding",
    ),
    "Атомарные утверждения": (
        "Atom ID", "Requirement ID", "Ordinal", "Формулировка атома",
        "Исходная цитата", "Обязательный", "Состояние: код",
        "Поддержка: код", "Responsibility IDs", "Подтверждённые аспекты",
        "Неподтверждённые аспекты", "Диагностика",
        "Obligation ID", "Source spans", "Source SHA-256", "Binding decision",
    ),
    "Ответственность": (
        "Record ID", "Requirement ID", "Atom ID", "Контур", "Component ref",
        "Executor ref", "Target contour", "Target ref", "Action ref", "Effect ref",
        "Lifecycle phase", "Version scope", "Evidence IDs", "Поддержка: код",
        "Related record IDs", "Procedure step IDs", "Диагностика",
        "Obligation ID", "Predicate IDs (atom)", "Uncovered spans (atom)",
    ),
    "Процедуры": (
        "Graph ID", "Requirement ID", "Template ID", "Graph diagnostics", "Step ID",
        "Phase", "Contour", "Executor ref", "Target ref", "Action ref",
        "Preconditions", "Success criteria", "Evidence IDs", "Depends on",
        "Rollback step ID",
    ),
    "Доказательства": (
        "Evidence ID", "Claim", "Claim kind", "Polarity", "Strength", "Source ID",
        "Locator", "Version constraint", "Applicable contours",
        "Supports entity refs", "Local excerpt", "Review state",
    ),
    "Диагностика": (
        "Порядок", "Scope", "Entity ID", "Requirement ID", "Текст требования",
        "Atom ID", "Исходная цитата", "Diagnostic",
    ),
    "Запуск": ("Параметр", "Значение"),
}


def _rows(path: Path, sheet: str) -> list[tuple[object, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return list(workbook[sheet].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()


def test_deep_workbook_has_exact_order_headers_and_canonical_rows(
    tmp_path: Path,
) -> None:
    run = deep_run()
    path = tmp_path / "result.xlsx"

    digest = write_deep_xlsx(run, path)

    assert digest == hashlib.sha256(path.read_bytes()).hexdigest()
    assert SHEET_HEADERS == EXPECTED_HEADERS
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == list(EXPECTED_HEADERS)
        for name, headers in EXPECTED_HEADERS.items():
            assert next(workbook[name].iter_rows(values_only=True)) == headers
            assert workbook[name].sheet_state == "visible"
        assert workbook["Требования"].max_row - 1 == len(run.requirements)
        assert workbook["Атомарные утверждения"].max_row - 1 == 1
        assert workbook["Ответственность"].max_row - 1 == 1
        assert workbook["Процедуры"].max_row - 1 == 1
        assert workbook["Доказательства"].max_row - 1 == 1
    finally:
        workbook.close()

    responsibility = _rows(path, "Ответственность")[0]
    assert responsibility[0:11] == (
        "REQ-0001-A001-R001", "REQ-0001", "REQ-0001-A001",
        "openstack_runtime", "nova", "ACTOR-NOVA-API", "openstack_runtime",
        "TARGET-NOVA-SERVER", "ACTION-NOVA-CREATE",
        "EFFECT-NOVA-SERVER-ACTIVE", "runtime",
    )
    assert responsibility[11] == (
        '{"host_profile":"rocky_linux_9","kolla_ansible_release":"2025.1",'
        '"source_release":"2025.1","target_release":"2025.1",'
        '"version_constraint":"2025.1"}'
    )
    assert responsibility[12] == '["EV-NOVA-CREATE"]'
    procedure = _rows(path, "Процедуры")[0]
    assert procedure[0:10] == (
        "REQ-0001-P001", "REQ-0001", "PROC-NOVA", "[]",
        "REQ-0001-P001-S001", "runtime", "openstack_runtime",
        "ACTOR-NOVA-API", "TARGET-NOVA-SERVER", "ACTION-NOVA-CREATE",
    )
    run_rows = dict(_rows(path, "Запуск"))
    assert tuple(run_rows) == (
        "run_id", "schema_version", "run_status", "reqmap_version",
        "analysis_profile", "model", "seed", "top_k", "input_sha256",
        "snapshot_id", "key_id", "manifest_sha256", "signer_identity",
        "prompt_versions", "release_profile", "retry_counts", "groups", "counts",
    )
    assert run_rows["counts"] == (
        '{"atoms":1,"diagnostics":0,"evidence":1,"procedure_steps":1,'
        '"requirements":1,"responsibilities":1}'
    )
    verify_deep_xlsx(path, run)


@pytest.mark.parametrize("literal", ("=1+1", "+SUM(A1:A2)", "-1+1", "@SUM(A1:A2)"))
def test_deep_xlsx_preserves_formula_like_source_values_as_literals(
    tmp_path: Path, literal: str
) -> None:
    run = deep_run()
    result = run.requirements[0]
    atom_result = result.atom_results[0]
    changed_requirement = replace(result.requirement, source_id=literal, text=literal)
    changed_atom = replace(atom_result.atom, text=literal, source_quote=literal)
    changed_result = replace(
        result,
        requirement=changed_requirement,
        atom_results=(replace(atom_result, atom=changed_atom),),
    )
    changed = replace(run, requirements=(changed_result,))
    path = tmp_path / "literal.xlsx"

    write_deep_xlsx(changed, path)

    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        assert workbook["Требования"].cell(2, 2).value == literal
        assert workbook["Требования"].cell(2, 2).data_type == "s"
        assert workbook["Атомарные утверждения"].cell(2, 4).value == literal
        assert all(
            cell.data_type != "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def test_deep_xlsx_is_byte_deterministic(tmp_path: Path) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"

    first_hash = write_deep_xlsx(deep_run(), first)
    second_hash = write_deep_xlsx(deep_run(), second)

    assert first_hash == second_hash
    assert first.read_bytes() == second.read_bytes()


def test_deep_xlsx_rejects_symlink_without_overwriting_target(tmp_path: Path) -> None:
    target = tmp_path / "target.xlsx"
    target.write_bytes(b"keep")
    link = tmp_path / "result.xlsx"
    link.symlink_to(target)

    with pytest.raises(ValueError, match="symlink"):
        write_deep_xlsx(deep_run(), link)

    assert target.read_bytes() == b"keep"
