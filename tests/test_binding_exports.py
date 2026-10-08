"""Every subject artifact carries the same immutable binding and decision."""
from dataclasses import replace
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from reqmap.agent_knowledge import VerifiedKnowledge
from reqmap.agent_service import AgentService
from reqmap.binding_catalog import knowledge_digest
from reqmap.config import AnalysisProfile
from reqmap.knowledge import KnowledgeBase
from reqmap.models import to_dict
from tests.agent_support import make_agent_config
from tests.test_agent_session import start
from tests.test_binding_sessions import select_source
from tests.test_binding_mapping import setup_binding, kb


@pytest.mark.parametrize("restart", [False, True])
@pytest.mark.parametrize("suffix,status", [("", "supported"), (" за одну миллисекунду", "insufficient_evidence")])
def test_all_exports_preserve_exact_gap_and_predicate_refs(tmp_path, monkeypatch, setup_binding, restart, suffix, status):
    knowledge, catalog, raw = setup_binding
    deep = type(knowledge) is not KnowledgeBase
    profile = AnalysisProfile.DEEP if deep else AnalysisProfile.LEGACY
    verified = VerifiedKnowledge(knowledge, knowledge_digest(knowledge), knowledge.snapshot_id if deep else None, catalog)
    monkeypatch.setattr("reqmap.agent_service.load_agent_knowledge", lambda config: verified)
    text = "Nova должна создавать ВМ через API" + suffix
    from tests.source_context_support import with_reviewed_texts
    config = with_reviewed_texts(replace(make_agent_config(tmp_path/"agent"), analysis_profile=profile), (text,))
    svc = AgentService(config)
    sid = start(svc, (text,))
    context = select_source(svc, sid, text)
    assert context.ok, context
    proposal = {**raw, "proposal_schema_version":2, "obligation_id":"REQ-0001-O001", "predicate_ids":["P-NOVA-CREATE"]}
    submitted = svc.call("reqmap_submit_mapping", dict(session_id=sid, atom_id="REQ-0001-A001",
        context_id=context.data["context_id"], proposal=proposal, request_id="map", expected_revision=1))
    assert submitted.ok, submitted
    result = submitted.data["result"]["atom_result"] if deep else submitted.data["result"]
    decision = result["binding_decision"]
    assert result["support_status"] == status
    if restart:
        svc = AgentService(svc.config)
    from reqmap.agent_finalize import build_agent_run
    from reqmap.publication import publish_artifacts
    view, verified = svc.verified_view(svc.store.read(sid))
    run = build_agent_run(view, verified)
    if restart:
        reply = svc.call("reqmap_finalize", dict(session_id=sid, request_id="finish", expected_revision=2, allow_partial=True))
        assert reply.ok, reply
        output = Path(reply.data["artifacts"]["result.json"]).parent
        assert svc.call("reqmap_get_result", dict(session_id=sid)).ok
    else:
        output = tmp_path/"after-submit"
        output.mkdir()
        assert not publish_artifacts(run, output)
    payload = json.loads((output/"result.json").read_text())
    assert payload["schema_version"] == ("2.2" if deep else "1.2")
    row = payload["requirements"][0]
    assert row["source_binding"] == context.data["payload"]["source_binding"]
    assert row["atom_results"][0]["binding_decision"] == decision
    assert row["atom_results"][0]["atom"]["source_spans"] == result["atom"]["source_spans"]
    book = load_workbook(output/"result.xlsx")
    assert len(book.sheetnames) == (7 if deep else 5)
    atoms = book["Атомарные утверждения"]
    headers = [c.value for c in atoms[1]]
    cell = atoms.cell(2, headers.index("Binding decision")+1)
    assert json.loads(cell.value) == decision
    report = (output/"report.md").read_text()
    from reqmap.export_json import canonical_json_bytes
    assert canonical_json_bytes(decision).decode().strip() in report
    manifest = json.loads((output/"manifest.json").read_text())
    assert manifest["binding_contract"] == payload["metadata"]["binding_contract"]
    journal = [json.loads(line) for line in (output/"run.jsonl").read_text().splitlines()]
    assert any(event.get("binding_contract") == manifest["binding_contract"] for event in journal)
    for name, digest in manifest["artifact_hashes"].items():
        assert hashlib.sha256((output/name).read_bytes()).hexdigest() == digest
    from reqmap.crosscheck import crosscheck
    from reqmap.crosscheck_deep import crosscheck_deep
    check = crosscheck_deep if deep else crosscheck
    assert not check(run, output/"result.json", output/"result.xlsx", output/"report.md")
    # A changed gap is detected even when all status/count columns are unchanged.
    forged = dict(decision)
    forged["uncovered"] = [] if decision["uncovered"] else [{"start":0,"end":1,"quote":"N"}]
    cell.value = json.dumps(forged, ensure_ascii=False)
    book.save(output/"result.xlsx")
    book.close()
    assert check(run, output/"result.json", output/"result.xlsx", output/"report.md")


@pytest.mark.parametrize('profile', list(AnalysisProfile))
@pytest.mark.parametrize('kind', ['binding', 'all_proof_fields'])
def test_long_unknown_source_finalizes_and_every_xlsx_fragment_is_verified(tmp_path, profile, kind):
    from tests.test_binding_sessions import insufficient_proposal
    from reqmap.agent_finalize import build_agent_run
    from reqmap.export_json import canonical_json_bytes
    from reqmap.crosscheck import crosscheck
    from reqmap.crosscheck_deep import crosscheck_deep
    # The second variant also overflows span/decision/aspect JSON and tests astral Unicode.
    text = 'Требование '+('дополнительное условие '*350 if kind=='binding' else '😀"'*16000)
    config=make_agent_config(tmp_path/'agent',profile)
    if kind=='all_proof_fields':
        config=replace(config,limits=replace(config.limits,max_response_bytes=4*1024*1024))
    svc=AgentService(config)
    sid=start(svc,(text,))
    context=select_source(svc,sid,text)
    assert context.ok,context
    reply=svc.call('reqmap_submit_mapping',dict(session_id=sid,atom_id='REQ-0001-A001',context_id=context.data['context_id'],proposal=insufficient_proposal(context.data['payload'],profile is AnalysisProfile.DEEP),request_id='map',expected_revision=1))
    assert reply.ok,reply
    view,knowledge=svc.verified_view(svc.store.read(sid))
    run=build_agent_run(view,knowledge)
    final=svc.call('reqmap_finalize',dict(session_id=sid,request_id='finish',expected_revision=2,allow_partial=True))
    assert final.ok,final
    output=Path(final.data['artifacts']['result.json']).parent
    assert AgentService(svc.config).call('reqmap_get_result',dict(session_id=sid)).ok
    payload=json.loads((output/'result.json').read_bytes())
    book=load_workbook(output/'result.xlsx')
    runsheet=book['Запуск'];parts={row[0].value:row[1].value for row in list(runsheet.rows)[1:]}
    def restored(sheet,header):
        headers=[c.value for c in book[sheet][1]]
        value=book[sheet].cell(2,headers.index(header)+1).value
        ref=json.loads(value)
        if type(ref) is dict and set(ref)=={'xlsx_text_key','parts','sha256'}:
            value=''.join(parts[f"{ref['xlsx_text_key']}:{i}"] for i in range(1,ref['parts']+1))
            assert hashlib.sha256(value.encode()).hexdigest()==ref['sha256']
        return json.loads(value)
    assert restored('Требования','Source binding')==payload['requirements'][0]['source_binding']
    assert restored('Атомарные утверждения','Binding decision')==payload['requirements'][0]['atom_results'][0]['binding_decision']
    assert restored('Атомарные утверждения','Source spans')==payload['requirements'][0]['atom_results'][0]['atom']['source_spans']
    assert len(book.sheetnames)==(7 if profile is AnalysisProfile.DEEP else 5)
    assert all(len(c.value.encode('utf-16-le'))//2<=32767 for sh in book for row in sh for c in row if isinstance(c.value,str))
    check=crosscheck_deep if profile is AnalysisProfile.DEEP else crosscheck
    assert not check(run,output/'result.json',output/'result.xlsx',output/'report.md')
    fragment=next(row[1] for row in list(runsheet.rows)[1:] if str(row[0].value).startswith('xlsx-text:'))
    fragment.value+='tampered'
    book.save(output/'result.xlsx');book.close()
    assert check(run,output/'result.json',output/'result.xlsx',output/'report.md')
