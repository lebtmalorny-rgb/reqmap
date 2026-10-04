"""Публичный CLI reqmap: источники, артефакты и коды завершения."""

from __future__ import annotations

from dataclasses import replace
import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
import pytest
from reqmap import publication

import reqmap.cli as cli
from reqmap.aggregation import aggregate_groups
from reqmap.crosscheck import CrosscheckIssue
from reqmap.models import AnalysisState, RequirementResult
from tests.test_json_manifest import completed_run
from tests.test_pipeline import config_for


def run_with_status(status: str, *, preflight_ok: bool = True):
    run = completed_run()
    if status == "SUCCESS":
        return replace(
            run,
            metadata={**run.metadata, "preflight_ok": preflight_ok},
        )
    source_result = run.requirements[0]
    failed = RequirementResult(
        requirement=replace(
            source_result.requirement,
            requirement_id=("REQ-0002" if status == "PARTIAL" else "REQ-0001"),
            ordinal=(2 if status == "PARTIAL" else 1),
            group_ids=(),
            text="Требование с ошибкой обработки",
        ),
        analysis_state=AnalysisState.MODEL_FAILED,
        support_status=None,
        atom_results=(),
        mappings=(),
        diagnostics=("Модель недоступна",),
    )
    requirements = (
        (source_result, failed) if status == "PARTIAL" else (failed,)
    )
    return replace(
        run,
        run_status=status,
        requirements=requirements,
        groups=aggregate_groups(requirements),
        evidence=(run.evidence if status == "PARTIAL" else ()),
        metadata={**run.metadata, "preflight_ok": preflight_ok},
        diagnostics=("Модель недоступна",),
    )


def install_fake_pipeline(
    monkeypatch: pytest.MonkeyPatch,
    *,
    status: str = "SUCCESS",
    preflight_ok: bool = True,
) -> SimpleNamespace:
    captured = SimpleNamespace(request=None, config=None, model=None)
    config = config_for()
    monkeypatch.setattr(cli, "load_config", lambda path, environ: config, raising=False)
    monkeypatch.setattr(
        cli,
        "OpenAICompatibleClient",
        lambda model_config: SimpleNamespace(model_config=model_config),
        raising=False,
    )

    def fake_analyze(request, actual_config, model):
        captured.request = request
        captured.config = actual_config
        captured.model = model
        run = run_with_status(status, preflight_ok=preflight_ok)
        if not preflight_ok:
            request.output_dir.mkdir(mode=0o700, parents=True, exist_ok=True)
            (request.output_dir / "run.jsonl").write_text(
                '{"event":"preflight_failed"}\n',
                encoding="utf-8",
            )
            (request.output_dir / "manifest.json").write_text(
                '{"run_status":"FAILED"}\n',
                encoding="utf-8",
            )
            run = replace(
                run,
                metadata={
                    **run.metadata,
                    "preflight_artifacts_written": True,
                },
            )
        return run

    monkeypatch.setattr(cli, "analyze", fake_analyze, raising=False)
    return captured


def analyze_args(output: Path) -> list[str]:
    return [
        "analyze",
        "--requirement",
        "Создать VM",
        "--config",
        "config.example.yaml",
        "--output",
        str(output),
    ]


def test_cli_multiple_requirements_invokes_one_pipeline_and_writes_artifacts(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    captured = install_fake_pipeline(monkeypatch)
    output = tmp_path / "result"
    values = ("Создать VM", "Создать сеть")

    code = cli.main(
        [
            "analyze",
            "--requirement",
            values[0],
            "--requirement",
            values[1],
            "--config",
            "config.example.yaml",
            "--output",
            str(output),
        ]
    )

    assert code == 0
    assert [item.text for item in captured.request.requirements] == list(values)
    expected_input = (
        json.dumps(values, ensure_ascii=False, separators=(",", ":")) + "\n"
    ).encode("utf-8")
    assert captured.request.input_sha256 == hashlib.sha256(expected_input).hexdigest()
    names = {
        "result.json",
        "result.xlsx",
        "report.md",
        "run.jsonl",
        "manifest.json",
    }
    assert {item.name for item in output.iterdir() if item.is_file()} == names
    terminal = capsys.readouterr()
    assert "Профиль анализа: legacy" in terminal.out
    for name in names:
        assert str((output / name).resolve()) in terminal.out
    assert terminal.out.count("SHA-256") == 5
    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    for name, digest in manifest["artifact_hashes"].items():
        assert digest == hashlib.sha256((output / name).read_bytes()).hexdigest()
    assert "top-secret" not in "".join(
        path.read_text(encoding="utf-8", errors="ignore")
        for path in output.iterdir()
        if path.suffix != ".xlsx"
    )


def test_cli_stdin_lines_preserves_rows_and_hashes_raw_utf8(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    captured = install_fake_pipeline(monkeypatch)
    source = "Первое\n\nВторое\n"
    monkeypatch.setattr(cli.sys, "stdin", io.StringIO(source))

    code = cli.main(
        [
            "analyze",
            "--stdin",
            "--text-mode",
            "lines",
            "--config",
            "config.example.yaml",
            "--output",
            str(tmp_path / "result"),
        ]
    )

    assert code == 0
    assert [item.text for item in captured.request.requirements] == [
        "Первое",
        "Второе",
    ]
    assert [item.coordinate.row for item in captured.request.requirements] == [1, 3]
    assert captured.request.input_sha256 == hashlib.sha256(
        source.encode("utf-8")
    ).hexdigest()


def test_cli_xlsx_input_hashes_original_bytes_without_modifying_source(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    captured = install_fake_pipeline(monkeypatch)
    source_path = tmp_path / "requirements.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("ID", "Требование"))
    sheet.append(("source-1", "Создать VM"))
    workbook.save(source_path)
    workbook.close()
    before = source_path.read_bytes()

    code = cli.main(
        [
            "analyze",
            str(source_path),
            "--config",
            "config.example.yaml",
            "--output",
            str(tmp_path / "result"),
        ]
    )

    assert code == 0
    assert captured.request.source_path == source_path
    assert captured.request.input_kind == "xlsx"
    assert captured.request.requirements[0].text == "Создать VM"
    assert captured.request.input_sha256 == hashlib.sha256(before).hexdigest()
    assert source_path.read_bytes() == before


def test_cli_xlsx_hash_and_parser_use_the_same_byte_snapshot(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    captured = install_fake_pipeline(monkeypatch)
    source_path = tmp_path / "requirements.xlsx"
    replacement_path = tmp_path / "replacement.xlsx"
    for path, text in (
        (source_path, "Исходный снимок"),
        (replacement_path, "Подменённый файл"),
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(("ID", "Требование"))
        sheet.append(("source-1", text))
        workbook.save(path)
        workbook.close()
    original_bytes = source_path.read_bytes()
    replacement_bytes = replacement_path.read_bytes()
    original_read_bytes = Path.read_bytes
    first_source_read = True

    def replace_after_first_read(path: Path) -> bytes:
        nonlocal first_source_read
        data = original_read_bytes(path)
        if path == source_path and first_source_read:
            first_source_read = False
            path.write_bytes(replacement_bytes)
        return data

    monkeypatch.setattr(Path, "read_bytes", replace_after_first_read)

    code = cli.main(
        [
            "analyze",
            str(source_path),
            "--config",
            "config.example.yaml",
            "--output",
            str(tmp_path / "result"),
        ]
    )

    assert code == 0
    assert captured.request.requirements[0].text == "Исходный снимок"
    assert captured.request.input_sha256 == hashlib.sha256(original_bytes).hexdigest()


@pytest.mark.parametrize(
    "artifact_name",
    [
        "result.json",
        "result.xlsx",
        "report.md",
        "run.jsonl",
        "manifest.json",
    ],
)
def test_cli_rejects_source_collision_with_every_published_artifact(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    artifact_name: str,
) -> None:
    captured = install_fake_pipeline(monkeypatch)
    output = tmp_path / "result"
    output.mkdir()
    source_path = output / artifact_name
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("ID", "Требование"))
    sheet.append(("source-1", "Исходный текст"))
    workbook.save(source_path)
    workbook.close()
    before = source_path.read_bytes()

    code = cli.main(
        [
            "analyze",
            str(source_path),
            "--config",
            "config.example.yaml",
            "--output",
            str(output),
        ]
    )

    assert code == 2
    assert captured.request is None
    assert source_path.read_bytes() == before


@pytest.mark.parametrize(
    "arguments",
    [
        ["analyze", "--config", "c", "--output", "o"],
        [
            "analyze",
            "input.xlsx",
            "--requirement",
            "Создать VM",
            "--config",
            "c",
            "--output",
            "o",
        ],
        [
            "analyze",
            "--requirement",
            "",
            "--config",
            "c",
            "--output",
            "o",
        ],
    ],
)
def test_cli_rejects_missing_conflicting_or_empty_input_in_russian(
    monkeypatch: pytest.MonkeyPatch,
    arguments: list[str],
    capsys: pytest.CaptureFixture[str],
) -> None:
    captured = install_fake_pipeline(monkeypatch)

    assert cli.main(arguments) == 2

    assert captured.request is None
    terminal = capsys.readouterr()
    assert "ошиб" in terminal.err.lower()
    assert "Traceback" not in terminal.err


@pytest.mark.parametrize(
    ("status", "preflight_ok", "expected_code"),
    [
        ("PARTIAL", True, 4),
        ("FAILED", True, 6),
        ("FAILED", False, 3),
    ],
)
def test_cli_returns_documented_status_codes(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    status: str,
    preflight_ok: bool,
    expected_code: int,
) -> None:
    install_fake_pipeline(
        monkeypatch,
        status=status,
        preflight_ok=preflight_ok,
    )
    output = tmp_path / f"result-{expected_code}"

    code = cli.main(analyze_args(output))

    assert code == expected_code
    names = {item.name for item in output.iterdir() if item.is_file()}
    if preflight_ok:
        assert names == {
            "result.json",
            "result.xlsx",
            "report.md",
            "run.jsonl",
            "manifest.json",
        }
    else:
        assert names == {"run.jsonl", "manifest.json"}


def test_cli_crosscheck_failure_returns_five_with_five_artifacts(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    install_fake_pipeline(monkeypatch)
    monkeypatch.setattr(
        publication,
        "crosscheck",
        lambda *args: (
            CrosscheckIssue(
                "CROSSCHECK_TEST",
                "Синтетическое расхождение артефактов.",
            ),
        ),
        raising=False,
    )
    output = tmp_path / "result"

    assert cli.main(analyze_args(output)) == 5

    assert {item.name for item in output.iterdir() if item.is_file()} == {
        "result.json",
        "result.xlsx",
        "report.md",
        "run.jsonl",
        "manifest.json",
    }


def test_cli_preflight_failure_does_not_report_stale_files_through_output_symlink(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    captured = install_fake_pipeline(monkeypatch)
    target = tmp_path / "old-output"
    target.mkdir()
    for name in ("run.jsonl", "manifest.json"):
        (target / name).write_text("stale\n", encoding="utf-8")
    output = tmp_path / "linked-output"
    output.symlink_to(target, target_is_directory=True)
    monkeypatch.setattr(
        cli,
        "analyze",
        lambda request, config, model: run_with_status(
            "FAILED",
            preflight_ok=False,
        ),
    )

    code = cli.main(analyze_args(output))

    assert code == 3
    assert captured.request is None
    assert "SHA-256" not in capsys.readouterr().out


def test_cli_does_not_report_stale_files_when_preflight_published_nothing(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    install_fake_pipeline(monkeypatch)
    output = tmp_path / "old-output"
    output.mkdir()
    for name in ("run.jsonl", "manifest.json"):
        (output / name).write_text("stale\n", encoding="utf-8")
    failed = run_with_status("FAILED", preflight_ok=False)
    failed = replace(
        failed,
        metadata={
            **failed.metadata,
            "preflight_artifacts_written": False,
        },
    )
    monkeypatch.setattr(cli, "analyze", lambda *args: failed)

    assert cli.main(analyze_args(output)) == 3

    assert "SHA-256" not in capsys.readouterr().out


def test_cli_hides_traceback_unless_debug_is_enabled(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    install_fake_pipeline(monkeypatch)

    def broken_analyze(*args):
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr(cli, "analyze", broken_analyze, raising=False)
    arguments = analyze_args(tmp_path / "normal")
    assert cli.main(arguments) == 6
    normal = capsys.readouterr()
    assert "Traceback" not in normal.err
    assert "ошиб" in normal.err.lower()

    debug_arguments = [*analyze_args(tmp_path / "debug"), "--debug"]
    assert cli.main(debug_arguments) == 6
    debug = capsys.readouterr()
    assert "Traceback" in debug.err


def test_cli_knowledge_validate_uses_local_snapshot(
    capsys: pytest.CaptureFixture[str],
) -> None:
    code = cli.main(
        [
            "knowledge",
            "validate",
            "--path",
            "knowledge/epoxy-2025.1",
        ]
    )

    assert code == 0
    assert "2025.1" in capsys.readouterr().out


def test_cli_help_is_fully_russian(capsys: pytest.CaptureFixture[str]) -> None:
    assert cli.main(["analyze", "--help"]) == 0

    help_text = capsys.readouterr().out.lower()
    assert "использование:" in help_text
    assert "позиционные аргументы" in help_text
    assert "параметры" in help_text
    assert "show this help" not in help_text


@pytest.mark.parametrize(
    "arguments",
    [
        ["analyze", "--requirement", "Создать VM"],
        [
            "analyze",
            "--requirement",
            "Создать VM",
            "--config",
            "c",
            "--output",
            "o",
            "--unknown",
        ],
        [
            "analyze",
            "--requirement",
            "Создать VM",
            "--config",
            "c",
            "--output",
            "o",
            "--text-mode",
            "many",
        ],
        [
            "analyze",
            "--requirement",
            "Создать VM",
            "--stdin",
            "--config",
            "c",
            "--output",
            "o",
        ],
    ],
)
def test_argparse_user_errors_do_not_leak_english_boilerplate(
    arguments: list[str],
    capsys: pytest.CaptureFixture[str],
) -> None:
    assert cli.main(arguments) == 2

    terminal = capsys.readouterr()
    message = terminal.err.lower()
    assert "ошибка" in message
    assert not any(
        phrase in message
        for phrase in (
            "required",
            "unrecognized",
            "invalid choice",
            "not allowed",
            "expected one argument",
        )
    )
