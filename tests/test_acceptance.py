"""Переносимый acceptance-контракт полного публичного pipeline reqmap."""

from __future__ import annotations

from collections.abc import Iterator, Mapping
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import re
from threading import Thread
from typing import cast

from openpyxl import load_workbook
import pytest

from reqmap.cli import main
from reqmap.knowledge import load_knowledge


ROOT = Path(__file__).resolve().parents[1]
INPUT_FIXTURE = ROOT / "tests" / "fixtures" / "requirements_synthetic.json"
ARTIFACT_NAMES = (
    "result.json",
    "result.xlsx",
    "report.md",
    "run.jsonl",
    "manifest.json",
)


class _FakeOpenAIServer(ThreadingHTTPServer):
    model_name = "acceptance-local-model"

    def __init__(self, response_factory=None) -> None:
        self.response_factory = response_factory or _model_response
        super().__init__(("127.0.0.1", 0), _FakeOpenAIHandler)
        self.requests: list[tuple[str, str, object | None]] = []
        self.errors: list[str] = []

    @property
    def base_url(self) -> str:
        host, port = self.server_address
        return f"http://{host}:{port}/v1"


class _FakeOpenAIHandler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    @property
    def fake_server(self) -> _FakeOpenAIServer:
        return cast(_FakeOpenAIServer, self.server)

    def do_GET(self) -> None:  # noqa: N802 - имя задаётся stdlib protocol
        self.fake_server.requests.append(("GET", self.path, None))
        if self.path != "/v1/models":
            self._send_json(404, {"error": "unknown endpoint"})
            return
        self._send_json(
            200,
            {"data": [{"id": self.fake_server.model_name}]},
        )

    def do_POST(self) -> None:  # noqa: N802 - имя задаётся stdlib protocol
        try:
            length = int(self.headers.get("Content-Length", "0"))
            body = json.loads(self.rfile.read(length).decode("utf-8"))
            self.fake_server.requests.append(("POST", self.path, body))
            if self.path != "/v1/chat/completions":
                self._send_json(404, {"error": "unknown endpoint"})
                return
            prompt = _prompt_payload(body)
            response = self.fake_server.response_factory(prompt)
            self._send_json(
                200,
                {
                    "choices": [
                        {
                            "message": {
                                "content": json.dumps(
                                    response,
                                    ensure_ascii=False,
                                    separators=(",", ":"),
                                )
                            }
                        }
                    ]
                },
            )
        except Exception as exc:  # pragma: no cover - отображается через server.errors
            self.fake_server.errors.append(f"{type(exc).__name__}: {exc}")
            self._send_json(500, {"error": "fake model failure"})

    def _send_json(self, status: int, payload: object) -> None:
        body = json.dumps(
            payload,
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        del format, args


@pytest.fixture
def fake_openai_server() -> Iterator[_FakeOpenAIServer]:
    server = _FakeOpenAIServer()
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield server
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


@pytest.fixture
def epoxy_kb() -> Path:
    path = ROOT / "knowledge" / "epoxy-2025.1"
    assert load_knowledge(path).release == "2025.1"
    return path


def _prompt_payload(body: object) -> dict[str, object]:
    assert isinstance(body, dict)
    assert body["model"] == _FakeOpenAIServer.model_name
    assert body["temperature"] == 0
    assert body["response_format"] == {"type": "json_object"}
    messages = body["messages"]
    assert isinstance(messages, list) and len(messages) == 2
    user_message = messages[1]
    assert isinstance(user_message, dict)
    assert user_message["role"] == "user"
    payload = json.loads(cast(str, user_message["content"]))
    assert isinstance(payload, dict)
    return payload


def _model_response(payload: Mapping[str, object]) -> dict[str, object]:
    if "requirement_text" in payload:
        return _decomposition_response(cast(str, payload["requirement_text"]))
    atom = payload.get("atom")
    if isinstance(atom, dict):
        return _mapping_response(cast(str, atom["source_quote"]))
    raise AssertionError("Неизвестный prompt fake LLM")


def _decomposition_response(requirement_text: str) -> dict[str, object]:
    if "создание виртуальной машины" in requirement_text:
        quotes = (
            "создание виртуальной машины через API",
            "подключение сети",
        )
    elif "sysctl" in requirement_text:
        quotes = (requirement_text,)
    elif "Квантовый режим X" in requirement_text:
        quotes = (requirement_text,)
    else:
        raise AssertionError(f"Неизвестное acceptance-требование: {requirement_text}")
    return {
        "atoms": [
            {
                "text": quote,
                "source_quote": quote,
                "mandatory": True,
            }
            for quote in quotes
        ]
    }


def _mapping_response(source_quote: str) -> dict[str, object]:
    if "создание виртуальной машины" in source_quote:
        return _supported_nova_mapping()
    if "sysctl" in source_quote:
        return _supported_host_sysctl_bundle()
    return {
        "support_status": "insufficient_evidence",
        "supported_aspects": [],
        "unconfirmed_aspects": [source_quote],
        "mappings": [],
    }


def _supported_nova_mapping() -> dict[str, object]:
    # Intentional proposed overclaim: the shipped scope evidence must downgrade it.
    role = "Compute instances через REST API"
    return {
        "support_status": "supported",
        "supported_aspects": [role],
        "unconfirmed_aspects": [],
        "mappings": [
            {
                "component_id": "nova",
                "role_ru": role,
                "relation": "implements",
                "phase": "runtime",
                "implementation_source": "upstream",
                "mechanism": "openstack_api",
                "steps": [
                    {
                        "action_ru": role,
                        "mechanism": "openstack_api",
                        "command": None,
                        "api_operation": role,
                    }
                ],
                "evidence_ids": ["E-NOVA-SCOPE-001"],
                "support_status": "supported",
                "reason_ru": "Nova документирует compute instances через REST API.",
            }
        ],
    }


def _supported_host_sysctl_bundle() -> dict[str, object]:
    host_role = "Автоматизация kernel sysctl хоста"
    kolla_role = "Design-time reconfigure конфигурации OpenStack"
    return {
        "support_status": "supported",
        "supported_aspects": [host_role, kolla_role],
        "unconfirmed_aspects": [],
        "mappings": [
            {
                "component_id": "host_os_kernel_sysctl",
                "role_ru": host_role,
                "relation": "host_os_change",
                "phase": "designtime",
                "implementation_source": "kolla_ansible",
                "mechanism": "kolla_ansible",
                "steps": [
                    {
                        "action_ru": host_role,
                        "mechanism": "ansible role",
                        "command": None,
                        "api_operation": None,
                    },
                    {
                        "action_ru": kolla_role,
                        "mechanism": "kolla_ansible",
                        "command": "kolla-ansible reconfigure",
                        "api_operation": None,
                    },
                ],
                "evidence_ids": ["E-HOST-SYSCTL-OFFICIAL-001"],
                "support_status": "supported",
                "reason_ru": "Официальный snapshot содержит role sysctl.",
            },
            {
                "component_id": "kolla_ansible",
                "role_ru": kolla_role,
                "relation": "configures",
                "phase": "designtime",
                "implementation_source": "kolla_ansible",
                "mechanism": "kolla_ansible",
                "steps": [
                    {
                        "action_ru": kolla_role,
                        "mechanism": "kolla_ansible",
                        "command": "kolla-ansible reconfigure",
                        "api_operation": None,
                    }
                ],
                "evidence_ids": ["E-KOLLA-RECONFIGURE-001"],
                "support_status": "supported",
                "reason_ru": "Kolla-Ansible документирует reconfigure.",
            },
        ],
    }


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _result_counts(result: Mapping[str, object]) -> dict[str, int]:
    requirements = cast(list[dict[str, object]], result["requirements"])
    return {
        "requirements": len(requirements),
        "atoms": sum(
            len(cast(list[object], item["atom_results"]))
            for item in requirements
        ),
        "mappings": sum(
            len(cast(list[object], item["mappings"]))
            for item in requirements
        ),
        "evidence": len(cast(list[object], result["evidence"])),
    }


def _assert_cross_artifact_contract(
    output: Path,
    result: Mapping[str, object],
) -> None:
    counts = _result_counts(result)
    report = (output / "report.md").read_text(encoding="utf-8")
    marker = re.search(r"^<!-- reqmap-counts:(\{[^\r\n]+\}) -->$", report, re.MULTILINE)
    assert marker is not None
    assert json.loads(marker.group(1)) == counts

    workbook = load_workbook(output / "result.xlsx", read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Требования",
            "Атомарные утверждения",
            "Сопоставления",
            "Доказательства",
            "Запуск",
        ]
        assert workbook["Требования"].max_row - 1 == counts["requirements"]
        assert workbook["Атомарные утверждения"].max_row - 1 == counts["atoms"]
        assert workbook["Сопоставления"].max_row - 1 == counts["mappings"]
        assert workbook["Доказательства"].max_row - 1 == counts["evidence"]
    finally:
        workbook.close()

    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["run_id"] == result["run_id"]
    assert manifest["run_status"] == result["run_status"]
    assert "manifest.json" not in manifest["artifact_hashes"]
    assert manifest["artifact_hashes"] == {
        name: _sha256(output / name)
        for name in ARTIFACT_NAMES
        if name != "manifest.json"
    }


def test_acceptance_contract(
    tmp_path: Path,
    fake_openai_server: _FakeOpenAIServer,
    epoxy_kb: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    fixture = json.loads(INPUT_FIXTURE.read_text(encoding="utf-8"))
    assert set(fixture) == {"requirements"}
    requirements = cast(list[str], fixture["requirements"])
    config_path = tmp_path / "acceptance-config.json"
    config_path.write_text(
        json.dumps(
            {
                "model": {
                    "base_url": fake_openai_server.base_url,
                    "model": fake_openai_server.model_name,
                    "timeout_seconds": 5,
                    "retries": 0,
                    "supports_response_format": True,
                    "seed": 19,
                },
                "knowledge_path": str(epoxy_kb),
                "top_k": 8,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    output = tmp_path / "acceptance-result"
    arguments = ["analyze"]
    for requirement in requirements:
        arguments.extend(("--requirement", requirement))
    arguments.extend(("--config", str(config_path), "--output", str(output)))

    exit_code = main(arguments)
    captured = capsys.readouterr()

    assert exit_code == 0, (captured.out, captured.err, fake_openai_server.errors)
    assert fake_openai_server.errors == []
    assert "Статус запуска: SUCCESS" in captured.out
    for name in ARTIFACT_NAMES:
        path = output / name
        assert path.is_file() and not path.is_symlink()
        assert f"{path.resolve()} SHA-256 {_sha256(path)}" in captured.out
    assert captured.out.count("SHA-256") == len(ARTIFACT_NAMES)

    result = json.loads((output / "result.json").read_text(encoding="utf-8"))
    result_requirements = result["requirements"]
    assert [item["requirement"]["text"] for item in result_requirements] == requirements
    assert all(item["analysis_state"] == "completed" for item in result_requirements)
    assert len(result_requirements[0]["atom_results"]) == 2
    assert any(
        len(atom_result["mappings"]) >= 2
        for item in result_requirements
        for atom_result in item["atom_results"]
    )

    host_mappings = {
        item["component_id"]: item
        for item in result_requirements[1]["mappings"]
    }
    assert {"host_os_kernel_sysctl", "kolla_ansible"}.issubset(host_mappings)
    host = host_mappings["host_os_kernel_sysctl"]
    kolla = host_mappings["kolla_ansible"]
    assert host["relation"] == "host_os_change"
    assert host["phase"] == kolla["phase"] == "designtime"
    assert host["implementation_source"] == kolla["implementation_source"] == (
        "kolla_ansible"
    )
    assert any(
        step["command"] == "kolla-ansible reconfigure"
        for step in host["steps"]
    )

    unknown = result_requirements[2]
    assert unknown["support_status"] == "insufficient_evidence"
    assert unknown["mappings"] == []
    assert result["run_status"] == "SUCCESS"
    _assert_cross_artifact_contract(output, result)

    requests = fake_openai_server.requests
    assert [(method, path) for method, path, _body in requests].count(
        ("GET", "/v1/models")
    ) == 1
    assert all(
        path in {"/v1/models", "/v1/chat/completions"}
        for _method, path, _body in requests
    )
    assert sum(method == "POST" for method, _path, _body in requests) == 7
