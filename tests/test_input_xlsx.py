import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook

from reqmap.config import InputProfile
from reqmap.input_xlsx import InputProfileError, detect_profile, load_xlsx
from reqmap.models import SourceField, SourceHint


def make_workbook(
    tmp_path: Path,
    *,
    headers: list[str],
    rows: list[list[object]],
    sheet_name: str = "Требования",
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    path = tmp_path / "requirements.xlsx"
    workbook.save(path)
    return path


def profile(
    id_column: str,
    text_column: str,
    *,
    include_sheets: tuple[str, ...] = ("Требования",),
    priority_column: str | None = None,
    expected_result_column: str | None = None,
    parent_column: str | None = None,
    hint_columns: tuple[str, ...] = (),
) -> InputProfile:
    return InputProfile(
        include_sheets=include_sheets,
        exclude_sheets=(),
        header_row=1,
        id_column=id_column,
        text_column=text_column,
        priority_column=priority_column,
        expected_result_column=expected_result_column,
        parent_column=parent_column,
        hint_columns=hint_columns,
    )


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_xlsx_import_does_not_modify_source(tmp_path: Path) -> None:
    """Открытие в режиме записи изменило бы контрольную сумму пользовательского файла."""
    path = make_workbook(
        tmp_path, headers=["ID", "Требование"], rows=[["1", "Исходный текст"]]
    )
    before = sha256_file(path)

    result = load_xlsx(path, profile("ID", "Требование"))

    assert result[0].text == "Исходный текст"
    assert sha256_file(path) == before


def test_repeated_source_ids_get_distinct_technical_ids(tmp_path: Path) -> None:
    """Дедупликация source_id потеряла бы самостоятельную исходную строку."""
    path = make_workbook(
        tmp_path,
        headers=["ID", "Требование"],
        rows=[["1.1", "Первое"], ["1.1", "Второе"]],
    )

    result = load_xlsx(path, profile("ID", "Требование"))

    assert [item.source_id for item in result] == ["1.1", "1.1"]
    assert [item.requirement_id for item in result] == ["REQ-0001", "REQ-0002"]


def test_xlsx_keeps_coordinate_and_configured_source_context(tmp_path: Path) -> None:
    """Сведение настроенных столбцов к тексту лишило бы результат трассируемости."""
    path = make_workbook(
        tmp_path,
        headers=["ID", "Требование", "Приоритет", "Ожидаемый", "Родитель", "Подсказка"],
        rows=[["7", "  Текст  ", "P1", "Результат", "5", "Не доверять"]],
    )

    result = load_xlsx(
        path,
        profile(
            "ID",
            "Требование",
            priority_column="Приоритет",
            expected_result_column="Ожидаемый",
            parent_column="Родитель",
            hint_columns=("Подсказка",),
        ),
    )

    assert result[0].text == "  Текст  "
    assert result[0].coordinate.source_name == "requirements.xlsx"
    assert result[0].coordinate.sheet == "Требования"
    assert result[0].coordinate.row == 2
    assert result[0].parent_id == "5"
    assert result[0].source_fields == (
        SourceField(column="Приоритет", value="P1"),
        SourceField(column="Ожидаемый", value="Результат"),
        SourceField(column="Родитель", value="5"),
    )
    assert result[0].source_hints == (SourceHint(value="Не доверять", column="Подсказка"),)


def test_detect_profile_accepts_exactly_one_visible_sheet_match(tmp_path: Path) -> None:
    """Автодетект должен выбрать только однозначную структуру известного формата."""
    path = make_workbook(tmp_path, headers=["Код", "Описание требования"], rows=[["1", "Текст"]])
    workbook = Workbook()
    workbook.close()
    from openpyxl import load_workbook

    loaded = load_workbook(path, read_only=True, data_only=False)
    try:
        result = detect_profile(loaded)
    finally:
        loaded.close()

    assert result == profile("Код", "Описание требования")


def test_detect_profile_rejects_ambiguous_text_columns(tmp_path: Path) -> None:
    """Выбор одной из двух текстовых колонок молча исказил бы импорт."""
    path = make_workbook(
        tmp_path, headers=["ID", "Требование", "Описание требования"], rows=[["1", "A", "B"]]
    )
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        with pytest.raises(InputProfileError) as error:
            detect_profile(workbook)
    finally:
        workbook.close()

    assert error.value.code == "XLSX_PROFILE_AMBIGUOUS"
    assert "неоднозначна" in error.value.message_ru


def test_xlsx_rejects_formula_in_requirement_text_cell(tmp_path: Path) -> None:
    """Импорт вычисленного значения вместо формулы сделал бы исходный текст неподтверждаемым."""
    path = make_workbook(tmp_path, headers=["ID", "Требование"], rows=[["1", "=CONCAT(\"A\", \"B\")"]])

    with pytest.raises(InputProfileError) as error:
        load_xlsx(path, profile("ID", "Требование"))

    assert error.value.code == "XLSX_TEXT_FORMULA"
    assert "формул" in error.value.message_ru


def test_explicit_profile_ignores_duplicate_unconfigured_headers(tmp_path: Path) -> None:
    """Отказ из-за несвязанного столбца сделал бы явный профиль чрезмерно строгим."""
    path = make_workbook(
        tmp_path,
        headers=["ID", "Требование", "Комментарий", "Комментарий"],
        rows=[["1", "Текст", "Первый", "Второй"]],
    )

    result = load_xlsx(path, profile("ID", "Требование"))

    assert [item.text for item in result] == ["Текст"]


def test_hidden_sheet_is_imported_only_when_explicitly_included(tmp_path: Path) -> None:
    """Обход скрытого листа без явного профиля раскрыл бы исключенные данные."""
    path = make_workbook(tmp_path, headers=["ID", "Требование"], rows=[["1", "Открытый"]])
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    hidden = workbook.create_sheet("Скрытый")
    hidden.append(["ID", "Требование"])
    hidden.append(["2", "Скрытый"])
    hidden.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()

    auto_detected = load_xlsx(path, None)
    explicit = load_xlsx(path, profile("ID", "Требование", include_sheets=("Скрытый",)))

    assert [item.text for item in auto_detected] == ["Открытый"]
    assert [item.text for item in explicit] == ["Скрытый"]
