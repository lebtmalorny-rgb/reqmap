"""Independent deep JSON/XLSX/Markdown cross-artifact verification."""

from pathlib import Path

from openpyxl import load_workbook

from reqmap.crosscheck_deep import crosscheck_deep
from reqmap.export_deep_json import write_deep_canonical_json
from reqmap.export_deep_markdown import write_deep_markdown
from reqmap.export_deep_xlsx import write_deep_xlsx
from tests.test_export_deep_json import deep_run


def _artifacts(tmp_path: Path):
    run = deep_run()
    json_path = tmp_path / "result.json"
    xlsx_path = tmp_path / "result.xlsx"
    markdown_path = tmp_path / "report.md"
    write_deep_canonical_json(run, json_path)
    write_deep_xlsx(run, xlsx_path)
    write_deep_markdown(run, markdown_path)
    return run, json_path, xlsx_path, markdown_path


def _codes(issues) -> set[str]:
    return {issue.code for issue in issues}


def test_deep_crosscheck_accepts_consistent_artifacts(tmp_path: Path) -> None:
    run, json_path, xlsx_path, markdown_path = _artifacts(tmp_path)

    assert crosscheck_deep(run, json_path, xlsx_path, markdown_path) == ()


def test_deep_crosscheck_detects_duplicate_json_keys_and_schema_deviation(
    tmp_path: Path,
) -> None:
    run, json_path, xlsx_path, markdown_path = _artifacts(tmp_path)
    payload = json_path.read_text(encoding="utf-8")
    json_path.write_text('{"run_id":"forged",' + payload[1:], encoding="utf-8")

    codes = _codes(crosscheck_deep(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_DEEP_JSON_DUPLICATE_KEY" in codes
    assert "CROSSCHECK_DEEP_JSON_BYTES" in codes


def test_deep_crosscheck_collects_independent_xlsx_relationship_mismatches(
    tmp_path: Path,
) -> None:
    run, json_path, xlsx_path, markdown_path = _artifacts(tmp_path)
    workbook = load_workbook(xlsx_path)
    workbook["Ответственность"].cell(2, 6).value = "forged-executor"
    workbook["Процедуры"].cell(2, 14).value = '["forged-step"]'
    workbook["Доказательства"].cell(2, 10).value = '["forged-entity"]'
    workbook.save(xlsx_path)
    workbook.close()

    codes = _codes(crosscheck_deep(run, json_path, xlsx_path, markdown_path))

    assert {
        "CROSSCHECK_DEEP_XLSX_RESPONSIBILITY",
        "CROSSCHECK_DEEP_XLSX_PROCEDURE",
        "CROSSCHECK_DEEP_XLSX_EVIDENCE",
    }.issubset(codes)


def test_deep_crosscheck_detects_formula_hidden_sheet_and_reordered_schema(
    tmp_path: Path,
) -> None:
    run, json_path, xlsx_path, markdown_path = _artifacts(tmp_path)
    workbook = load_workbook(xlsx_path)
    workbook["Требования"].cell(2, 7).value = "=1+1"
    workbook["Диагностика"].sheet_state = "hidden"
    workbook._sheets[0], workbook._sheets[1] = workbook._sheets[1], workbook._sheets[0]
    workbook.save(xlsx_path)
    workbook.close()

    codes = _codes(crosscheck_deep(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_DEEP_XLSX_FORMULA" in codes
    assert "CROSSCHECK_DEEP_XLSX_HIDDEN" in codes
    assert "CROSSCHECK_DEEP_XLSX_SHEETS" in codes


def test_deep_crosscheck_detects_duplicate_ids_and_deleted_rows(tmp_path: Path) -> None:
    run, json_path, xlsx_path, markdown_path = _artifacts(tmp_path)
    workbook = load_workbook(xlsx_path)
    sheet = workbook["Ответственность"]
    sheet.append([cell.value for cell in sheet[2]])
    workbook["Процедуры"].delete_rows(2)
    workbook.save(xlsx_path)
    workbook.close()

    codes = _codes(crosscheck_deep(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_DEEP_XLSX_DUPLICATE_ID" in codes
    assert "CROSSCHECK_DEEP_XLSX_RESPONSIBILITY" in codes
    assert "CROSSCHECK_DEEP_XLSX_PROCEDURE" in codes


def test_deep_crosscheck_detects_markdown_marker_duplication_and_mutation(
    tmp_path: Path,
) -> None:
    run, json_path, xlsx_path, markdown_path = _artifacts(tmp_path)
    text = markdown_path.read_text(encoding="utf-8")
    marker = next(line for line in text.splitlines() if line.startswith("<!-- reqmap-counts:"))
    markdown_path.write_text(
        text.replace(marker, marker.replace('"atoms":1', '"atoms":99') + "\n" + marker),
        encoding="utf-8",
    )

    codes = _codes(crosscheck_deep(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_DEEP_MARKDOWN_MARKER" in codes
    assert "CROSSCHECK_DEEP_MARKDOWN_BYTES" in codes


def test_deep_crosscheck_returns_independent_read_issues_for_bad_artifacts(
    tmp_path: Path,
) -> None:
    run, json_path, xlsx_path, markdown_path = _artifacts(tmp_path)
    json_path.write_bytes(b"\xff")
    xlsx_path.write_bytes(b"not an xlsx")
    markdown_path.write_bytes(b"\xff")

    codes = _codes(crosscheck_deep(run, json_path, xlsx_path, markdown_path))

    assert "CROSSCHECK_DEEP_JSON_READ" in codes
    assert "CROSSCHECK_DEEP_XLSX_READ" in codes
    assert "CROSSCHECK_DEEP_MARKDOWN_READ" in codes
