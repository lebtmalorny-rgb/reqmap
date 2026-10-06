"""Scope descriptions must never prove a concrete implementation claim."""
from copy import deepcopy
from dataclasses import replace
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from reqmap.knowledge import KnowledgeError, load_knowledge, validate_knowledge
from reqmap.mapping import accept_mapping
from reqmap.models import EvidenceClaimScope, EvidencePolarity, SupportStatus
from reqmap.retrieval import retrieve
from tests.test_acceptance import _supported_nova_mapping
from tests.test_mapping import atom, kb, raw_mapping, response, step
from tests.test_agent_session import service, start, atoms
from tests.test_knowledge import kb_path


def _rewrite_scope(path, value, *, remove=False):
    record = json.loads((path / "evidence.jsonl").read_text().splitlines()[0])
    if remove:
        record.pop("claim_scope", None)
    else:
        record["claim_scope"] = value
    (path / "evidence.jsonl").write_text(json.dumps(record, ensure_ascii=False) + "\n")


@pytest.mark.parametrize("scope", ["context", "specific"])
def test_loader_keeps_explicit_claim_scope_in_evidence_context(kb_path, scope):
    _rewrite_scope(kb_path, scope)
    kb = load_knowledge(kb_path, verify_snapshot_hash=False)
    from reqmap.mapping import prepare_mapping
    text = "Nova сервер"
    context = prepare_mapping(atom(text), retrieve(kb, text, (), 8), kb)
    assert context["evidence"][0]["claim_scope"] == scope


@pytest.mark.parametrize("scope", [None, True, 1, [], {}, "", "direct", "Specific"])
def test_loader_rejects_invalid_claim_scope(kb_path, scope):
    _rewrite_scope(kb_path, scope)
    with pytest.raises(KnowledgeError) as error:
        load_knowledge(kb_path, verify_snapshot_hash=False)
    assert error.value.code == "EVIDENCE_CLAIM_SCOPE"


def test_unclassified_old_evidence_defaults_to_context_without_losing_candidates(kb_path):
    _rewrite_scope(kb_path, None, remove=True)
    kb = load_knowledge(kb_path, verify_snapshot_hash=False)
    text = "Nova должна обеспечивать квантовую телепортацию."
    candidates = retrieve(kb, text, (), 8)
    proposal = _supported_nova_mapping()
    role = kb.capabilities["CAP-NOVA-SERVER-API"].name_ru
    proposal["supported_aspects"] = [role]
    mapping = proposal["mappings"][0]
    mapping.update(role_ru=role, evidence_ids=["E-NOVA-API-001"])
    mapping["steps"][0].update(action_ru=role, api_operation=role)

    result = accept_mapping(atom(text), candidates, kb, proposal)

    assert any("E-NOVA-API-001" in c.evidence_ids for c in candidates)
    assert result.support_status is SupportStatus.INSUFFICIENT_EVIDENCE


SERVICES = (
    "aodh", "barbican", "ceilometer", "cinder", "cloudkitty", "designate",
    "glance", "gnocchi", "heat", "horizon", "ironic", "keystone", "magnum",
    "manila", "masakari", "neutron", "nova", "octavia", "placement", "swift",
    "trove", "watcher",
)


@pytest.mark.parametrize("component_id", SERVICES)
def test_scope_label_cannot_prove_unrelated_obligation(component_id):
    kb = load_knowledge(Path("knowledge/epoxy-2025.1"))
    capability = next(c for c in kb.capabilities.values() if c.component_id == component_id)
    text = f"{component_id} должен обеспечивать квантовую телепортацию."
    proposal = deepcopy(_supported_nova_mapping())
    proposal["supported_aspects"] = [capability.name_ru]
    mapping = proposal["mappings"][0]
    mapping.update(component_id=component_id, role_ru=capability.name_ru,
                   evidence_ids=[e.evidence_id for e in kb.evidence.values()
                                 if e.capability_id == capability.capability_id])
    mapping["steps"][0].update(action_ru=capability.name_ru, api_operation=capability.name_ru)

    result = accept_mapping(atom(text), retrieve(kb, text, (), 8), kb, proposal)

    assert result.support_status is SupportStatus.INSUFFICIENT_EVIDENCE
    assert result.supported_aspects == ()
    assert result.atom.source_quote == text
    assert any("EVIDENCE_CONTEXT_ONLY" in d for d in result.diagnostics)


def test_scope_label_cannot_prove_latency_guarantee():
    kb = load_knowledge(Path("knowledge/epoxy-2025.1"))
    text = "Nova должна гарантировать создание виртуальной машины за одну миллисекунду."

    result = accept_mapping(atom(text), retrieve(kb, text, (), 8), kb, _supported_nova_mapping())

    assert result.support_status is SupportStatus.INSUFFICIENT_EVIDENCE
    assert result.supported_aspects == ()


def test_specific_operation_still_supports_its_mapping(kb, tmp_path):
    from tests.binding_factories import binding_case, loaded_catalog
    atom, context, proposal = binding_case(kb, loaded_catalog(tmp_path, kb), response(raw_mapping()))
    result = accept_mapping(atom, retrieve(kb, atom.text, (), 8), kb, proposal, binding_context=context)
    assert result.support_status is SupportStatus.SUPPORTED


def test_direct_knowledge_validation_rejects_untyped_claim_scope(kb_path):
    kb = load_knowledge(kb_path)
    invalid = replace(kb.evidence["E-NOVA-API-001"], claim_scope="specific")
    kb = replace(kb, evidence={**kb.evidence, invalid.evidence_id: invalid})
    assert any(i.code == "EVIDENCE_CLAIM_SCOPE" and i.object_id == "E-NOVA-API-001"
               for i in validate_knowledge(kb))


def test_context_text_cannot_launder_operation_beside_specific_evidence(kb):
    contextual = replace(kb.evidence["E-NOVA"], evidence_id="E-CONTEXT",
                         claim_scope=EvidenceClaimScope.CONTEXT,
                         claim_ru="Nova API описывает POST /quantum-teleportation.")
    kb = replace(kb, evidence={**kb.evidence, contextual.evidence_id: contextual})
    text = "Nova должна обеспечивать квантовую телепортацию."
    proposal = response(raw_mapping(evidence_ids=("E-NOVA", "E-CONTEXT"),
                                   steps=[step("runtime", api_operation="POST /quantum-teleportation")]))
    result = accept_mapping(atom(text), retrieve(kb, text, (), 8), kb, proposal)
    assert result.support_status is SupportStatus.INSUFFICIENT_EVIDENCE


@pytest.mark.parametrize("version", ["2025.1", "2024.2"])
def test_negative_context_cannot_prove_unsupported_or_version_conflict(kb, version):
    contextual = replace(kb.evidence["E-NOVA"], claim_scope=EvidenceClaimScope.CONTEXT,
                         polarity=EvidencePolarity.NEGATIVE, version_constraint=version)
    kb = replace(kb, evidence={**kb.evidence, contextual.evidence_id: contextual})
    text = "Nova API"
    proposal = response(raw_mapping(support_status="not_supported"), status="not_supported")
    result = accept_mapping(atom(text), retrieve(kb, text, (), 8), kb, proposal)
    assert result.support_status is SupportStatus.INSUFFICIENT_EVIDENCE


def test_scope_only_proposal_remains_insufficient_after_restart_and_publication(tmp_path):
    svc = service(tmp_path)
    text = "Nova должна поддерживать квантовую телепортацию виртуальной машины."
    sid = start(svc, (text,))
    assert atoms(svc, sid, text).ok
    context = svc.call("reqmap_get_atom_context", {"session_id": sid, "atom_id": "REQ-0001-A001"})
    result = svc.call("reqmap_submit_mapping", dict(
        session_id=sid, atom_id="REQ-0001-A001", context_id=context.data["context_id"],
        proposal=_supported_nova_mapping(), request_id="map", expected_revision=1,
    ))
    assert result.ok, result
    assert result.data["result"]["support_status"] == "insufficient_evidence"
    svc = type(svc)(svc.config)
    final = svc.call("reqmap_finalize", dict(
        session_id=sid, request_id="final", expected_revision=2, allow_partial=False,
    ))
    assert final.ok, final
    fetched = svc.call("reqmap_get_result", {"session_id": sid})
    assert fetched.ok, fetched
    paths = final.data["artifacts"]
    assert len(paths) == 5
    result = json.loads(Path(paths["result.json"]).read_text())
    assert result["requirements"][0]["support_status"] == "insufficient_evidence"
    assert result["requirements"][0]["requirement"]["text"] == text
    assert result["evidence"][0]["claim_scope"] == "context"
    workbook = load_workbook(paths["result.xlsx"], read_only=True)
    try:
        rows = list(workbook["Требования"].iter_rows(values_only=True))
        assert rows[1][rows[0].index("Поддержка: код")] == "insufficient_evidence"
    finally:
        workbook.close()
    assert "EVIDENCE_CONTEXT_ONLY" in Path(paths["report.md"]).read_text()
