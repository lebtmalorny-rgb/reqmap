"""Нормализованный русский XLSX и независимая OOXML-проверка."""

from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
import hashlib
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
import openpyxl.writer.excel as openpyxl_excel_writer
import pytest

from reqmap.export_xlsx import (
    ExportError,
    SHEET_HEADERS,
    verify_xlsx,
    write_xlsx,
)
from reqmap.models import (
    AnalysisState,
    EvidencePolarity,
    EvidenceStrength,
    ImplementationSource,
    SupportStatus,
)
from tests.test_json_manifest import completed_run


def mixed_run():
    run = completed_run()
    source_result = run.requirements[0]
    source_atom = source_result.atom_results[0]
    first_mapping = source_atom.mappings[0]
    second_mapping = replace(
        first_mapping,
        mapping_id="REQ-0001-A001-M002",
        component_id="neutron",
        evidence_ids=("evidence-2",),
        role_ru="Neutron network API",
    )
    atom_result = replace(
        source_atom,
        mappings=(first_mapping, second_mapping),
    )
    requirement_result = replace(
        source_result,
        atom_results=(atom_result,),
        mappings=(first_mapping, second_mapping),
    )
    group = replace(
        run.groups[0],
        component_ids=("neutron", "nova"),
        mapping_ids=(first_mapping.mapping_id, second_mapping.mapping_id),
    )
    evidence = replace(
        run.evidence[0],
        evidence_id="evidence-2",
        component_id="neutron",
        capability_id="CAP-NEUTRON",
        source_id="SRC-NEUTRON",
        local_path="sources/neutron.md",
        claim_ru="Neutron подтверждает синтетическую возможность.",
    )
    return replace(
        run,
        requirements=(requirement_result,),
        groups=(group,),
        evidence=(*run.evidence, evidence),
    )


def test_xlsx_contains_five_normalized_sheets_and_many_to_many_rows(
    tmp_path: Path,
) -> None:
    run = mixed_run()
    path = tmp_path / "result.xlsx"

    digest = write_xlsx(run, path)

    assert digest == hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=False)
    assert workbook.sheetnames == [
        "Требования",
        "Атомарные утверждения",
        "Сопоставления",
        "Доказательства",
        "Запуск",
    ]
    for name, expected_headers in SHEET_HEADERS.items():
        values = next(workbook[name].iter_rows(values_only=True))
        assert values == expected_headers
    assert SHEET_HEADERS["Сопоставления"] == (
        "Mapping ID",
        "Atom ID",
        "Component ID",
        "Роль",
        "Тип связи: код",
        "Тип связи: русский",
        "Фаза: код",
        "Фаза: русский",
        "Источник реализации: код",
        "Источник реализации: русский",
        "Механизм",
        "Шаги",
        "Evidence IDs",
        "Поддержка: код",
        "Поддержка: русский",
        "Обоснование",
        "Obligation ID", "Predicate IDs (atom)", "Uncovered spans (atom)",
    )
    assert SHEET_HEADERS["Доказательства"] == (
        "Evidence ID",
        "Component ID",
        "Capability ID",
        "Полярность: код",
        "Полярность: русский",
        "Сила: код",
        "Сила: русский",
        "Утверждение",
        "URL",
        "Локальный путь",
        "Локатор",
        "Версия",
        "SHA-256",
        "Происхождение",
    )
    assert workbook["Требования"].max_row == (
        1 + len(run.requirements) + len(run.groups)
    )
    assert workbook["Атомарные утверждения"].max_row == 2
    assert workbook["Сопоставления"].max_row == 3
    assert workbook["Доказательства"].max_row == 3
    mapping_rows = list(
        workbook["Сопоставления"].iter_rows(min_row=2, values_only=True)
    )
    assert [row[0] for row in mapping_rows] == [
        "REQ-0001-A001-M001",
        "REQ-0001-A001-M002",
    ]
    assert mapping_rows[0][8:10] == ("upstream", "базовый OpenStack")
    evidence_rows = list(
        workbook["Доказательства"].iter_rows(min_row=2, values_only=True)
    )
    assert evidence_rows[0][3:7] == (
        "positive",
        "положительное",
        "direct",
        "прямое",
    )
    requirement_rows = list(
        workbook["Требования"].iter_rows(min_row=2, values_only=True)
    )
    assert [row[0] for row in requirement_rows] == ["requirement", "group"]
    assert requirement_rows[0][12:16] == (
        "completed",
        "обработка завершена",
        "supported",
        "поддерживается",
    )
    workbook.close()
    verify_xlsx(path, run)


@pytest.mark.parametrize(
    "literal",
    ("=1+1", "+SUM(A1:A2)", "-1+1", "@SUM(A1:A2)"),
)
def test_xlsx_preserves_formula_like_user_text_as_literal_string(
    tmp_path: Path,
    literal: str,
) -> None:
    run = completed_run()
    source_result = run.requirements[0]
    source_atom = source_result.atom_results[0]
    source = replace(
        source_result.requirement,
        source_id=literal,
        text=literal,
    )
    claim = replace(source_atom.atom, text=literal, source_quote=literal)
    atom_result = replace(source_atom, atom=claim)
    requirement_result = replace(
        source_result,
        requirement=source,
        atom_results=(atom_result,),
    )
    run = replace(run, requirements=(requirement_result,))
    path = tmp_path / "literal.xlsx"

    write_xlsx(run, path)

    workbook = load_workbook(path, read_only=False, data_only=False)
    source_id = workbook["Требования"].cell(row=2, column=3)
    source_text = workbook["Требования"].cell(row=2, column=7)
    assert source_id.value == literal
    assert source_text.value == literal
    assert source_id.data_type == "s"
    assert source_text.data_type == "s"
    assert all(
        cell.data_type != "f"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    workbook.close()
    verify_xlsx(path, run)


def test_xlsx_output_is_deterministic(tmp_path: Path) -> None:
    run = mixed_run()
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"

    first_hash = write_xlsx(run, first)
    second_hash = write_xlsx(run, second)

    assert first_hash == second_hash
    assert first.read_bytes() == second.read_bytes()


def test_xlsx_output_is_independent_from_openpyxl_save_clock(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    run = mixed_run()
    first = tmp_path / "first-clock.xlsx"
    second = tmp_path / "second-clock.xlsx"

    class FirstClock(datetime):
        @classmethod
        def now(cls, tz=None):
            value = cls(2026, 1, 1, 0, 0, 0)
            return value.replace(tzinfo=tz) if tz is not None else value

    class SecondClock(datetime):
        @classmethod
        def now(cls, tz=None):
            value = cls(2027, 2, 2, 3, 4, 5)
            return value.replace(tzinfo=tz) if tz is not None else value

    monkeypatch.setattr(
        openpyxl_excel_writer,
        "datetime",
        SimpleNamespace(datetime=FirstClock, timezone=timezone),
    )
    first_hash = write_xlsx(run, first)
    monkeypatch.setattr(
        openpyxl_excel_writer,
        "datetime",
        SimpleNamespace(datetime=SecondClock, timezone=timezone),
    )
    second_hash = write_xlsx(run, second)

    assert first_hash == second_hash
    assert first.read_bytes() == second.read_bytes()


def test_verify_xlsx_rejects_formula_and_corrupt_zip(tmp_path: Path) -> None:
    run = completed_run()
    formula_path = tmp_path / "formula.xlsx"
    write_xlsx(run, formula_path)
    workbook = load_workbook(formula_path)
    workbook["Требования"].cell(row=2, column=7).value = "=1+1"
    workbook.save(formula_path)
    workbook.close()

    with pytest.raises(ExportError, match="формул"):
        verify_xlsx(formula_path, run)

    corrupt_path = tmp_path / "corrupt.xlsx"
    corrupt_path.write_bytes(b"not-an-xlsx")
    with pytest.raises(ExportError, match="поврежд"):
        verify_xlsx(corrupt_path, run)


@pytest.mark.parametrize(
    ("sheet_name", "row", "column", "replacement"),
    [
        ("Требования", 3, 17, '["forged-component"]'),
        ("Атомарные утверждения", 2, 3, "Подменённый атом"),
        ("Сопоставления", 2, 11, "forged mechanism"),
        ("Доказательства", 2, 8, "Подменённое доказательство"),
        ("Запуск", 10, 2, 999),
    ],
)
def test_verify_xlsx_rejects_any_tampered_exported_value(
    tmp_path: Path,
    sheet_name: str,
    row: int,
    column: int,
    replacement: object,
) -> None:
    run = completed_run()
    path = tmp_path / f"tampered-{sheet_name}.xlsx"
    write_xlsx(run, path)
    workbook = load_workbook(path)
    workbook[sheet_name].cell(row=row, column=column).value = replacement
    workbook.save(path)
    workbook.close()

    with pytest.raises(ExportError, match="не совпада.*с RunResult"):
        verify_xlsx(path, run)


def test_xlsx_russian_enum_labels_cover_failure_and_insufficient(
    tmp_path: Path,
) -> None:
    run = completed_run()
    source_result = run.requirements[0]
    failed = replace(
        source_result,
        analysis_state=AnalysisState.MODEL_FAILED,
        support_status=None,
        atom_results=(),
        mappings=(),
        diagnostics=("model_failed",),
    )
    run = replace(
        run,
        run_status="FAILED",
        requirements=(failed,),
        groups=(
            replace(
                run.groups[0],
                support_status=None,
                component_ids=(),
                mapping_ids=(),
                analysis_states=(AnalysisState.MODEL_FAILED,),
            ),
        ),
        evidence=(),
    )
    path = tmp_path / "failed.xlsx"

    write_xlsx(run, path)

    workbook = load_workbook(path, read_only=True, data_only=False)
    row = next(
        workbook["Требования"].iter_rows(min_row=2, max_row=2, values_only=True)
    )
    assert row[12:16] == (
        "model_failed",
        "ошибка модели",
        None,
        None,
    )
    workbook.close()


@pytest.mark.parametrize(
    ("status", "expected_ru"),
    [
        (SupportStatus.PARTIAL, "поддерживается частично"),
        (SupportStatus.NOT_SUPPORTED, "не поддерживается"),
        (SupportStatus.INSUFFICIENT_EVIDENCE, "недостаточно доказательств"),
        (SupportStatus.NOT_APPLICABLE, "неприменимо"),
    ],
)
def test_support_status_russian_labels_are_stable(
    status: SupportStatus,
    expected_ru: str,
) -> None:
    from reqmap.export_xlsx import support_status_ru

    assert support_status_ru(status) == expected_ru


@pytest.mark.parametrize(
    ("value", "expected_ru"),
    [
        (ImplementationSource.UPSTREAM, "базовый OpenStack"),
        (ImplementationSource.KOLLA_ANSIBLE, "Kolla-Ansible"),
        (ImplementationSource.PRODUCT_EXTENSION, "расширение продукта"),
        (ImplementationSource.EXTERNAL_COMPONENT, "внешний компонент"),
    ],
)
def test_implementation_source_russian_labels_are_stable(
    value: ImplementationSource,
    expected_ru: str,
) -> None:
    from reqmap.export_xlsx import implementation_source_ru

    assert implementation_source_ru(value) == expected_ru


@pytest.mark.parametrize(
    ("value", "expected_ru"),
    [
        (EvidencePolarity.POSITIVE, "положительное"),
        (EvidencePolarity.NEGATIVE, "отрицательное"),
    ],
)
def test_evidence_polarity_russian_labels_are_stable(
    value: EvidencePolarity,
    expected_ru: str,
) -> None:
    from reqmap.export_xlsx import evidence_polarity_ru

    assert evidence_polarity_ru(value) == expected_ru


@pytest.mark.parametrize(
    ("value", "expected_ru"),
    [
        (EvidenceStrength.DIRECT, "прямое"),
        (EvidenceStrength.INDIRECT, "косвенное"),
        (EvidenceStrength.NONE, "отсутствует"),
    ],
)
def test_evidence_strength_russian_labels_are_stable(
    value: EvidenceStrength,
    expected_ru: str,
) -> None:
    from reqmap.export_xlsx import evidence_strength_ru

    assert evidence_strength_ru(value) == expected_ru
