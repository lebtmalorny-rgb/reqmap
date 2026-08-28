"""Independent cross-artifact verification for canonical deep results."""

from __future__ import annotations

import json
from pathlib import Path
import re
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from reqmap.crosscheck import CrosscheckIssue
from reqmap.deep_models import DeepRunResult, ResponsibilityContour
from reqmap.export_deep_json import _deep_run_payload, validate_deep_run_result
from reqmap.export_deep_xlsx import SHEET_HEADERS
from reqmap.export_json import canonical_json_bytes, symlink_component
from reqmap.models import AnalysisState, SupportStatus


_MARKER_PATTERN = re.compile(r"^<!-- reqmap-counts:(\{[^\r\n]*\}) -->$", re.MULTILINE)
_MARKER_PREFIX = "<!-- reqmap-counts:"
_MARKDOWN_HEADINGS = (
    "Сводка",
    "Требования и группы",
    "Атомарные утверждения",
    "Контуры ответственности",
    "Executor → target",
    "Версии и область применимости",
    "Процедуры",
    "Каталог evidence",
    "Нормализованная диагностика",
    "Конфликты evidence",
    "Предупреждения procedures и rollback",
    "Проблемные требования и атомы",
    "Ошибки обработки",
)
_PROBLEM_STATUSES = frozenset(
    {SupportStatus.PARTIAL, SupportStatus.NOT_SUPPORTED, SupportStatus.INSUFFICIENT_EVIDENCE}
)
_SHEET_CODES = {
    "Требования": "CROSSCHECK_DEEP_XLSX_REQUIREMENT",
    "Атомарные утверждения": "CROSSCHECK_DEEP_XLSX_ATOM",
    "Ответственность": "CROSSCHECK_DEEP_XLSX_RESPONSIBILITY",
    "Процедуры": "CROSSCHECK_DEEP_XLSX_PROCEDURE",
    "Доказательства": "CROSSCHECK_DEEP_XLSX_EVIDENCE",
    "Диагностика": "CROSSCHECK_DEEP_XLSX_DIAGNOSTIC",
    "Запуск": "CROSSCHECK_DEEP_XLSX_RUN",
}
_ID_COLUMNS = {
    "Требования": 0,
    "Атомарные утверждения": 0,
    "Ответственность": 0,
    "Процедуры": 4,
    "Доказательства": 0,
}
_XLSX_ERRORS = (
    BadZipFile, InvalidFileException, ParseError, OSError, ValueError,
    KeyError, TypeError, IndexError,
)


class _DuplicateKey(ValueError):
    pass


def crosscheck_deep(
    run: DeepRunResult,
    json_path: Path,
    xlsx_path: Path,
    markdown_path: Path,
) -> tuple[CrosscheckIssue, ...]:
    """Collect all independently observable deep artifact mismatches."""
    validate_deep_run_result(run)
    for label, path in (("JSON", json_path), ("XLSX", xlsx_path), ("Markdown", markdown_path)):
        if not isinstance(path, Path):
            raise ValueError(f"Путь {label} должен быть Path.")

    issues: list[CrosscheckIssue] = []
    json_bytes = _read(json_path, "JSON", issues)
    xlsx_bytes = _read(xlsx_path, "XLSX", issues)
    markdown_bytes = _read(markdown_path, "MARKDOWN", issues)
    if json_bytes is not None:
        _check_json(run, json_bytes, issues)
    if xlsx_bytes is not None:
        _check_xlsx(run, xlsx_path, issues)
    if markdown_bytes is not None:
        _check_markdown(run, markdown_bytes, issues)
    return tuple(issues)


def _read(path: Path, label: str, issues: list[CrosscheckIssue]) -> bytes | None:
    if symlink_component(path) is not None:
        _issue(issues, f"CROSSCHECK_DEEP_{label}_READ", f"Артефакт {label} не должен быть symlink.")
        return None
    try:
        return path.read_bytes()
    except OSError:
        _issue(issues, f"CROSSCHECK_DEEP_{label}_READ", f"Артефакт {label} недоступен для чтения.")
        return None


def _check_json(run: DeepRunResult, payload: bytes, issues: list[CrosscheckIssue]) -> None:
    expected = _deep_run_payload(run)
    if payload != canonical_json_bytes(expected):
        _issue(issues, "CROSSCHECK_DEEP_JSON_BYTES", "JSON не совпадает с каноническими байтами DeepRunResult.")
    try:
        text = payload.decode("utf-8")
        parsed = json.loads(text, object_pairs_hook=_unique_object)
    except _DuplicateKey:
        _issue(issues, "CROSSCHECK_DEEP_JSON_DUPLICATE_KEY", "JSON содержит повторяющийся ключ.")
        return
    except (UnicodeDecodeError, json.JSONDecodeError):
        _issue(issues, "CROSSCHECK_DEEP_JSON_READ", "JSON не является строгим UTF-8 JSON-документом.")
        return
    if not isinstance(parsed, dict):
        _issue(issues, "CROSSCHECK_DEEP_JSON_SCHEMA", "Корень JSON должен быть объектом.")
        return
    if parsed != expected:
        _issue(issues, "CROSSCHECK_DEEP_JSON_VALUES", "JSON schema или значения не совпадают с DeepRunResult.")


def _check_xlsx(run: DeepRunResult, path: Path, issues: list[CrosscheckIssue]) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except _XLSX_ERRORS:
        _issue(issues, "CROSSCHECK_DEEP_XLSX_READ", "XLSX не читается как OOXML workbook.")
        return
    try:
        expected_rows = _expected_rows(run)
        if workbook.sheetnames != list(SHEET_HEADERS):
            _issue(issues, "CROSSCHECK_DEEP_XLSX_SHEETS", "Набор или порядок листов XLSX неканоничен.")
        for name, headers in SHEET_HEADERS.items():
            if name not in workbook.sheetnames:
                _issue(issues, _SHEET_CODES[name], f"Обязательный лист {name} отсутствует.")
                continue
            sheet = workbook[name]
            if sheet.sheet_state != "visible":
                _issue(issues, "CROSSCHECK_DEEP_XLSX_HIDDEN", f"Лист {name} скрыт.")
            actual_headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
            if actual_headers != headers:
                _issue(issues, "CROSSCHECK_DEEP_XLSX_HEADERS", f"Заголовки листа {name} изменены.")
            actual_rows = tuple(tuple(row) for row in sheet.iter_rows(min_row=2, values_only=True))
            if actual_rows != expected_rows[name]:
                _issue(issues, _SHEET_CODES[name], f"Строки листа {name} не совпадают с DeepRunResult.")
            if name in _ID_COLUMNS:
                column = _ID_COLUMNS[name]
                identifiers = tuple(row[column] for row in actual_rows if len(row) > column)
                if len(identifiers) != len(set(identifiers)):
                    _issue(issues, "CROSSCHECK_DEEP_XLSX_DUPLICATE_ID", f"Лист {name} содержит повторяющиеся ID.")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        _issue(issues, "CROSSCHECK_DEEP_XLSX_FORMULA", f"Лист {name} содержит формулу.")
    except _XLSX_ERRORS:
        _issue(issues, "CROSSCHECK_DEEP_XLSX_READ", "Нормализованная структура XLSX повреждена.")
    finally:
        workbook.close()


def _check_markdown(run: DeepRunResult, payload: bytes, issues: list[CrosscheckIssue]) -> None:
    try:
        text = payload.decode("utf-8")
    except UnicodeDecodeError:
        _issue(issues, "CROSSCHECK_DEEP_MARKDOWN_READ", "Markdown не является UTF-8 текстом.")
        return
    matches = _MARKER_PATTERN.findall(text)
    marker_valid = text.count(_MARKER_PREFIX) == 1 and len(matches) == 1
    if marker_valid:
        try:
            parsed = json.loads(matches[0], object_pairs_hook=_unique_object)
            canonical = json.dumps(
                _counts(run), ensure_ascii=False, sort_keys=True, separators=(",", ":")
            )
            marker_valid = parsed == _counts(run) and matches[0] == canonical
        except (_DuplicateKey, json.JSONDecodeError):
            marker_valid = False
    if not marker_valid:
        _issue(issues, "CROSSCHECK_DEEP_MARKDOWN_MARKER", "Markdown должен содержать один точный канонический marker counts.")

    headings = tuple(
        line[3:]
        for line in text.splitlines()
        if line.startswith("## ") and not line.startswith("### ")
    )
    if headings != _MARKDOWN_HEADINGS:
        _issue(
            issues,
            "CROSSCHECK_DEEP_MARKDOWN_STRUCTURE",
            "Порядок или набор Markdown-разделов не совпадает с контрактом.",
        )
    tables, malformed = _parse_markdown_tables(text)
    if malformed:
        _issue(
            issues,
            "CROSSCHECK_DEEP_MARKDOWN_STRUCTURE",
            "Markdown-таблицы имеют неканоническую структуру или escaping.",
        )
    expected = _markdown_expected_tables(run)
    expected_keys = set(expected)
    if set(tables) != expected_keys:
        _issue(
            issues,
            "CROSSCHECK_DEEP_MARKDOWN_STRUCTURE",
            "Набор Markdown-таблиц не совпадает с контрактом.",
        )
    for key, (code, rows) in expected.items():
        actual = tables.get(key)
        encoded = tuple(tuple(_markdown_cell(value) for value in row) for row in rows)
        if actual != encoded:
            _issue(issues, code, f"Таблица Markdown {key[0]} не совпадает с DeepRunResult.")
    if text.count(f"Статус запуска: `{run.run_status}`") != 1:
        _issue(
            issues,
            "CROSSCHECK_DEEP_MARKDOWN_SUMMARY",
            "Статус запуска Markdown не совпадает с DeepRunResult.",
        )


def _parse_markdown_tables(
    text: str,
) -> tuple[
    dict[tuple[str, tuple[str, ...]], tuple[tuple[str, ...], ...]],
    bool,
]:
    lines = text.splitlines()
    tables: dict[tuple[str, tuple[str, ...]], tuple[tuple[str, ...], ...]] = {}
    current_heading = ""
    malformed = False
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("## ") and not line.startswith("### "):
            current_heading = line[3:]
            index += 1
            continue
        if not line.startswith("| "):
            index += 1
            continue
        try:
            headers = _table_cells(line)
        except ValueError:
            malformed = True
            index += 1
            continue
        if index + 1 >= len(lines) or not _valid_alignment(lines[index + 1], len(headers)):
            malformed = True
            index += 1
            continue
        index += 2
        rows: list[tuple[str, ...]] = []
        while index < len(lines) and lines[index].startswith("| "):
            try:
                row = _table_cells(lines[index])
            except ValueError:
                malformed = True
                index += 1
                continue
            if len(row) != len(headers):
                malformed = True
            else:
                rows.append(row)
            index += 1
        key = (current_heading, headers)
        if key in tables or not current_heading:
            malformed = True
        else:
            tables[key] = tuple(rows)
    return tables, malformed


def _table_cells(line: str) -> tuple[str, ...]:
    if not line.startswith("| ") or not line.endswith(" |"):
        raise ValueError("noncanonical table row")
    body = line[2:-2]
    cells = tuple(body.split(" | "))
    if any(_decode_markdown_cell(cell) is None for cell in cells):
        raise ValueError("invalid Markdown escape")
    return cells


def _valid_alignment(line: str, columns: int) -> bool:
    if not line.startswith("|") or not line.endswith("|"):
        return False
    cells = line[1:-1].split("|")
    return len(cells) == columns and all(
        re.fullmatch(r":?-{3,}:?", cell) is not None
        for cell in cells
    )


def _decode_markdown_cell(value: str) -> str | None:
    decoded: list[str] = []
    index = 0
    while index < len(value):
        character = value[index]
        if character != "\\":
            decoded.append(character)
            index += 1
            continue
        if index + 1 >= len(value) or value[index + 1] not in {"\\", "|"}:
            return None
        decoded.append(value[index + 1])
        index += 2
    return "".join(decoded)


def _markdown_expected_tables(
    run: DeepRunResult,
) -> dict[
    tuple[str, tuple[str, ...]],
    tuple[str, tuple[tuple[object, ...], ...]],
]:
    counts = _counts(run)
    payload = _deep_run_payload(run)
    metadata = payload["metadata"]
    assert isinstance(metadata, dict)
    trust = metadata["knowledge_trust"]
    assert isinstance(trust, dict)
    records_by_id = {item.record_id: item for item in run.responsibility_records}
    evidence_by_id = {item.evidence_id: item for item in run.evidence}
    diagnostic_rows = _diagnostic_rows(run)

    summary_headers = ("Показатель", "Количество")
    summary_rows = (
        ("Исходные требования", counts["requirements"]),
        ("Атомарные утверждения", counts["atoms"]),
        ("Записи ответственности", counts["responsibilities"]),
        ("Шаги процедур", counts["procedure_steps"]),
        ("Доказательства", counts["evidence"]),
        ("Диагностические записи", counts["diagnostics"]),
    )
    trust_headers = ("Параметр", "Значение")
    trust_rows = (
        ("Профиль", metadata["analysis_profile"]),
        ("Модель", metadata["model"]),
        ("Snapshot ID", metadata["snapshot_id"]),
        ("Key ID", trust["key_id"]),
        ("Signer identity", trust["signer_identity"]),
        ("Manifest SHA-256", trust["manifest_sha256"]),
    )
    requirement_headers = (
        "Type", "Entity ID", "Source requirement IDs", "Group IDs",
        "Analysis state", "Support status", "Atom IDs", "Responsibility IDs",
        "Procedure graph IDs", "Component refs", "Diagnostics",
        "Full requirement text",
    )
    requirement_rows: list[tuple[object, ...]] = []
    for result in run.requirements:
        components = tuple(dict.fromkeys(
            records_by_id[item].component_ref for item in result.responsibility_ids
        ))
        requirement_rows.append((
            "requirement", result.requirement.requirement_id,
            _json_text((result.requirement.requirement_id,)),
            _json_text(result.requirement.group_ids), result.analysis_state.value,
            None if result.support_status is None else result.support_status.value,
            _json_text(tuple(item.atom.atom_id for item in result.atom_results)),
            _json_text(result.responsibility_ids), _json_text(result.procedure_graph_ids),
            _json_text(components), _json_text(result.diagnostics), result.requirement.text,
        ))
    requirement_rows.extend(
        (
            "group", group.group_id, _json_text(group.source_requirement_ids), "[]",
            _json_text(tuple(item.value for item in group.analysis_states)),
            None if group.support_status is None else group.support_status.value,
            "[]", _json_text(group.responsibility_ids), "[]",
            _json_text(group.component_refs), "[]", None,
        )
        for group in run.groups
    )
    atom_headers = (
        "Atom ID", "Requirement ID", "Ordinal", "Text", "Source quote",
        "Mandatory", "Analysis state", "Support status", "Responsibility IDs",
        "Supported aspects", "Unconfirmed aspects", "Diagnostics",
    )
    atom_rows = tuple(
        (
            item.atom.atom_id, item.atom.requirement_id, item.atom.ordinal,
            item.atom.text, item.atom.source_quote, item.atom.mandatory,
            item.analysis_state.value,
            None if item.support_status is None else item.support_status.value,
            _json_text(item.responsibility_ids), _json_text(item.supported_aspects),
            _json_text(item.unconfirmed_aspects), _json_text(item.diagnostics),
        )
        for result in run.requirements for item in result.atom_results
    )
    responsibility_headers = (
        "Record ID", "Requirement ID", "Atom ID", "Contour", "Component",
        "Executor", "Target contour", "Target", "Action", "Effect", "Phase",
        "Version scope", "Evidence IDs", "Support status", "Related records",
        "Procedure step IDs", "Diagnostics",
    )
    responsibility_rows = tuple(
        (
            item.record_id, item.requirement_id, item.atomic_claim_id,
            item.contour.value, item.component_ref, item.executor_ref,
            item.target_contour.value, item.target_ref, item.action_ref,
            item.effect_ref, item.lifecycle_phase.value,
            _json_text({
                "source_release": item.version_scope.source_release,
                "target_release": item.version_scope.target_release,
                "kolla_ansible_release": item.version_scope.kolla_ansible_release,
                "host_profile": item.version_scope.host_profile,
                "version_constraint": item.version_scope.version_constraint,
            }),
            _json_text(item.evidence_ids), item.support_status.value,
            _json_text(item.related_record_ids), _json_text(item.procedure_step_ids),
            _json_text(item.diagnostics),
        )
        for item in run.responsibility_records
    )
    contour_headers = ("Контур", "Количество записей")
    contour_rows = tuple(
        (contour.value, sum(item.contour is contour for item in run.responsibility_records))
        for contour in ResponsibilityContour
    )
    executor_headers = (
        "Record ID", "Связь", "Target contour", "Action", "Effect", "Related records",
    )
    executor_rows = tuple(
        (
            item.record_id, f"{item.executor_ref} → {item.target_ref}",
            item.target_contour.value, item.action_ref, item.effect_ref,
            _json_text(item.related_record_ids),
        )
        for item in run.responsibility_records
    )
    version_headers = (
        "Record ID", "Phase", "Source release", "Target release", "Kolla-Ansible",
        "Host profile", "Constraint",
    )
    version_rows = tuple(
        (
            item.record_id, item.lifecycle_phase.value,
            item.version_scope.source_release, item.version_scope.target_release,
            item.version_scope.kolla_ansible_release, item.version_scope.host_profile,
            item.version_scope.version_constraint,
        )
        for item in run.responsibility_records
    )
    procedure_headers = (
        "Graph ID", "Requirement ID", "Template ID", "Graph diagnostics",
        "Step ID", "Phase", "Contour", "Executor", "Target", "Action",
        "Preconditions", "Success criteria", "Evidence IDs", "Evidence versions",
        "Depends on", "Rollback step ID",
    )
    procedure_rows = tuple(
        (
            graph.graph_id, graph.requirement_id, graph.template_id,
            _json_text(graph.diagnostics), step.step_id, step.phase.value,
            step.contour.value, step.executor_ref, step.target_ref, step.action_ref,
            _json_text(step.preconditions), _json_text(step.success_criteria),
            _json_text(step.evidence_ids),
            _json_text(tuple(evidence_by_id[item].version_constraint for item in step.evidence_ids)),
            _json_text(step.depends_on), step.rollback_step_id,
        )
        for graph in run.procedure_graphs for step in graph.steps
    )
    evidence_headers = (
        "Evidence ID", "Claim", "Claim kind", "Polarity", "Strength", "Source ID",
        "Locator", "Version constraint", "Applicable contours",
        "Supports entity refs", "Local excerpt", "Review state",
    )
    evidence_rows = tuple(
        (
            item.evidence_id, item.claim, item.claim_kind, item.polarity.value,
            item.strength.value, item.source_id, item.locator, item.version_constraint,
            _json_text(tuple(contour.value for contour in item.applicable_contours)),
            _json_text(item.supports_entity_refs), item.local_excerpt, item.review_state,
        )
        for item in run.evidence
    )
    diagnostic_headers = (
        "Order", "Scope", "Entity ID", "Requirement ID", "Full requirement text",
        "Atom ID", "Source quote", "Diagnostic",
    )
    conflict_diagnostic_headers = (
        "Scope", "Entity ID", "Requirement ID", "Полный текст требования",
        "Atom ID", "Исходная цитата", "Диагностика",
    )
    conflict_diagnostics = tuple(
        row[1:] for row in diagnostic_rows if row[7] == "evidence_conflict"
    )
    conflict_headers = (
        "Evidence ID", "Polarity", "Strength", "Source ID", "Locator",
        "Version constraint", "Claim",
    )
    conflict_ids: list[str] = []
    for result in run.requirements:
        for atom in result.atom_results:
            if "evidence_conflict" not in atom.diagnostics:
                continue
            for record_id in atom.responsibility_ids:
                for evidence_id in records_by_id[record_id].evidence_ids:
                    if evidence_id not in conflict_ids:
                        conflict_ids.append(evidence_id)
    conflict_rows = tuple(
        (
            item.evidence_id, item.polarity.value, item.strength.value,
            item.source_id, item.locator, item.version_constraint, item.claim,
        )
        for item in (evidence_by_id[item] for item in conflict_ids if item in evidence_by_id)
        if item.strength.value == "direct" and item.polarity.value in {"positive", "negative"}
    )
    warning_headers = (
        "Scope", "Entity ID", "Requirement ID", "Полный текст требования",
        "Atom ID", "Исходная цитата", "Диагностика",
    )
    warning_rows = tuple(
        row[1:] for row in diagnostic_rows
        if str(row[7]).startswith(("procedure_gap", "rollback_unverified"))
    )
    problem_headers = (
        "Scope", "Entity ID", "Полный текст требования", "Atom ID",
        "Исходная цитата", "Статус", "Диагностика",
    )
    problem_rows: list[tuple[object, ...]] = []
    for result in run.requirements:
        if result.support_status in _PROBLEM_STATUSES or result.diagnostics:
            problem_rows.append((
                "requirement", result.requirement.requirement_id,
                result.requirement.text, None, None,
                None if result.support_status is None else result.support_status.value,
                _json_text(result.diagnostics),
            ))
        for atom in result.atom_results:
            if atom.support_status in _PROBLEM_STATUSES or atom.diagnostics:
                problem_rows.append((
                    "atom", atom.atom.atom_id, result.requirement.text,
                    atom.atom.atom_id, atom.atom.source_quote,
                    None if atom.support_status is None else atom.support_status.value,
                    _json_text(atom.diagnostics),
                ))
    failure_headers = (
        "Scope", "Entity ID", "Requirement ID", "Полный текст требования",
        "Atom ID", "Исходная цитата", "Причина",
    )
    failure_rows: list[tuple[object, ...]] = [
        row[1:] for row in diagnostic_rows
        if not str(row[7]).startswith(("evidence_conflict", "procedure_gap", "rollback_unverified"))
    ]
    failure_rows.extend(
        (
            "requirement_state", result.requirement.requirement_id,
            result.requirement.requirement_id, result.requirement.text,
            None, None, result.analysis_state.value,
        )
        for result in run.requirements
        if result.analysis_state is not AnalysisState.COMPLETED
    )
    return {
        ("Сводка", summary_headers): ("CROSSCHECK_DEEP_MARKDOWN_SUMMARY", summary_rows),
        ("Сводка", trust_headers): ("CROSSCHECK_DEEP_MARKDOWN_TRUST", trust_rows),
        ("Требования и группы", requirement_headers): ("CROSSCHECK_DEEP_MARKDOWN_REQUIREMENTS_GROUPS", tuple(requirement_rows)),
        ("Атомарные утверждения", atom_headers): ("CROSSCHECK_DEEP_MARKDOWN_ATOMS", atom_rows),
        ("Контуры ответственности", responsibility_headers): ("CROSSCHECK_DEEP_MARKDOWN_RESPONSIBILITIES", responsibility_rows),
        ("Контуры ответственности", contour_headers): ("CROSSCHECK_DEEP_MARKDOWN_CONTOURS", contour_rows),
        ("Executor → target", executor_headers): ("CROSSCHECK_DEEP_MARKDOWN_EXECUTOR_TARGET", executor_rows),
        ("Версии и область применимости", version_headers): ("CROSSCHECK_DEEP_MARKDOWN_VERSIONS", version_rows),
        ("Процедуры", procedure_headers): ("CROSSCHECK_DEEP_MARKDOWN_PROCEDURES", procedure_rows),
        ("Каталог evidence", evidence_headers): ("CROSSCHECK_DEEP_MARKDOWN_EVIDENCE", evidence_rows),
        ("Нормализованная диагностика", diagnostic_headers): ("CROSSCHECK_DEEP_MARKDOWN_DIAGNOSTICS", diagnostic_rows),
        ("Конфликты evidence", conflict_diagnostic_headers): ("CROSSCHECK_DEEP_MARKDOWN_CONFLICTS", conflict_diagnostics),
        ("Конфликты evidence", conflict_headers): ("CROSSCHECK_DEEP_MARKDOWN_CONFLICTS", conflict_rows),
        ("Предупреждения procedures и rollback", warning_headers): ("CROSSCHECK_DEEP_MARKDOWN_WARNINGS", warning_rows),
        ("Проблемные требования и атомы", problem_headers): ("CROSSCHECK_DEEP_MARKDOWN_PROBLEMS", tuple(problem_rows)),
        ("Ошибки обработки", failure_headers): ("CROSSCHECK_DEEP_MARKDOWN_FAILURES", tuple(failure_rows)),
    }


def _markdown_cell(value: object) -> str:
    text = "" if value is None else str(value)
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\r", "<br>")
        .replace("\n", "<br>")
    )


def _unique_object(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise _DuplicateKey(key)
        result[key] = value
    return result


def _counts(run: DeepRunResult) -> dict[str, int]:
    return {
        "requirements": len(run.requirements),
        "atoms": sum(len(item.atom_results) for item in run.requirements),
        "responsibilities": len(run.responsibility_records),
        "procedure_steps": sum(len(graph.steps) for graph in run.procedure_graphs),
        "evidence": len(run.evidence),
        "diagnostics": len(_diagnostic_rows(run)),
    }


def _expected_rows(run: DeepRunResult) -> dict[str, tuple[tuple[object, ...], ...]]:
    payload = _deep_run_payload(run)
    requirements = payload["requirements"]
    records = payload["responsibility_records"]
    graphs = payload["procedure_graphs"]
    evidence = payload["evidence"]
    metadata = payload["metadata"]
    groups = payload["groups"]
    assert isinstance(requirements, list)
    assert isinstance(records, list)
    assert isinstance(graphs, list)
    assert isinstance(evidence, list)
    assert isinstance(metadata, dict)
    trust = metadata["knowledge_trust"]
    assert isinstance(trust, dict)

    requirement_rows = []
    atom_rows = []
    for item in requirements:
        source = item["requirement"]
        coordinate = source["coordinate"]
        atom_results = item["atom_results"]
        requirement_rows.append((
            source["requirement_id"], source["source_id"], source["ordinal"],
            coordinate["source_name"], coordinate["sheet"], coordinate["row"],
            source["text"], source["parent_id"], _json_text(source["group_ids"]),
            _json_text(source["source_fields"]), _json_text(source["source_hints"]),
            item["analysis_state"], item["support_status"],
            _json_text([atom["atom"]["atom_id"] for atom in atom_results]),
            _json_text(item["responsibility_ids"]), _json_text(item["procedure_graph_ids"]),
            _json_text(item["diagnostics"]),
        ))
        for atom_result in atom_results:
            atom = atom_result["atom"]
            atom_rows.append((
                atom["atom_id"], atom["requirement_id"], atom["ordinal"], atom["text"],
                atom["source_quote"], atom["mandatory"], atom_result["analysis_state"],
                atom_result["support_status"], _json_text(atom_result["responsibility_ids"]),
                _json_text(atom_result["supported_aspects"]),
                _json_text(atom_result["unconfirmed_aspects"]),
                _json_text(atom_result["diagnostics"]),
            ))
    responsibility_rows = tuple((
        item["record_id"], item["requirement_id"], item["atomic_claim_id"], item["contour"],
        item["component_ref"], item["executor_ref"], item["target_contour"], item["target_ref"],
        item["action_ref"], item["effect_ref"], item["lifecycle_phase"],
        _json_text(item["version_scope"]), _json_text(item["evidence_ids"]),
        item["support_status"], _json_text(item["related_record_ids"]),
        _json_text(item["procedure_step_ids"]), _json_text(item["diagnostics"]),
    ) for item in records)
    procedure_rows = tuple((
        graph["graph_id"], graph["requirement_id"], graph["template_id"],
        _json_text(graph["diagnostics"]), step["step_id"], step["phase"], step["contour"],
        step["executor_ref"], step["target_ref"], step["action_ref"],
        _json_text(step["preconditions"]), _json_text(step["success_criteria"]),
        _json_text(step["evidence_ids"]), _json_text(step["depends_on"]),
        step["rollback_step_id"],
    ) for graph in graphs for step in graph["steps"])
    evidence_rows = tuple((
        item["evidence_id"], item["claim"], item["claim_kind"], item["polarity"],
        item["strength"], item["source_id"], item["locator"], item["version_constraint"],
        _json_text(item["applicable_contours"]), _json_text(item["supports_entity_refs"]),
        item["local_excerpt"], item["review_state"],
    ) for item in evidence)
    run_values = (
        ("run_id", payload["run_id"]), ("schema_version", payload["schema_version"]),
        ("run_status", payload["run_status"]), ("reqmap_version", metadata["reqmap_version"]),
        ("analysis_profile", metadata["analysis_profile"]), ("model", metadata["model"]),
        ("seed", metadata["seed"]), ("top_k", metadata["top_k"]),
        ("input_sha256", metadata["input_sha256"]), ("snapshot_id", metadata["snapshot_id"]),
        ("key_id", trust["key_id"]), ("manifest_sha256", trust["manifest_sha256"]),
        ("signer_identity", trust["signer_identity"]),
        ("prompt_versions", _json_text(metadata["prompt_versions"])),
        ("release_profile", _json_text(metadata["release_profile"])),
        ("retry_counts", _json_text(metadata["retry_counts"])),
        ("groups", _json_text(groups)), ("counts", _json_text(_counts(run))),
    )
    return {
        "Требования": tuple(requirement_rows),
        "Атомарные утверждения": tuple(atom_rows),
        "Ответственность": responsibility_rows,
        "Процедуры": procedure_rows,
        "Доказательства": evidence_rows,
        "Диагностика": _diagnostic_rows(run),
        "Запуск": run_values,
    }


def _diagnostic_rows(run: DeepRunResult) -> tuple[tuple[object, ...], ...]:
    rows: list[tuple[object, ...]] = []
    requirements = {item.requirement.requirement_id: item for item in run.requirements}
    atoms = {atom.atom.atom_id: atom for result in run.requirements for atom in result.atom_results}

    def append(scope: str, entity_id: str, requirement_id: str | None,
               atom_id: str | None, diagnostic: str) -> None:
        requirement = requirements.get(requirement_id or "")
        atom = atoms.get(atom_id or "")
        rows.append((
            len(rows) + 1, scope, entity_id, requirement_id,
            None if requirement is None else requirement.requirement.text,
            atom_id, None if atom is None else atom.atom.source_quote, diagnostic,
        ))

    for diagnostic in run.diagnostics:
        append("run", run.run_id, None, None, diagnostic)
    for result in run.requirements:
        requirement_id = result.requirement.requirement_id
        for diagnostic in result.diagnostics:
            append("requirement", requirement_id, requirement_id, None, diagnostic)
        for atom in result.atom_results:
            for diagnostic in atom.diagnostics:
                append("atom", atom.atom.atom_id, requirement_id, atom.atom.atom_id, diagnostic)
    for record in run.responsibility_records:
        for diagnostic in record.diagnostics:
            append("responsibility", record.record_id, record.requirement_id,
                   record.atomic_claim_id, diagnostic)
    for graph in run.procedure_graphs:
        for diagnostic in graph.diagnostics:
            append("procedure_graph", graph.graph_id, graph.requirement_id, None, diagnostic)
    return tuple(rows)


def _json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _issue(issues: list[CrosscheckIssue], code: str, message: str) -> None:
    issue = CrosscheckIssue(code, message)
    if issue not in issues:
        issues.append(issue)
