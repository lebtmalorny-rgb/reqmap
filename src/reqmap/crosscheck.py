"""Сверка JSON, XLSX и Markdown с canonical RunResult."""

from __future__ import annotations

from collections.abc import Mapping as MappingABC
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re
import tempfile
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from reqmap.export_json import (
    canonical_json_bytes,
    symlink_component,
    validate_run_result,
)
from reqmap.export_markdown import _canonical_counts, _markdown_bytes
from reqmap.export_xlsx import ExportError, verify_xlsx, write_xlsx
from reqmap.models import RunResult, to_dict


_COUNTS_PATTERN = re.compile(
    r"^<!-- reqmap-counts:(\{[^\r\n]*\}) -->$",
    re.MULTILINE,
)
_COUNT_CODES = {
    "requirements": "CROSSCHECK_REQUIREMENT_COUNT",
    "atoms": "CROSSCHECK_ATOM_COUNT",
    "mappings": "CROSSCHECK_MAPPING_COUNT",
    "evidence": "CROSSCHECK_EVIDENCE_COUNT",
}
_XLSX_ARTIFACT_ERRORS = (
    ExportError,
    BadZipFile,
    InvalidFileException,
    ParseError,
    OSError,
    ValueError,
    KeyError,
    TypeError,
)


@dataclass(frozen=True)
class CrosscheckIssue:
    code: str
    message_ru: str


def crosscheck(
    run: RunResult,
    json_path: Path,
    xlsx_path: Path,
    markdown_path: Path,
) -> tuple[CrosscheckIssue, ...]:
    """Вернуть все найденные расхождения артефактов, не скрывая соседние."""
    validate_run_result(run)
    for label, path in (
        ("JSON", json_path),
        ("XLSX", xlsx_path),
        ("Markdown", markdown_path),
    ):
        if not isinstance(path, Path):
            raise ValueError(f"Путь {label} должен быть Path.")

    issues: list[CrosscheckIssue] = []
    json_bytes = _read_artifact(json_path, "JSON", issues)
    xlsx_bytes = _read_artifact(xlsx_path, "XLSX", issues)
    markdown_bytes = _read_artifact(markdown_path, "MARKDOWN", issues)

    if json_bytes is not None:
        _check_json(run, json_bytes, issues)
    if xlsx_bytes is not None:
        _check_xlsx(run, xlsx_path, xlsx_bytes, issues)
    if markdown_bytes is not None:
        _check_markdown(run, markdown_bytes, issues)
    return tuple(issues)


def crosscheck_counts(
    run: RunResult,
    xlsx_counts: MappingABC[str, int],
    markdown_counts: MappingABC[str, int],
) -> tuple[CrosscheckIssue, ...]:
    """Сравнить машинные counts XLSX/Markdown с canonical RunResult."""
    validate_run_result(run)
    expected = _canonical_counts(run)
    issues: list[CrosscheckIssue] = []
    if dict(xlsx_counts) != expected:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_XLSX_COUNTS",
                "Количество записей XLSX не совпадает с JSON.",
            )
        )
        for key, expected_value in expected.items():
            if xlsx_counts.get(key) != expected_value:
                issues.append(
                    CrosscheckIssue(
                        _COUNT_CODES[key],
                        f"Количество {key} в XLSX не совпадает с JSON.",
                    )
                )
    if dict(markdown_counts) != expected:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_MARKDOWN_COUNTS",
                "Количество записей Markdown не совпадает с JSON.",
            )
        )
    return tuple(issues)


def _read_artifact(
    path: Path,
    label: str,
    issues: list[CrosscheckIssue],
) -> bytes | None:
    if symlink_component(path) is not None:
        issues.append(
            CrosscheckIssue(
                f"CROSSCHECK_{label}_READ",
                f"Артефакт {label} не должен быть symlink.",
            )
        )
        return None
    try:
        return path.read_bytes()
    except OSError:
        issues.append(
            CrosscheckIssue(
                f"CROSSCHECK_{label}_READ",
                f"Не удалось прочитать артефакт {label}.",
            )
        )
        return None


def _check_json(
    run: RunResult,
    payload: bytes,
    issues: list[CrosscheckIssue],
) -> None:
    expected_bytes = canonical_json_bytes(to_dict(run))
    if hashlib.sha256(payload).digest() != hashlib.sha256(expected_bytes).digest():
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_JSON_HASH",
                "SHA-256 JSON не совпадает с canonical RunResult.",
            )
        )
    try:
        parsed = json.loads(payload.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_JSON_READ",
                "JSON не является корректным UTF-8 JSON-документом.",
            )
        )
        return
    if not isinstance(parsed, dict):
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_JSON_SCHEMA",
                "Корень JSON должен быть объектом.",
            )
        )
        return

    if parsed.get("run_status") != run.run_status:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_JSON_RUN_STATUS",
                "Итоговый статус JSON не совпадает с RunResult.",
            )
        )
    json_counts = _json_counts(parsed)
    if json_counts != _canonical_counts(run):
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_JSON_COUNTS",
                "Количество записей JSON не совпадает с RunResult.",
            )
        )
    if _json_ids(parsed) != _canonical_ids(run):
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_JSON_IDS",
                "Идентификаторы JSON не совпадают с RunResult.",
            )
        )
    if _json_statuses(parsed) != _canonical_statuses(run):
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_JSON_STATUSES",
                "Статусы поддержки JSON не совпадают с RunResult.",
            )
        )


def _check_xlsx(
    run: RunResult,
    path: Path,
    payload: bytes,
    issues: list[CrosscheckIssue],
) -> None:
    try:
        verify_xlsx(path, run)
    except _XLSX_ARTIFACT_ERRORS:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_XLSX_VALUES",
                "Структура или значения XLSX не совпадают с RunResult.",
            )
        )
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        required = {
            "Требования",
            "Атомарные утверждения",
            "Сопоставления",
            "Доказательства",
            "Запуск",
        }
        if not required.issubset(workbook.sheetnames):
            raise ValueError("missing sheets")
        xlsx_counts = _xlsx_counts(workbook)
        issues.extend(crosscheck_counts(run, xlsx_counts, _canonical_counts(run)))
        if _xlsx_ids(workbook) != _canonical_ids(run):
            issues.append(
                CrosscheckIssue(
                    "CROSSCHECK_XLSX_IDS",
                    "Идентификаторы XLSX не совпадают с JSON.",
                )
            )
        if _xlsx_statuses(workbook) != _canonical_statuses(run):
            issues.append(
                CrosscheckIssue(
                    "CROSSCHECK_XLSX_STATUSES",
                    "Статусы поддержки XLSX не совпадают с JSON.",
                )
            )
        if _xlsx_run_status(workbook) != run.run_status:
            issues.append(
                CrosscheckIssue(
                    "CROSSCHECK_XLSX_RUN_STATUS",
                    "Итоговый статус XLSX не совпадает с JSON.",
                )
            )
    except _XLSX_ARTIFACT_ERRORS:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_XLSX_READ",
                "Не удалось прочитать нормализованную структуру XLSX.",
            )
        )
    finally:
        if workbook is not None:
            workbook.close()

    try:
        expected_digest = _expected_xlsx_digest(run)
    except (OSError, ValueError):
        expected_digest = None
    if expected_digest is None or hashlib.sha256(payload).hexdigest() != expected_digest:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_XLSX_HASH",
                "SHA-256 XLSX не совпадает с воспроизводимым артефактом.",
            )
        )


def _check_markdown(
    run: RunResult,
    payload: bytes,
    issues: list[CrosscheckIssue],
) -> None:
    expected = _markdown_bytes(run)
    if hashlib.sha256(payload).digest() != hashlib.sha256(expected).digest():
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_MARKDOWN_HASH",
                "SHA-256 Markdown не совпадает с воспроизводимым отчётом.",
            )
        )
    try:
        text = payload.decode("utf-8")
    except UnicodeDecodeError:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_MARKDOWN_READ",
                "Markdown не является корректным UTF-8 текстом.",
            )
        )
        return
    matches = _COUNTS_PATTERN.findall(text)
    marker_counts: object = None
    if len(matches) == 1:
        try:
            marker_counts = json.loads(matches[0])
        except json.JSONDecodeError:
            marker_counts = None
    issues.extend(
        crosscheck_counts(
            run,
            _canonical_counts(run),
            marker_counts if isinstance(marker_counts, dict) else {},
        )
    )
    if f"Статус запуска: `{run.run_status}`" not in text:
        issues.append(
            CrosscheckIssue(
                "CROSSCHECK_MARKDOWN_RUN_STATUS",
                "Итоговый статус Markdown не совпадает с JSON.",
            )
        )


def _expected_xlsx_digest(run: RunResult) -> str:
    temporary_root = Path(tempfile.gettempdir()).resolve()
    with tempfile.TemporaryDirectory(
        prefix="reqmap-crosscheck-",
        dir=temporary_root,
    ) as directory:
        return write_xlsx(run, Path(directory) / "expected.xlsx")


def _json_counts(payload: dict[str, object]) -> dict[str, int] | None:
    requirements = payload.get("requirements")
    evidence = payload.get("evidence")
    if not isinstance(requirements, list) or not isinstance(evidence, list):
        return None
    atoms = 0
    mappings = 0
    for result in requirements:
        if not isinstance(result, dict):
            return None
        atom_results = result.get("atom_results")
        result_mappings = result.get("mappings")
        if not isinstance(atom_results, list) or not isinstance(result_mappings, list):
            return None
        atoms += len(atom_results)
        mappings += len(result_mappings)
    return {
        "requirements": len(requirements),
        "atoms": atoms,
        "mappings": mappings,
        "evidence": len(evidence),
    }


def _canonical_ids(run: RunResult) -> dict[str, tuple[str, ...]]:
    return {
        "requirements": tuple(
            item.requirement.requirement_id for item in run.requirements
        ),
        "groups": tuple(item.group_id for item in run.groups),
        "atoms": tuple(
            atom_result.atom.atom_id
            for result in run.requirements
            for atom_result in result.atom_results
        ),
        "mappings": tuple(
            item.mapping_id
            for result in run.requirements
            for item in result.mappings
        ),
        "evidence": tuple(item.evidence_id for item in run.evidence),
    }


def _json_ids(payload: dict[str, object]) -> dict[str, tuple[str, ...]] | None:
    requirements = payload.get("requirements")
    groups = payload.get("groups")
    evidence = payload.get("evidence")
    if not all(isinstance(item, list) for item in (requirements, groups, evidence)):
        return None
    try:
        return {
            "requirements": tuple(
                item["requirement"]["requirement_id"] for item in requirements
            ),
            "groups": tuple(item["group_id"] for item in groups),
            "atoms": tuple(
                atom_result["atom"]["atom_id"]
                for item in requirements
                for atom_result in item["atom_results"]
            ),
            "mappings": tuple(
                mapping["mapping_id"]
                for item in requirements
                for mapping in item["mappings"]
            ),
            "evidence": tuple(item["evidence_id"] for item in evidence),
        }
    except (KeyError, TypeError):
        return None


def _canonical_statuses(run: RunResult) -> dict[str, tuple[tuple[str, str | None], ...]]:
    return {
        "requirements": tuple(
            (
                item.requirement.requirement_id,
                None if item.support_status is None else item.support_status.value,
            )
            for item in run.requirements
        ),
        "groups": tuple(
            (
                item.group_id,
                None if item.support_status is None else item.support_status.value,
            )
            for item in run.groups
        ),
        "atoms": tuple(
            (
                item.atom.atom_id,
                None if item.support_status is None else item.support_status.value,
            )
            for result in run.requirements
            for item in result.atom_results
        ),
        "mappings": tuple(
            (item.mapping_id, item.support_status.value)
            for result in run.requirements
            for item in result.mappings
        ),
    }


def _json_statuses(
    payload: dict[str, object],
) -> dict[str, tuple[tuple[str, str | None], ...]] | None:
    requirements = payload.get("requirements")
    groups = payload.get("groups")
    if not isinstance(requirements, list) or not isinstance(groups, list):
        return None
    try:
        return {
            "requirements": tuple(
                (
                    item["requirement"]["requirement_id"],
                    item["support_status"],
                )
                for item in requirements
            ),
            "groups": tuple(
                (item["group_id"], item["support_status"]) for item in groups
            ),
            "atoms": tuple(
                (
                    atom_result["atom"]["atom_id"],
                    atom_result["support_status"],
                )
                for item in requirements
                for atom_result in item["atom_results"]
            ),
            "mappings": tuple(
                (mapping["mapping_id"], mapping["support_status"])
                for item in requirements
                for mapping in item["mappings"]
            ),
        }
    except (KeyError, TypeError):
        return None


def _xlsx_counts(workbook) -> dict[str, int]:
    requirement_rows = tuple(
        row
        for row in workbook["Требования"].iter_rows(
            min_row=2,
            values_only=True,
        )
        if row[0] == "requirement"
    )
    return {
        "requirements": len(requirement_rows),
        "atoms": max(0, workbook["Атомарные утверждения"].max_row - 1),
        "mappings": max(0, workbook["Сопоставления"].max_row - 1),
        "evidence": max(0, workbook["Доказательства"].max_row - 1),
    }


def _xlsx_ids(workbook) -> dict[str, tuple[str, ...]]:
    requirement_rows = tuple(
        row
        for row in workbook["Требования"].iter_rows(
            min_row=2,
            values_only=True,
        )
    )
    return {
        "requirements": tuple(row[1] for row in requirement_rows if row[0] == "requirement"),
        "groups": tuple(row[1] for row in requirement_rows if row[0] == "group"),
        "atoms": _first_column(workbook["Атомарные утверждения"]),
        "mappings": _first_column(workbook["Сопоставления"]),
        "evidence": _first_column(workbook["Доказательства"]),
    }


def _xlsx_statuses(workbook) -> dict[str, tuple[tuple[str, str | None], ...]]:
    requirement_sheet = workbook["Требования"]
    requirement_headers = _header_columns(requirement_sheet)
    requirement_rows = tuple(
        row
        for row in requirement_sheet.iter_rows(min_row=2, values_only=True)
    )
    atom_sheet = workbook["Атомарные утверждения"]
    atom_headers = _header_columns(atom_sheet)
    mapping_sheet = workbook["Сопоставления"]
    mapping_headers = _header_columns(mapping_sheet)
    return {
        "requirements": tuple(
            (
                row[requirement_headers["Технический ID"]],
                row[requirement_headers["Поддержка: код"]],
            )
            for row in requirement_rows
            if row[requirement_headers["Тип строки"]] == "requirement"
        ),
        "groups": tuple(
            (
                row[requirement_headers["Технический ID"]],
                row[requirement_headers["Поддержка: код"]],
            )
            for row in requirement_rows
            if row[requirement_headers["Тип строки"]] == "group"
        ),
        "atoms": tuple(
            (
                row[atom_headers["Atom ID"]],
                row[atom_headers["Поддержка: код"]],
            )
            for row in atom_sheet.iter_rows(min_row=2, values_only=True)
        ),
        "mappings": tuple(
            (
                row[mapping_headers["Mapping ID"]],
                row[mapping_headers["Поддержка: код"]],
            )
            for row in mapping_sheet.iter_rows(min_row=2, values_only=True)
        ),
    }


def _xlsx_run_status(workbook) -> object:
    return dict(
        workbook["Запуск"].iter_rows(min_row=2, values_only=True)
    ).get("run_status")


def _first_column(sheet) -> tuple[str, ...]:
    return tuple(
        row[0]
        for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True)
    )


def _header_columns(sheet) -> dict[object, int]:
    return {
        cell.value: cell.column - 1
        for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    }
