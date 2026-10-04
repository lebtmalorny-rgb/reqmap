"""Deterministic normalized XLSX export for canonical schema-v2 results."""

from __future__ import annotations

from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from reqmap.deep_models import DeepRunResult
from reqmap.export_deep_json import _deep_run_payload, validate_deep_run_result
from reqmap.export_json import atomic_write_bytes, ensure_secure_directory, symlink_component
from reqmap.export_xlsx import (
    ExportError,
    _normalized_zip_bytes,
    _style_sheet,
    _temporary_path,
    write_literal,
)


SHEET_HEADERS = {
    "Требования": (
        "Requirement ID", "Source ID", "Ordinal", "Файл", "Лист", "Строка",
        "Текст требования", "Parent ID", "Group IDs", "Source fields",
        "Source hints", "Состояние: код", "Поддержка: код", "Atom IDs",
        "Responsibility IDs", "Procedure graph IDs", "Диагностика",
    ),
    "Атомарные утверждения": (
        "Atom ID", "Requirement ID", "Ordinal", "Формулировка атома",
        "Исходная цитата", "Обязательный", "Состояние: код",
        "Поддержка: код", "Responsibility IDs", "Подтверждённые аспекты",
        "Неподтверждённые аспекты", "Диагностика",
    ),
    "Ответственность": (
        "Record ID", "Requirement ID", "Atom ID", "Контур", "Component ref",
        "Executor ref", "Target contour", "Target ref", "Action ref", "Effect ref",
        "Lifecycle phase", "Version scope", "Evidence IDs", "Поддержка: код",
        "Related record IDs", "Procedure step IDs", "Диагностика",
    ),
    "Процедуры": (
        "Graph ID", "Requirement ID", "Template ID", "Graph diagnostics", "Step ID",
        "Phase", "Contour", "Executor ref", "Target ref", "Action ref",
        "Preconditions", "Success criteria", "Evidence IDs", "Depends on",
        "Rollback step ID",
    ),
    "Доказательства": (
        "Evidence ID", "Claim", "Claim kind", "Polarity", "Strength", "Source ID",
        "Locator", "Version constraint", "Applicable contours",
        "Supports entity refs", "Local excerpt", "Review state",
    ),
    "Диагностика": (
        "Порядок", "Scope", "Entity ID", "Requirement ID", "Текст требования",
        "Atom ID", "Исходная цитата", "Diagnostic",
    ),
    "Запуск": ("Параметр", "Значение"),
}

RUN_KEYS = (
    "run_id", "schema_version", "run_status", "reqmap_version",
    "analysis_profile", "model", "seed", "top_k", "input_sha256",
    "snapshot_id", "key_id", "manifest_sha256", "signer_identity",
    "prompt_versions", "release_profile", "retry_counts", "groups", "counts",
)

_ID_SHEETS = {
    "Требования": 0,
    "Атомарные утверждения": 0,
    "Ответственность": 0,
    "Процедуры": 4,
    "Доказательства": 0,
}


def write_deep_xlsx(run: DeepRunResult, path: Path) -> str:
    """Validate, independently verify and atomically publish a deep workbook."""
    validate_deep_run_result(run)
    if not isinstance(path, Path):
        raise ValueError("Выходной XLSX path должен быть Path.")
    symlink = symlink_component(path)
    if symlink is not None:
        raise ValueError(f"Выходной XLSX не может проходить через symlink: {symlink.name}")
    ensure_secure_directory(path.parent)
    workbook = _build_workbook(run)
    raw_path = _temporary_path(path.parent, ".deep-xlsx-raw-")
    verified_path = _temporary_path(path.parent, ".deep-xlsx-verify-")
    try:
        workbook.save(raw_path)
        normalized = _normalized_zip_bytes(raw_path)
        verified_path.write_bytes(normalized)
        os.chmod(verified_path, 0o600)
        verify_deep_xlsx(verified_path, run)
        atomic_write_bytes(path, normalized)
    finally:
        workbook.close()
        for temporary in (raw_path, verified_path):
            try:
                temporary.unlink()
            except OSError:
                pass
    return hashlib.sha256(normalized).hexdigest()


def verify_deep_xlsx(path: Path, run: DeepRunResult) -> None:
    """Reopen and verify the complete deep XLSX schema and canonical values."""
    validate_deep_run_result(run)
    if not isinstance(path, Path):
        raise ValueError("Проверяемый XLSX path должен быть Path.")
    if symlink_component(path) is not None:
        _fail("DEEP_XLSX_SYMLINK", "Выходной XLSX не может быть symlink.")
    try:
        with ZipFile(path) as archive:
            if archive.testzip() is not None:
                _fail("DEEP_XLSX_ZIP", "Выходной XLSX повреждён.")
    except (BadZipFile, OSError) as exc:
        raise ExportError("DEEP_XLSX_ZIP", "Выходной XLSX повреждён или недоступен.") from exc
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ExportError("DEEP_XLSX_OOXML", "Выходной XLSX не читается parser-ом.") from exc
    try:
        if workbook.sheetnames != list(SHEET_HEADERS):
            _fail("DEEP_XLSX_SHEETS", "Набор или порядок листов не соответствует схеме.")
        expected_rows = _data_rows(run)
        for name, headers in SHEET_HEADERS.items():
            sheet = workbook[name]
            if sheet.sheet_state != "visible":
                _fail("DEEP_XLSX_HIDDEN", f"Лист {name} не должен быть скрыт.")
            if tuple(cell.value for cell in sheet[1]) != headers:
                _fail("DEEP_XLSX_HEADERS", f"Заголовки листа {name} не соответствуют схеме.")
            actual = tuple(
                tuple(row)
                for row in sheet.iter_rows(min_row=2, values_only=True)
            )
            if actual != expected_rows[name]:
                _fail("DEEP_XLSX_VALUES", f"Значения листа {name} неканоничны.")
        _verify_no_formulas(workbook)
        _verify_unique_ids(workbook)
    finally:
        workbook.close()


def canonical_counts(run: DeepRunResult) -> dict[str, int]:
    """Return exact public deep artifact counts, including diagnostic occurrences."""
    return {
        "requirements": len(run.requirements),
        "atoms": sum(len(item.atom_results) for item in run.requirements),
        "responsibilities": len(run.responsibility_records),
        "procedure_steps": sum(len(graph.steps) for graph in run.procedure_graphs),
        "evidence": len(run.evidence),
        "diagnostics": len(_diagnostic_rows(run)),
    }


def _build_workbook(run: DeepRunResult) -> Workbook:
    workbook = Workbook()
    workbook.properties.creator = "reqmap"
    workbook.properties.lastModifiedBy = "reqmap"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    workbook.active.title = next(iter(SHEET_HEADERS))
    for name in tuple(SHEET_HEADERS)[1:]:
        workbook.create_sheet(name)
    for name, headers in SHEET_HEADERS.items():
        _append_row(workbook[name], headers)
    for name, rows in _data_rows(run).items():
        for row in rows:
            _append_row(workbook[name], row)
    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    return workbook


def _data_rows(run: DeepRunResult) -> dict[str, tuple[tuple[object, ...], ...]]:
    payload = _deep_run_payload(run)
    requirements = payload["requirements"]
    responsibilities = payload["responsibility_records"]
    graphs = payload["procedure_graphs"]
    evidence = payload["evidence"]
    assert isinstance(requirements, list)
    assert isinstance(responsibilities, list)
    assert isinstance(graphs, list)
    assert isinstance(evidence, list)
    return {
        "Требования": tuple(_requirement_row(item) for item in requirements),
        "Атомарные утверждения": tuple(
            _atom_row(atom)
            for item in requirements
            for atom in item["atom_results"]
        ),
        "Ответственность": tuple(_responsibility_row(item) for item in responsibilities),
        "Процедуры": tuple(
            _procedure_row(graph, step)
            for graph in graphs
            for step in graph["steps"]
        ),
        "Доказательства": tuple(_evidence_row(item) for item in evidence),
        "Диагностика": _diagnostic_rows(run),
        "Запуск": _run_rows(payload, run),
    }


def _requirement_row(item: dict[str, object]) -> tuple[object, ...]:
    source = item["requirement"]
    assert isinstance(source, dict)
    coordinate = source["coordinate"]
    assert isinstance(coordinate, dict)
    atom_results = item["atom_results"]
    assert isinstance(atom_results, list)
    return (
        source["requirement_id"], source["source_id"], source["ordinal"],
        coordinate["source_name"], coordinate["sheet"], coordinate["row"],
        source["text"], source["parent_id"], _json_text(source["group_ids"]),
        _json_text(source["source_fields"]), _json_text(source["source_hints"]),
        item["analysis_state"], item["support_status"],
        _json_text([atom["atom"]["atom_id"] for atom in atom_results]),
        _json_text(item["responsibility_ids"]), _json_text(item["procedure_graph_ids"]),
        _json_text(item["diagnostics"]),
    )


def _atom_row(item: dict[str, object]) -> tuple[object, ...]:
    atom = item["atom"]
    assert isinstance(atom, dict)
    return (
        atom["atom_id"], atom["requirement_id"], atom["ordinal"], atom["text"],
        atom["source_quote"], atom["mandatory"], item["analysis_state"],
        item["support_status"], _json_text(item["responsibility_ids"]),
        _json_text(item["supported_aspects"]), _json_text(item["unconfirmed_aspects"]),
        _json_text(item["diagnostics"]),
    )


def _responsibility_row(item: dict[str, object]) -> tuple[object, ...]:
    return (
        item["record_id"], item["requirement_id"], item["atomic_claim_id"],
        item["contour"], item["component_ref"], item["executor_ref"],
        item["target_contour"], item["target_ref"], item["action_ref"],
        item["effect_ref"], item["lifecycle_phase"], _json_text(item["version_scope"]),
        _json_text(item["evidence_ids"]), item["support_status"],
        _json_text(item["related_record_ids"]), _json_text(item["procedure_step_ids"]),
        _json_text(item["diagnostics"]),
    )


def _procedure_row(graph: dict[str, object], step: dict[str, object]) -> tuple[object, ...]:
    return (
        graph["graph_id"], graph["requirement_id"], graph["template_id"],
        _json_text(graph["diagnostics"]), step["step_id"], step["phase"],
        step["contour"], step["executor_ref"], step["target_ref"], step["action_ref"],
        _json_text(step["preconditions"]), _json_text(step["success_criteria"]),
        _json_text(step["evidence_ids"]), _json_text(step["depends_on"]),
        step["rollback_step_id"],
    )


def _evidence_row(item: dict[str, object]) -> tuple[object, ...]:
    return (
        item["evidence_id"], item["claim"], item["claim_kind"], item["polarity"],
        item["strength"], item["source_id"], item["locator"],
        item["version_constraint"], _json_text(item["applicable_contours"]),
        _json_text(item["supports_entity_refs"]), item["local_excerpt"],
        item["review_state"],
    )


def _diagnostic_rows(run: DeepRunResult) -> tuple[tuple[object, ...], ...]:
    rows: list[tuple[object, ...]] = []
    requirements = {item.requirement.requirement_id: item for item in run.requirements}
    atoms = {
        atom.atom.atom_id: atom
        for result in run.requirements
        for atom in result.atom_results
    }

    def append(scope: str, entity_id: str, requirement_id: str | None,
               atom_id: str | None, diagnostic: str) -> None:
        requirement = requirements.get(requirement_id or "")
        atom = atoms.get(atom_id or "")
        rows.append((
            len(rows) + 1, scope, entity_id, requirement_id,
            None if requirement is None else requirement.requirement.text,
            atom_id, None if atom is None else atom.atom.source_quote, diagnostic,
        ))

    for diagnostic in run.diagnostics:
        append("run", run.run_id, None, None, diagnostic)
    for result in run.requirements:
        requirement_id = result.requirement.requirement_id
        for diagnostic in result.diagnostics:
            append("requirement", requirement_id, requirement_id, None, diagnostic)
        for atom in result.atom_results:
            for diagnostic in atom.diagnostics:
                append("atom", atom.atom.atom_id, requirement_id, atom.atom.atom_id, diagnostic)
    for record in run.responsibility_records:
        for diagnostic in record.diagnostics:
            append("responsibility", record.record_id, record.requirement_id,
                   record.atomic_claim_id, diagnostic)
    for graph in run.procedure_graphs:
        for diagnostic in graph.diagnostics:
            append("procedure_graph", graph.graph_id, graph.requirement_id, None, diagnostic)
    return tuple(rows)


def _run_rows(payload: dict[str, object], run: DeepRunResult) -> tuple[tuple[str, object], ...]:
    metadata = payload["metadata"]
    groups = payload["groups"]
    assert isinstance(metadata, dict)
    trust = metadata["knowledge_trust"]
    assert isinstance(trust, dict)
    values: dict[str, object] = {
        "run_id": payload["run_id"],
        "schema_version": payload["schema_version"],
        "run_status": payload["run_status"],
        "reqmap_version": metadata["reqmap_version"],
        "analysis_profile": metadata["analysis_profile"],
        "model": metadata["model"],
        "seed": metadata["seed"],
        "top_k": metadata["top_k"],
        "input_sha256": metadata["input_sha256"],
        "snapshot_id": metadata["snapshot_id"],
        "key_id": trust["key_id"],
        "manifest_sha256": trust["manifest_sha256"],
        "signer_identity": trust["signer_identity"],
        "prompt_versions": _json_text(metadata["prompt_versions"]),
        "release_profile": _json_text(metadata["release_profile"]),
        "retry_counts": _json_text(metadata["retry_counts"]),
        "groups": _json_text(groups),
        "counts": _json_text(canonical_counts(run)),
    }
    rows = tuple((key, values[key]) for key in RUN_KEYS)
    return rows + ((("analysis_origin", _json_text(metadata["analysis_origin"])),) if "analysis_origin" in metadata else ())


def _append_row(sheet, values: tuple[object, ...]) -> None:
    row_number = 1 if sheet.max_row == 1 and sheet.cell(1, 1).value is None else sheet.max_row + 1
    for column, value in enumerate(values, 1):
        write_literal(sheet.cell(row_number, column), value)


def _json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _verify_no_formulas(workbook) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    _fail("DEEP_XLSX_FORMULA", "Выходной XLSX не должен содержать формул.")


def _verify_unique_ids(workbook) -> None:
    for sheet_name, column in _ID_SHEETS.items():
        identifiers = tuple(
            row[column]
            for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True)
        )
        if len(identifiers) != len(set(identifiers)):
            _fail("DEEP_XLSX_DUPLICATE_ID", f"Лист {sheet_name} содержит повторяющиеся ID.")


def _fail(code: str, message: str) -> None:
    raise ExportError(code, message)
