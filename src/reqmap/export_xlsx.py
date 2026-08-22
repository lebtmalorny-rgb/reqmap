"""Нормализованный русский XLSX exporter с независимой OOXML-проверкой."""

from __future__ import annotations

from datetime import datetime
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from reqmap.errors import ReqmapError
from reqmap.export_json import (
    atomic_write_bytes,
    ensure_secure_directory,
    symlink_component,
    validate_run_result,
)
from reqmap.models import (
    AnalysisState,
    AtomResult,
    Evidence,
    EvidencePolarity,
    EvidenceStrength,
    GroupResult,
    ImplementationSource,
    Mapping,
    Phase,
    RelationType,
    RequirementResult,
    RunResult,
    SupportStatus,
    to_dict,
)


SHEET_HEADERS = {
    "Требования": (
        "Тип строки",
        "Технический ID",
        "Исходный ID",
        "Файл",
        "Лист",
        "Строка",
        "Формулировка",
        "Исходные поля",
        "Parent ID",
        "Group IDs",
        "Source requirement IDs",
        "Source hints",
        "Состояние обработки: код",
        "Состояние обработки: русский",
        "Поддержка: код",
        "Поддержка: русский",
        "Компоненты",
    ),
    "Атомарные утверждения": (
        "Atom ID",
        "Requirement ID",
        "Формулировка атома",
        "Исходная цитата",
        "Обязательный",
        "Состояние: код",
        "Состояние: русский",
        "Поддержка: код",
        "Поддержка: русский",
        "Подтверждённые аспекты",
        "Неподтверждённые аспекты",
    ),
    "Сопоставления": (
        "Mapping ID",
        "Atom ID",
        "Component ID",
        "Роль",
        "Тип связи: код",
        "Тип связи: русский",
        "Фаза: код",
        "Фаза: русский",
        "Источник реализации: код",
        "Источник реализации: русский",
        "Механизм",
        "Шаги",
        "Evidence IDs",
        "Поддержка: код",
        "Поддержка: русский",
        "Обоснование",
    ),
    "Доказательства": (
        "Evidence ID",
        "Component ID",
        "Capability ID",
        "Полярность: код",
        "Полярность: русский",
        "Сила: код",
        "Сила: русский",
        "Утверждение",
        "URL",
        "Локальный путь",
        "Локатор",
        "Версия",
        "SHA-256",
        "Происхождение",
    ),
    "Запуск": ("Параметр", "Значение"),
}

_ANALYSIS_RU = {
    AnalysisState.COMPLETED: "обработка завершена",
    AnalysisState.MODEL_FAILED: "ошибка модели",
    AnalysisState.VALIDATION_FAILED: "ошибка валидации",
    AnalysisState.SKIPPED: "пропущено",
}
_SUPPORT_RU = {
    SupportStatus.SUPPORTED: "поддерживается",
    SupportStatus.PARTIAL: "поддерживается частично",
    SupportStatus.NOT_SUPPORTED: "не поддерживается",
    SupportStatus.INSUFFICIENT_EVIDENCE: "недостаточно доказательств",
    SupportStatus.NOT_APPLICABLE: "неприменимо",
}
_RELATION_RU = {
    RelationType.IMPLEMENTS: "реализует",
    RelationType.CONFIGURES: "настраивает",
    RelationType.PREREQUISITE: "предварительное условие",
    RelationType.INTEGRATES: "интегрирует",
    RelationType.HOST_OS_CHANGE: "изменение хостовой ОС",
}
_PHASE_RU = {
    Phase.RUNTIME: "выполнение во время эксплуатации",
    Phase.DESIGNTIME: "настройка при развёртывании",
}
_IMPLEMENTATION_SOURCE_RU = {
    ImplementationSource.UPSTREAM: "базовый OpenStack",
    ImplementationSource.KOLLA_ANSIBLE: "Kolla-Ansible",
    ImplementationSource.PRODUCT_EXTENSION: "расширение продукта",
    ImplementationSource.EXTERNAL_COMPONENT: "внешний компонент",
}
_EVIDENCE_POLARITY_RU = {
    EvidencePolarity.POSITIVE: "положительное",
    EvidencePolarity.NEGATIVE: "отрицательное",
}
_EVIDENCE_STRENGTH_RU = {
    EvidenceStrength.DIRECT: "прямое",
    EvidenceStrength.INDIRECT: "косвенное",
    EvidenceStrength.NONE: "отсутствует",
}
_CONFIRMED_STATUSES = frozenset(
    {
        SupportStatus.SUPPORTED,
        SupportStatus.PARTIAL,
        SupportStatus.NOT_SUPPORTED,
    }
)
_FIXED_ZIP_TIME = (2000, 1, 1, 0, 0, 0)
_FIXED_CORE_TIME = b"2000-01-01T00:00:00Z"
_CORE_MODIFIED_OPEN = b'<dcterms:modified xsi:type="dcterms:W3CDTF">'
_CORE_MODIFIED_CLOSE = b"</dcterms:modified>"


class ExportError(ReqmapError):
    """Выходная книга нарушает опубликованный XLSX/OOXML-контракт."""


def write_xlsx(run: RunResult, path: Path) -> str:
    """Создать, независимо проверить и атомарно опубликовать XLSX."""
    validate_run_result(run)
    symlink = symlink_component(path)
    if symlink is not None:
        raise ValueError(
            f"Выходной XLSX не может проходить через symlink: {symlink.name}"
        )
    ensure_secure_directory(path.parent)
    workbook = _build_workbook(run)
    raw_path = _temporary_path(path.parent, ".xlsx-raw-")
    verified_path = _temporary_path(path.parent, ".xlsx-verify-")
    try:
        workbook.save(raw_path)
        normalized = _normalized_zip_bytes(raw_path)
        verified_path.write_bytes(normalized)
        os.chmod(verified_path, 0o600)
        verify_xlsx(verified_path, run)
        atomic_write_bytes(path, normalized)
    finally:
        workbook.close()
        for temporary in (raw_path, verified_path):
            try:
                temporary.unlink()
            except OSError:
                pass
    return hashlib.sha256(normalized).hexdigest()


def verify_xlsx(path: Path, run: RunResult) -> None:
    """Проверить ZIP/OOXML, структуру, значения и отсутствие формул."""
    validate_run_result(run)
    if symlink_component(path) is not None:
        _fail("XLSX_SYMLINK", "Выходной XLSX не может быть symlink.")
    try:
        with ZipFile(path) as archive:
            if archive.testzip() is not None:
                _fail("XLSX_ZIP_INVALID", "Выходной XLSX повреждён.")
    except (BadZipFile, OSError) as exc:
        raise ExportError(
            "XLSX_ZIP_INVALID",
            "Выходной XLSX повреждён или недоступен.",
        ) from exc

    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except (OSError, ValueError, BadZipFile) as exc:
        raise ExportError(
            "XLSX_OOXML_INVALID",
            "Выходной XLSX не читается независимым parser-ом.",
        ) from exc
    try:
        _verify_sheet_names(workbook.sheetnames)
        _verify_headers(workbook)
        _verify_counts(workbook, run)
        _verify_no_formulas(workbook)
        _verify_rows(workbook, run)
    finally:
        workbook.close()


def write_literal(cell: Cell, value: object) -> None:
    """Записать пользовательскую строку literal без Excel formula semantics."""
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
        return
    cell.value = value


def analysis_state_ru(state: AnalysisState) -> str:
    if type(state) is not AnalysisState:
        raise ValueError("Неканонический analysis_state.")
    return _ANALYSIS_RU[state]


def support_status_ru(status: SupportStatus) -> str:
    if type(status) is not SupportStatus:
        raise ValueError("Неканонический support_status.")
    return _SUPPORT_RU[status]


def implementation_source_ru(source: ImplementationSource) -> str:
    if type(source) is not ImplementationSource:
        raise ValueError("Неканонический implementation_source.")
    return _IMPLEMENTATION_SOURCE_RU[source]


def evidence_polarity_ru(polarity: EvidencePolarity) -> str:
    if type(polarity) is not EvidencePolarity:
        raise ValueError("Неканоническая evidence polarity.")
    return _EVIDENCE_POLARITY_RU[polarity]


def evidence_strength_ru(strength: EvidenceStrength) -> str:
    if type(strength) is not EvidenceStrength:
        raise ValueError("Неканоническая evidence strength.")
    return _EVIDENCE_STRENGTH_RU[strength]


def _build_workbook(run: RunResult) -> Workbook:
    workbook = Workbook()
    workbook.properties.creator = "reqmap"
    workbook.properties.lastModifiedBy = "reqmap"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    first = workbook.active
    first.title = next(iter(SHEET_HEADERS))
    for name in tuple(SHEET_HEADERS)[1:]:
        workbook.create_sheet(name)

    for name, headers in SHEET_HEADERS.items():
        _append_row(workbook[name], headers)
    for sheet_name, rows in _data_rows(run).items():
        for row in rows:
            _append_row(workbook[sheet_name], row)

    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    return workbook


def _requirement_row(result: RequirementResult) -> tuple[object, ...]:
    requirement = result.requirement
    component_ids = tuple(
        sorted(
            {
                item.component_id
                for item in result.mappings
                if item.support_status in _CONFIRMED_STATUSES
            }
        )
    )
    return (
        "requirement",
        requirement.requirement_id,
        requirement.source_id,
        requirement.coordinate.source_name,
        requirement.coordinate.sheet,
        requirement.coordinate.row,
        requirement.text,
        _json_text(requirement.source_fields),
        requirement.parent_id,
        _json_text(requirement.group_ids),
        None,
        _json_text(requirement.source_hints),
        result.analysis_state.value,
        analysis_state_ru(result.analysis_state),
        _status_code(result.support_status),
        _status_ru(result.support_status),
        _json_text(component_ids),
    )


def _group_row(group: GroupResult) -> tuple[object, ...]:
    return (
        "group",
        group.group_id,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        _json_text(group.source_requirement_ids),
        None,
        _json_text(tuple(item.value for item in group.analysis_states)),
        _json_text(tuple(analysis_state_ru(item) for item in group.analysis_states)),
        _status_code(group.support_status),
        _status_ru(group.support_status),
        _json_text(group.component_ids),
    )


def _atom_row(atom_result: AtomResult) -> tuple[object, ...]:
    return (
        atom_result.atom.atom_id,
        atom_result.atom.requirement_id,
        atom_result.atom.text,
        atom_result.atom.source_quote,
        atom_result.atom.mandatory,
        atom_result.analysis_state.value,
        analysis_state_ru(atom_result.analysis_state),
        _status_code(atom_result.support_status),
        _status_ru(atom_result.support_status),
        _json_text(atom_result.supported_aspects),
        _json_text(atom_result.unconfirmed_aspects),
    )


def _mapping_row(item: Mapping) -> tuple[object, ...]:
    return (
        item.mapping_id,
        item.atom_id,
        item.component_id,
        item.role_ru,
        item.relation.value,
        _RELATION_RU[item.relation],
        item.phase.value,
        _PHASE_RU[item.phase],
        item.implementation_source.value,
        implementation_source_ru(item.implementation_source),
        item.mechanism,
        _json_text(item.steps),
        _json_text(item.evidence_ids),
        item.support_status.value,
        support_status_ru(item.support_status),
        item.reason_ru,
    )


def _evidence_row(item: Evidence) -> tuple[object, ...]:
    return (
        item.evidence_id,
        item.component_id,
        item.capability_id,
        item.polarity.value,
        evidence_polarity_ru(item.polarity),
        item.strength.value,
        evidence_strength_ru(item.strength),
        item.claim_ru,
        item.source_url,
        item.local_path,
        item.locator,
        item.version_constraint,
        item.source_sha256,
        item.provenance,
    )


def _data_rows(run: RunResult) -> dict[str, tuple[tuple[object, ...], ...]]:
    return {
        "Требования": tuple(
            [*(_requirement_row(item) for item in run.requirements)]
            + [*(_group_row(item) for item in run.groups)]
        ),
        "Атомарные утверждения": tuple(
            _atom_row(atom_result)
            for result in run.requirements
            for atom_result in result.atom_results
        ),
        "Сопоставления": tuple(
            _mapping_row(item)
            for result in run.requirements
            for item in result.mappings
        ),
        "Доказательства": tuple(_evidence_row(item) for item in run.evidence),
        "Запуск": _run_rows(run),
    }


def _run_rows(run: RunResult) -> tuple[tuple[str, object], ...]:
    metadata = run.metadata
    return (
        ("schema_version", run.schema_version),
        ("reqmap_version", metadata.get("reqmap_version")),
        ("knowledge_version", metadata.get("knowledge_version")),
        ("knowledge_sha256", metadata.get("knowledge_sha256")),
        ("model", metadata.get("model")),
        ("endpoint_origin", metadata.get("endpoint_origin")),
        ("prompt_versions", _json_text(metadata.get("prompt_versions", {}))),
        ("input_sha256", metadata.get("input_sha256")),
        ("requirements_count", len(run.requirements)),
        (
            "atoms_count",
            sum(len(item.atom_results) for item in run.requirements),
        ),
        ("mappings_count", sum(len(item.mappings) for item in run.requirements)),
        ("groups_count", len(run.groups)),
        ("evidence_count", len(run.evidence)),
        ("run_status", run.run_status),
    )


def _append_row(sheet: Worksheet, values: tuple[object, ...]) -> None:
    row = (
        1
        if sheet.max_row == 1 and sheet.cell(row=1, column=1).value is None
        else sheet.max_row + 1
    )
    for column, value in enumerate(values, start=1):
        write_literal(sheet.cell(row=row, column=column), value)


def _style_sheet(sheet: Worksheet) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cells in sheet.columns:
        values = ("" if cell.value is None else str(cell.value) for cell in cells)
        width = min(60, max(12, max((len(value) for value in values), default=0) + 2))
        sheet.column_dimensions[cells[0].column_letter].width = width


def _json_text(value: object) -> str:
    return json.dumps(
        to_dict(value),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )


def _status_code(status: SupportStatus | None) -> str | None:
    return None if status is None else status.value


def _status_ru(status: SupportStatus | None) -> str | None:
    return None if status is None else support_status_ru(status)


def _temporary_path(parent: Path, prefix: str) -> Path:
    descriptor, name = tempfile.mkstemp(prefix=prefix, suffix=".xlsx", dir=parent)
    os.fchmod(descriptor, 0o600)
    os.close(descriptor)
    return Path(name)


def _normalized_zip_bytes(path: Path) -> bytes:
    output = BytesIO()
    with ZipFile(path) as source, ZipFile(
        output,
        mode="w",
        compression=ZIP_DEFLATED,
        compresslevel=9,
    ) as target:
        for original in sorted(source.infolist(), key=lambda item: item.filename):
            normalized = ZipInfo(original.filename, _FIXED_ZIP_TIME)
            normalized.compress_type = (
                ZIP_STORED if original.is_dir() else ZIP_DEFLATED
            )
            normalized.external_attr = original.external_attr
            normalized.flag_bits = original.flag_bits & 0x800
            normalized.create_system = 0
            payload = source.read(original.filename)
            if original.filename == "docProps/core.xml":
                payload = _normalized_core_properties(payload)
            target.writestr(normalized, payload)
    return output.getvalue()


def _normalized_core_properties(payload: bytes) -> bytes:
    if (
        payload.count(_CORE_MODIFIED_OPEN) != 1
        or payload.count(_CORE_MODIFIED_CLOSE) != 1
    ):
        raise ValueError(
            "OOXML core properties не содержит единственную modified timestamp."
        )
    before, _separator, remainder = payload.partition(_CORE_MODIFIED_OPEN)
    _timestamp, _separator, after = remainder.partition(_CORE_MODIFIED_CLOSE)
    return (
        before
        + _CORE_MODIFIED_OPEN
        + _FIXED_CORE_TIME
        + _CORE_MODIFIED_CLOSE
        + after
    )


def _verify_sheet_names(names: list[str]) -> None:
    if names != list(SHEET_HEADERS):
        _fail("XLSX_SHEETS", "Набор или порядок листов XLSX не соответствует схеме.")


def _verify_headers(workbook: Workbook) -> None:
    for name, expected in SHEET_HEADERS.items():
        actual = tuple(cell.value for cell in workbook[name][1])
        if actual != expected:
            _fail("XLSX_HEADERS", f"Заголовки листа {name} не соответствуют схеме.")


def _verify_counts(workbook: Workbook, run: RunResult) -> None:
    for name, rows in _data_rows(run).items():
        if workbook[name].max_row != 1 + len(rows):
            _fail("XLSX_COUNT", f"Число строк листа {name} не совпадает с RunResult.")


def _verify_rows(workbook: Workbook, run: RunResult) -> None:
    for sheet_name, expected in _data_rows(run).items():
        sheet = workbook[sheet_name]
        actual = tuple(
            tuple(
                sheet.cell(row=row, column=column).value
                for column in range(1, len(SHEET_HEADERS[sheet_name]) + 1)
            )
            for row in range(2, sheet.max_row + 1)
        )
        if actual != expected:
            _fail(
                "XLSX_VALUES",
                f"Значения листа {sheet_name} не совпадают с RunResult.",
            )


def _verify_no_formulas(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    _fail("XLSX_FORMULA", "Выходной XLSX не должен содержать формул.")


def _fail(code: str, message: str) -> None:
    raise ExportError(code, message)
