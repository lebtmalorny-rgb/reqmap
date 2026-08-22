"""Безопасная неизменяемая загрузка требований из XLSX."""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from reqmap.config import InputProfile
from reqmap.errors import InputProfileError
from reqmap.ids import generated_requirement_id
from reqmap.models import Requirement, SourceCoordinate, SourceField, SourceHint


_ID_HEADERS = frozenset({"ID", "Идентификатор", "Код"})
_TEXT_HEADERS = frozenset(
    {"Требование", "Формулировка требования", "Описание требования"}
)


def detect_profile(workbook: Workbook) -> InputProfile:
    """Автоматически определяет только однозначный профиль первого заголовка."""
    id_matches: list[tuple[str, str]] = []
    text_matches: list[tuple[str, str]] = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        for cell in next(sheet.iter_rows(min_row=1, max_row=1), ()):
            if cell.value in _ID_HEADERS:
                id_matches.append((sheet.title, cell.value))
            if cell.value in _TEXT_HEADERS:
                text_matches.append((sheet.title, cell.value))

    if (
        len(id_matches) != 1
        or len(text_matches) != 1
        or id_matches[0][0] != text_matches[0][0]
    ):
        raise _ambiguous_profile_error()

    sheet_name, id_column = id_matches[0]
    _, text_column = text_matches[0]
    return InputProfile(
        include_sheets=(sheet_name,),
        exclude_sheets=(),
        header_row=1,
        id_column=id_column,
        text_column=text_column,
    )


def load_xlsx(path: Path, profile: InputProfile | None) -> tuple[Requirement, ...]:
    """Загружает XLSX read-only, сохраняя порядок, координаты и исходный контекст."""
    source_bytes = path.read_bytes()
    requirements = load_xlsx_bytes(source_bytes, profile, path.name)
    if path.read_bytes() != source_bytes:
        raise InputProfileError(
            "XLSX_SOURCE_CHANGED",
            "Исходный XLSX изменился во время импорта; результат отклонён.",
        )
    return requirements


def load_xlsx_bytes(
    source_bytes: bytes,
    profile: InputProfile | None,
    source_name: str,
) -> tuple[Requirement, ...]:
    """Разбирает ровно тот byte snapshot XLSX, для которого вычисляется hash."""
    if type(source_bytes) is not bytes or not source_bytes:
        raise InputProfileError(
            "XLSX_SOURCE_INVALID",
            "Исходный XLSX должен содержать непустой byte snapshot.",
        )
    if type(source_name) is not str or not source_name:
        raise InputProfileError(
            "XLSX_SOURCE_INVALID",
            "Для исходного XLSX требуется непустое имя файла.",
        )
    workbook = load_workbook(
        BytesIO(source_bytes),
        read_only=True,
        data_only=False,
    )
    try:
        effective_profile = profile if profile is not None else detect_profile(workbook)
        requirements = _load_profiled_workbook(
            source_name,
            workbook,
            effective_profile,
        )
    finally:
        workbook.close()
    return tuple(requirements)


def _load_profiled_workbook(
    source_name: str,
    workbook: Workbook,
    profile: InputProfile,
) -> list[Requirement]:
    selected_sheets = _select_sheets(workbook, profile)
    requirements: list[Requirement] = []
    for sheet in selected_sheets:
        columns = _header_columns(sheet, profile)
        for row in sheet.iter_rows(min_row=profile.header_row + 1):
            text_cell = row[columns[profile.text_column]]
            if text_cell.data_type == "f":
                raise InputProfileError(
                    "XLSX_TEXT_FORMULA",
                    "В ячейке текста требования обнаружена формула; укажите исходный текст.",
                    {"sheet": sheet.title, "row": text_cell.row},
                )
            if text_cell.value is None or text_cell.value == "":
                continue

            ordinal = len(requirements) + 1
            source_id = _optional_value(row[columns[profile.id_column]].value)
            source_fields = tuple(
                SourceField(column=column, value=_string_value(row[columns[column]].value))
                for column in _source_field_columns(profile)
            )
            source_hints = tuple(
                SourceHint(value=_string_value(row[columns[column]].value), column=column)
                for column in profile.hint_columns
            )
            requirements.append(
                Requirement(
                    requirement_id=generated_requirement_id(ordinal),
                    source_id=source_id,
                    text=_string_value(text_cell.value),
                    ordinal=ordinal,
                    coordinate=SourceCoordinate(
                        source_name=source_name,
                        sheet=sheet.title,
                        row=text_cell.row,
                    ),
                    source_fields=source_fields,
                    source_hints=source_hints,
                )
            )
    return requirements


def _select_sheets(workbook: Workbook, profile: InputProfile) -> list[object]:
    requested = set(profile.include_sheets) - set(profile.exclude_sheets)
    sheets = [sheet for sheet in workbook.worksheets if sheet.title in requested]
    if len(sheets) != len(requested):
        missing = ", ".join(sorted(requested - {sheet.title for sheet in sheets}))
        raise InputProfileError(
            "XLSX_PROFILE_SHEET_MISSING",
            f"Во входном XLSX не найдены листы профиля: {missing}.",
        )
    return sheets


def _header_columns(sheet: object, profile: InputProfile) -> dict[str, int]:
    header_cells = next(
        sheet.iter_rows(min_row=profile.header_row, max_row=profile.header_row), ()
    )
    header_indexes: dict[str, list[int]] = {}
    for index, cell in enumerate(header_cells):
        if isinstance(cell.value, str):
            header_indexes.setdefault(cell.value, []).append(index)

    required = tuple(
        dict.fromkeys(
            [
                profile.id_column,
                profile.text_column,
                *_source_field_columns(profile),
                *profile.hint_columns,
            ]
        )
    )
    missing = [column for column in required if column not in header_indexes]
    if missing:
        raise InputProfileError(
            "XLSX_PROFILE_COLUMN_MISSING",
            "В листе "
            f"{sheet.title!r} не найдены заголовки профиля: {', '.join(missing)}.",
        )
    ambiguous = [column for column in required if len(header_indexes[column]) != 1]
    if ambiguous:
        raise InputProfileError(
            "XLSX_PROFILE_COLUMN_AMBIGUOUS",
            "В листе "
            f"{sheet.title!r} повторяются заголовки профиля: {', '.join(ambiguous)}.",
        )
    return {column: header_indexes[column][0] for column in required}


def _source_field_columns(profile: InputProfile) -> tuple[str, ...]:
    return tuple(
        column
        for column in (
            profile.priority_column,
            profile.expected_result_column,
            profile.parent_column,
        )
        if column is not None
    )


def _optional_value(value: object) -> str | None:
    if value is None or value == "":
        return None
    return _string_value(value)


def _string_value(value: object) -> str:
    return "" if value is None else str(value)


def _ambiguous_profile_error() -> InputProfileError:
    return InputProfileError(
        "XLSX_PROFILE_AMBIGUOUS",
        "Структура XLSX неоднозначна. Укажите лист, строку заголовка и колонки во входном профиле.",
    )
